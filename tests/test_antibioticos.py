"""Antibióticos: la lista automática y el libro de control para Salubridad."""

from datetime import date, timedelta
from decimal import Decimal as D

import pytest

from app.models import Entrada, EntradaRenglon, Lote, Producto, Proveedor
from app.services.antibioticos import es_antibiotico
from tests.test_devoluciones import cancelar, devolver, pieza
from tests.test_ventas import caja, producto, r, shampoo, vender  # noqa: F401  (fixtures)

HOY = date.today()


@pytest.mark.parametrize("nombre, si", [
    ("BRUBIOL TABS 500 MG C/10 (CIPROFLOXACINO)", True),
    ("AMOXIL 12H SUSP 400MG 50ML", True),
    ("3651 VALCLAN 500MG AMOX/AC CLAVULANICO C/10 TABS", True),
    ("*AMPICILINA 500 X 20 CAP. [SALUCOM]", True),
    ("Cefalexina 500 mg", True),
    ("GARAMICINA SOL INY 160MG C/5 AMP", True),
    ("SOLTRIM TABS 80MG/400MG C/20 (TRIMETROPRIMA/SULFAM", True),
    ("ERITROPOYETINA HUMANA RECOMBINANTE 2000UI/ML", False),
    ("FLOROGLUCINOL/TRIMETILFLOROGLUCINOL CAPS C/20", False),
    ("CIPROHEPTADINA JARABE", False),
    ("NORFENON TBS 300MG C/30", False),
    ("PARACETAMOL 500 MG", False),
])
def test_reconoce_antibioticos_por_nombre(nombre, si):
    assert es_antibiotico(nombre) is si


def test_se_marca_solo_al_crear_y_renombrar(db, negocio):
    p = Producto(negocio_id=negocio.id, nombre="AMOXICILINA 500 MG C/12")
    otro = Producto(negocio_id=negocio.id, nombre="PARACETAMOL 500 MG")
    db.add_all([p, otro])
    db.commit()
    assert (p.antibiotico, p.requiere_receta, otro.antibiotico) == (True, True, False)
    otro.nombre = "DOXICICLINA 100 MG"
    db.commit()
    assert otro.antibiotico is True


def test_lo_marcado_a_mano_se_respeta(como_admin, db, negocio):
    p = Producto(negocio_id=negocio.id, nombre="NEOMICINA CREMA")
    db.add(p)
    db.commit()
    assert p.antibiotico is True
    assert como_admin.put(f"/antibioticos/productos/{p.id}", json={"antibiotico": False}).json()["manual"] is True
    p.nombre = "NEOMICINA CREMA 20 G"
    db.commit()
    db.refresh(p)
    assert p.antibiotico is False
    assert como_admin.post("/antibioticos/detectar").json()["antibioticos"] == 0
    # Y uno que el sistema no reconoce se puede agregar a mano.
    q = Producto(negocio_id=negocio.id, nombre="XYZ ANTIBIOTICO RARO")
    db.add(q)
    db.commit()
    como_admin.put(f"/antibioticos/productos/{q.id}", json={"antibiotico": True})
    lista = como_admin.get("/antibioticos/productos").json()
    assert [x["nombre"] for x in lista] == ["XYZ ANTIBIOTICO RARO"]
    assert como_admin.get("/antibioticos/productos?q=neomicina").json()[0]["antibiotico"] is False


@pytest.fixture
def amoxicilina_ab(db, negocio):
    """Antibiótico con 10 piezas sin caducidad (existencia inicial)."""
    return producto(db, negocio, "AMOXICILINA 500MG C/12", "85.50", [(10, None, None)])


def entrada(db, negocio, admin, p, piezas, lote="L1", caducidad=date(2027, 5, 31)):
    prov = Proveedor(negocio_id=negocio.id, nombre="NADRO")
    db.add(prov)
    db.flush()
    e = Entrada(negocio_id=negocio.id, proveedor_id=prov.id, folio="A-77", fecha_recepcion=HOY, origen="manual",
                usuario_id=admin.id, subtotal=D(100))
    db.add(e)
    db.flush()
    l = Lote(negocio_id=negocio.id, producto_id=p.id, numero_lote=lote, caducidad=caducidad, cantidad=D(piezas))
    db.add(l)
    db.flush()
    db.add(EntradaRenglon(entrada_id=e.id, producto_id=p.id, cantidad=D(piezas), factor=D(1), piezas=D(piezas),
                          costo_unitario=D(20), costo_pieza=D(20), lote_id=l.id, numero_lote=lote, caducidad=caducidad))
    db.commit()


def libro(cliente, **extra):
    return cliente.get("/antibioticos/libro", params={"desde": HOY.isoformat(), "hasta": HOY.isoformat(), **extra}).json()


def test_libro_de_entradas_y_salidas(como_admin, db, negocio, admin, caja, amoxicilina_ab, shampoo):
    entrada(db, negocio, admin, amoxicilina_ab, 24)
    v1 = vender(como_admin, caja, [r(amoxicilina_ab, 3), r(shampoo, 1)], efectivo="500").json()
    v2 = vender(como_admin, caja, [r(amoxicilina_ab, 2)], efectivo="200").json()
    devolver(como_admin, v2, caja, [pieza(v2, 0, 1)])
    v3 = vender(como_admin, caja, [r(amoxicilina_ab, 1)], efectivo="100").json()
    cancelar(como_admin, v3, caja)

    l = libro(como_admin)
    assert l["antibioticos"] == 1
    res = l["resumen"][0]
    # Hoy: entraron 24 (compra) + 1 (devolución) + 1 (cancelación); salieron 3 + 2 + 1. Quedan 10 + 24 - 6 + 2 = 30.
    assert (res["inicial"], res["entradas"], res["salidas"], res["final"]) == ("10", "26", "6", "30")
    movs = l["movimientos"]
    assert [m["tipo"] for m in movs][:1] == ["entrada"] and movs[0]["documento"] == "Factura A-77"
    assert movs[0]["detalle"] == "NADRO" and movs[0]["numero_lote"] == "L1" and movs[0]["saldo"] == "34"
    assert movs[-1]["saldo"] == "30"
    ventas = [m for m in movs if m["tipo"] == "venta"]
    assert all(m["falta_receta"] for m in ventas) and len(ventas) == 3
    assert {m["tipo"] for m in movs} == {"entrada", "venta", "devolucion", "cancelacion"}

    # Recetas por capturar: la venta cancelada ya no la pide.
    pendientes = como_admin.get("/antibioticos/recetas-pendientes").json()
    assert sorted(p["folio"] for p in pendientes) == sorted([v1["folio"], v2["folio"]])
    assert [a["nombre"] for a in next(p for p in pendientes if p["folio"] == v1["folio"])["antibioticos"]] == \
        ["AMOXICILINA 500MG C/12"]  # el shampoo no cuenta


def test_capturar_receta_y_recordar_medico(como_mostrador, como_admin, db, negocio, caja, amoxicilina_ab):
    v = vender(como_mostrador, caja, [r(amoxicilina_ab, 2)], efectivo="200").json()
    datos = {"medico": "dra. ana lópez", "cedula": " 1234567 ", "domicilio": "Av. Juárez 10, Centro",
             "fecha_receta": HOY.isoformat()}
    res = como_mostrador.put(f"/antibioticos/recetas/{v['id']}", json=datos)
    assert res.status_code == 200, res.text
    assert (res.json()["medico"], res.json()["cedula"]) == ("DRA. ANA LÓPEZ", "1234567")
    assert como_mostrador.get("/antibioticos/recetas-pendientes").json() == []
    venta = next(m for m in libro(como_mostrador)["movimientos"] if m["tipo"] == "venta")
    assert venta["receta"]["cedula"] == "1234567" and venta["falta_receta"] is False
    assert como_mostrador.get("/antibioticos/medicos?q=123").json() == [
        {"cedula": "1234567", "nombre": "DRA. ANA LÓPEZ", "domicilio": "Av. Juárez 10, Centro"}]
    assert como_mostrador.get("/antibioticos/medicos?q=lopez").json()[0]["cedula"] == "1234567"
    # Validaciones.
    sin = como_mostrador.put(f"/antibioticos/recetas/{v['id']}", json={**datos, "cedula": ""})
    assert sin.status_code in (400, 409, 422)
    futura = como_mostrador.put(f"/antibioticos/recetas/{v['id']}",
                                json={**datos, "fecha_receta": (HOY + timedelta(days=3)).isoformat()})
    assert futura.status_code in (400, 409, 422) and "posterior" in futura.json()["detail"]


def test_saldo_de_periodos_anteriores(como_admin, db, negocio, caja, amoxicilina_ab):
    """Un periodo que ya pasó termina con la existencia de ese día, no la de hoy."""
    vender(como_admin, caja, [r(amoxicilina_ab, 4)], efectivo="400")
    ayer = (HOY - timedelta(days=1)).isoformat()
    l = como_admin.get("/antibioticos/libro", params={"desde": ayer, "hasta": ayer}).json()
    assert l["movimientos"] == [] and l["resumen"][0]["final"] == "10"  # la venta de hoy no cuenta


def test_excel(como_admin, db, negocio, caja, amoxicilina_ab):
    from io import BytesIO

    from openpyxl import load_workbook

    vender(como_admin, caja, [r(amoxicilina_ab, 1)], efectivo="100")
    res = como_admin.get("/antibioticos/libro/excel", params={"desde": HOY.isoformat(), "hasta": HOY.isoformat()})
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["Libro de control", "Resumen", "Recetas por capturar"]
    filas = [f for f in wb["Libro de control"].iter_rows(values_only=True)]
    assert any(f[2] == "AMOXICILINA 500MG C/12" and f[11] == "FALTA CAPTURAR" for f in filas)


def test_permisos(como_bodega, como_mostrador):
    assert como_bodega.get("/antibioticos/recetas-pendientes").status_code == 403
    assert como_mostrador.get("/antibioticos/productos").status_code == 403
