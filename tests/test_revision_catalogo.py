"""Reporte de revisión del catálogo: problemas de cada producto y clave SAT
sugerida por el nombre. No cambia nada."""

from decimal import Decimal as D
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Categoria, Producto
from app.services import revision_catalogo as rc


def test_las_claves_de_las_reglas_existen_en_el_sat():
    claves = rc.claves_sat()
    assert len(claves) > 50000
    usadas = {c for _, c in rc.SUSTANCIAS} | {c for _, c, _ in rc.TIPOS}
    assert usadas - claves.keys() == set()


@pytest.mark.parametrize("nombre, clave", [
    ("ADOPREN 400 MG TABLETAS C/10 (IBUPROFENO)", "51142100"),
    ("ALBOZ CAPS 20MG C/14 (OMEPRAZOL)", "51171900"),
    ("TELMISARTAN TABS 80MG C/28 AMSA", "51121700"),
    ("TENESCAN V CREMA 20 GRS (CLOTRIMAZOL)", "51101800"),
    ("LECHE DE MAGNESIA PHILLIPS 360ML", "51171600"),
    ("SH CAPRICE CONTROL CASPA 200 ML", "53131628"),
    ("CRA P/ PEINAR (NATURAL\"U) KERATINA 250 ML", "53131602"),
    ("PASTA DENTAL COLGATE TRIPLE ACCION 50ML", "53131502"),
    ("DES GILLETTE GEL COOL WAVE 113 G", "53131606"),
    ("PAÑAL HUGGIES ULTRACONFORT ET 3 C/40", "53102305"),
    ("BEBIDA COCA 500ML TAPARROSCA", "50202306"),
    ("BEBIDA JUMEX NECTAR CARTON MANGO 475ML", "50202304"),
    ("CERVEZA CORONA LIGHT 355ML", "50202201"),
    ("CURITAS FLEXIBLE XL C/10", "42311505"),
    ("LECHE NAN 1 400G", "50131704"),
    ('CHOCOLATE M&M"S CACAHUATE 37.5G', "50161813"),
    ("PROTEC SOLAR HAWAIIAN ISLAND SPORT SPRAY 170G", "53131609"),
])
def test_clave_sugerida_por_el_nombre(nombre, clave):
    s = rc.sugerir_clave(nombre)
    assert s.clave == clave, s
    assert s.descripcion == rc.claves_sat()[clave]


def test_medicina_sin_sustancia_conocida_no_se_inventa():
    s = rc.sugerir_clave("TREDA C/20 TAB")
    assert s.clave is None and s.tipo == rc.MEDICINA and s.iva_esperado == D(0)
    assert rc.sugerir_clave("COLLAR JMP").clave is None
    # Una crema con concentración o con óvulos es medicina, no cosmético.
    assert rc.sugerir_clave("QUIMARA CRA 5% 5G").tipo == rc.MEDICINA
    assert rc.sugerir_clave("CANESTEN V 3 DIAS DUAL CRA+OV").tipo == rc.MEDICINA
    # Antisépticos y vitaminas: su IVA depende del registro; no se marca.
    assert rc.sugerir_clave("ALCOHOL LASA 120ML").iva_esperado is None


def test_clave_mas_especifica_no_es_dudosa():
    p = Producto(nombre="ACIDO FOLICO 0.5 X 50 TABLETA (SYDENHAN)", clave="1", clave_sat="51131517",
                 costo=D(1), precio_venta=D(5), iva_porcentaje=D(0))
    assert "clave_sat_dudosa" not in rc.revisar(p, "").marcas


def _producto(db, negocio, nombre, **extra):
    _producto.n = getattr(_producto, "n", 0) + 1
    datos = {"clave": f"75010000{_producto.n:05d}", "costo": D(10), "precio_venta": D(20), "clave_sat": None, **extra}
    p = Producto(negocio_id=negocio.id, nombre=nombre, **datos)
    db.add(p)
    return p


def test_problemas_y_excel(como_admin, como_mostrador, db, negocio):
    db.add(Categoria(negocio_id=negocio.id, nombre="Depto 7"))
    bien = _producto(db, negocio, "ADOPREN 400 MG TABLETAS C/10 (IBUPROFENO)", clave_sat="51142100")
    _producto(db, negocio, "ALBOZ CAPS 20MG C/14 (OMEPRAZOL)", clave_sat="51101500", costo=None)  # clave de antibiótico
    _producto(db, negocio, "SH CAPRICE CONTROL CASPA 200 ML", iva_porcentaje=D(0))  # sin clave SAT, IVA dudoso
    _producto(db, negocio, "BEBIDA JUMEX AMI NARANJA 591 ML", clave_sat="84111506", iva_porcentaje=D(16))  # servicio
    _producto(db, negocio, "GASAS", clave=None, precio_venta=None, clave_sat="99999999")
    _producto(db, negocio, "VITAMINA C 1G", precio_venta=D(9), costo=D(10), clave_sat="51191905")
    _producto(db, negocio, "COLLAR JMP", clave="CO1", clave_sat="11121900", iva_porcentaje=D(16))
    _producto(db, negocio, "Collar  JMP", clave="CO2", clave_sat="11121900", iva_porcentaje=D(16))
    db.commit()

    revisiones, vacias = rc.revisar_catalogo(db, negocio.id)
    por_nombre = {r.producto.nombre: r for r in revisiones}
    assert bien.nombre not in por_nombre  # todo bien: no sale
    assert vacias == ["Depto 7"]
    assert set(por_nombre["ALBOZ CAPS 20MG C/14 (OMEPRAZOL)"].marcas) == {"sin_costo", "clave_sat_dudosa"}
    assert set(por_nombre["SH CAPRICE CONTROL CASPA 200 ML"].marcas) == {"sin_clave_sat", "iva_dudoso"}
    assert "no es de un producto" in por_nombre["BEBIDA JUMEX AMI NARANJA 591 ML"].problemas[0]
    assert set(por_nombre["GASAS"].marcas) == {"sin_precio", "clave_sat_dudosa", "sin_codigo", "iva_dudoso"}
    assert set(por_nombre["VITAMINA C 1G"].marcas) == {"precio_bajo"}
    assert set(por_nombre["COLLAR JMP"].marcas) == {"nombre_repetido"}
    assert "otros códigos: CO2" in por_nombre["COLLAR JMP"].problemas[0]

    assert como_mostrador.get("/productos/revision/excel").status_code == 403
    r = como_admin.get("/productos/revision/excel")
    assert r.status_code == 200 and "revision_catalogo_" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Resumen", "Productos", "Categorías vacías"]
    filas = list(wb["Productos"].iter_rows(values_only=True))
    omeprazol = next(f for f in filas if f[1] == "ALBOZ CAPS 20MG C/14 (OMEPRAZOL)")
    assert omeprazol[6:10] == ("51101500", "Antibióticos", "51171900",
                               "Fármacos antiúlcera y otros fármacos gastrointestinales (GI) relacionados")
    assert len(filas) == 1 + len(revisiones)
    db.expire_all()
    assert db.get(Producto, bien.id).clave_sat == "51142100"  # no cambia nada
