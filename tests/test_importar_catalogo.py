"""Cargar productos desde un Excel o CSV en la pantalla de Catálogo."""

import io
from decimal import Decimal as D

from openpyxl import Workbook, load_workbook

from app.models import AjusteInventario, Categoria, Lote, PrecioHistorial, Producto


def excel(filas, encabezado=("Clave", "Nombre", "Precio de venta", "Costo", "IVA %", "IEPS %", "Categoría",
                             "Laboratorio", "Clave SAT", "Mínimo", "Máximo", "Existencia", "Requiere receta"),
          antes=()):
    wb = Workbook()
    ws = wb.active
    for fila in antes:
        ws.append(fila)
    ws.append(list(encabezado))
    for fila in filas:
        ws.append(list(fila))
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def importar(cliente, datos, guardar=False):
    return cliente.post(f"/productos/importar?guardar={'true' if guardar else 'false'}", content=datos,
                        headers={"Content-Type": "application/octet-stream"})


def test_plantilla(como_admin):
    r = como_admin.get("/productos/importar/plantilla")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert [c.value for c in wb["Productos"][1]][:3] == ["Clave", "Nombre", "Precio de venta"]
    assert "Instrucciones" in wb.sheetnames
    # La plantilla tal cual se puede cargar (sus dos ejemplos son productos nuevos).
    assert importar(como_admin, r.content).json()["nuevos"] == 2


def test_revisar_no_guarda_y_guardar_si(como_admin, db, negocio):
    datos = excel([
        ("7501000000017", "Paracetamol 500 mg c/20", 35, 18.5, 0, None, "Patente", "GENOMMA", "51142106", 5, 20, 12, "no"),
        ("7509876543210", "CHOCOLATE 40 G", "$22.00", 14, "16%", 8, "Dulces", None, None, None, None, None, "NO"),
    ])
    r = importar(como_admin, datos).json()
    assert (r["guardado"], r["nuevos"], r["con_error"], r["piezas"]) == (False, 2, 0, "12")
    assert set(r["categorias_nuevas"]) == {"Patente", "Dulces"}
    assert db.query(Producto).count() == 0 and db.query(Categoria).count() == 0

    r = importar(como_admin, datos, guardar=True).json()
    assert r["guardado"] is True and r["nuevos"] == 2
    db.expire_all()
    p = db.query(Producto).filter_by(clave="7501000000017").one()
    assert (p.nombre, p.precio_venta, p.costo, p.iva_porcentaje, p.clave_sat, p.minimo, p.maximo) == \
        ("PARACETAMOL 500 MG C/20", D("35.00"), D("18.5"), D(0), "51142106", D(5), D(20))
    assert db.get(Categoria, p.categoria_id).nombre == "Patente"
    lote = db.query(Lote).filter_by(producto_id=p.id).one()
    assert (lote.cantidad, lote.caducidad) == (D(12), None)
    assert db.query(AjusteInventario).filter_by(producto_id=p.id).one().cantidad == D(12)
    chocolate = db.query(Producto).filter_by(clave="7509876543210").one()
    assert (chocolate.precio_venta, chocolate.iva_porcentaje, chocolate.ieps_porcentaje) == (D(22), D(16), D(8))
    assert db.query(Lote).filter_by(producto_id=chocolate.id).count() == 0
    assert db.query(PrecioHistorial).count() == 2


def test_actualiza_sin_borrar_lo_que_viene_vacio(como_admin, db, negocio):
    p = Producto(negocio_id=negocio.id, clave="070942302388", nombre="PALILLOS", precio_venta=D(25), costo=D(10),
                 iva_porcentaje=D(16), laboratorio="GUM")
    db.add(p)
    db.commit()
    # Clave sin el cero de la izquierda (como la deja Excel), solo con precio y existencia.
    datos = excel([("70942302388", None, 27, None, None, None, None, None, None, None, None, 50, None),
                   ("70942302388", None, 28, None, None, None, None, None, None, None, None, None, None)])
    r = importar(como_admin, datos, guardar=True).json()
    assert (r["actualizados"], r["nuevos"], r["existencia_ignorada"], r["con_error"]) == (1, 0, 1, 1)
    assert "precio 25.00 → 27.00" in r["muestra_actualizados"][0]["cambios"]
    assert "Repetido" in r["errores"][0]["error"]
    db.expire_all()
    assert (p.precio_venta, p.costo, p.laboratorio, p.nombre) == (D(27), D(10), "GUM", "PALILLOS")
    assert db.query(Lote).count() == 0  # la existencia de un producto que ya existía no se toca
    assert importar(como_admin, datos).json()["sin_cambios"] == 1


def test_errores_por_renglon(como_admin, db, negocio):
    datos = excel([
        ("1", "BUENO", 10, None, None, None, None, None, None, None, None, None, None),
        ("2", "IVA RARO", 10, None, 15, None, None, None, None, None, None, None, None),
        ("3", "PRECIO MAL", "diez", None, None, None, None, None, None, None, None, None, None),
        ("4", None, 10, None, None, None, None, None, None, None, None, None, None),
        ("5", "SAT MAL", 10, None, None, None, None, None, "123", None, None, None, None),
        ("6", "MIN MAYOR", 10, None, None, None, None, None, None, 9, 3, None, None),
        ("7", "RECETA MAL", 10, None, None, None, None, None, None, None, None, None, "tal vez"),
    ])
    r = importar(como_admin, datos, guardar=True).json()
    assert r["nuevos"] == 1 and r["con_error"] == 6
    textos = {e["renglon"]: e["error"] for e in r["errores"]}
    assert "0, 8 o 16" in textos[3] and "no es un número" in textos[4] and "Falta el nombre" in textos[5]
    assert "8 números" in textos[6] and "mínimo" in textos[7] and "SÍ o NO" in textos[8]
    db.expire_all()
    assert [p.nombre for p in db.query(Producto)] == ["BUENO"]


def test_otros_nombres_de_columnas_y_titulo_arriba(como_admin, db, negocio):
    datos = excel([("ABC-1", "Jabón neutro", 18.5, 4)], encabezado=("Código", "Descripción", "Precio", "Stock"),
                  antes=[("Reporte de artículos",), ()])
    r = importar(como_admin, datos).json()
    assert r["columnas"] == {"clave": "Código", "nombre": "Descripción", "precio": "Precio", "existencia": "Stock"}
    assert r["nuevos"] == 1 and r["muestra_nuevos"][0]["nombre"] == "JABÓN NEUTRO"


def test_csv(como_admin, db, negocio):
    # Excel en español: punto y coma, coma decimal y acentos en cp1252.
    datos = "Clave;Descripción;Precio;Costo\n111;AGUA 1 L;12,50;1,250.00\n".encode("cp1252")
    r = importar(como_admin, datos, guardar=True).json()
    assert r["nuevos"] == 1 and r["columnas"]["nombre"] == "Descripción"
    db.expire_all()
    agua = db.query(Producto).one()
    assert (agua.precio_venta, agua.costo) == (D("12.50"), D("1250.00"))
    datos = "Clave,Descripción,Precio\n222,HIELO,30\n".encode("utf-8")
    assert importar(como_admin, datos).json()["nuevos"] == 1


def test_archivos_que_no_sirven(como_admin):
    assert importar(como_admin, b"").status_code in (400, 409, 422)
    r = importar(como_admin, excel([("1", 2)], encabezado=("Clave", "Algo")))
    assert r.status_code in (400, 409, 422) and "nombre" in r.json()["detail"]
    assert "xlsx" in importar(como_admin, b"\xd0\xcf\x11\xe0" + b"0" * 100).json()["detail"]


def test_solo_admin(como_mostrador):
    assert como_mostrador.post("/productos/importar", content=b"x").status_code == 403
