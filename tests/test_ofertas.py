"""Ofertas sugeridas: candidatos exactos (sin IA), revisión de lo que propone
el asistente y el API (solo admin)."""

import json
from datetime import timedelta
from decimal import Decimal as D
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.asistente.consultas import hoy
from app.models import Lote, Producto
from app.services import configuracion_ia, ofertas
from app.services.errores import OperacionInvalida


def producto(db, negocio, nombre, costo, precio, existencia, caducidad=None, iva=D(0), receta=False, maximo=None):
    p = Producto(negocio_id=negocio.id, nombre=nombre, clave=nombre[:6], costo=D(costo), precio_venta=D(precio),
                 iva_porcentaje=iva, requiere_receta=receta, maximo=D(maximo) if maximo is not None else None)
    db.add(p)
    db.flush()
    db.add(Lote(negocio_id=negocio.id, producto_id=p.id, caducidad=caducidad, cantidad=D(existencia)))
    db.commit()
    return p


@pytest.fixture
def datos(db, negocio):
    negocio.redondeo_precio_venta = D(1)
    en_dos_meses = hoy() + timedelta(days=60)
    return {
        "jarabe": producto(db, negocio, "JARABE TOS", 50, 80, 8, caducidad=en_dos_meses),
        "antibiotico": producto(db, negocio, "AMOXICILINA", 100, 150, 4, caducidad=en_dos_meses, receta=True),
        "panales": producto(db, negocio, "PAÑALES", 100, 250, 30, iva=D(16), maximo=10),
        "normal": producto(db, negocio, "GASAS", 10, 20, 5, maximo=10),  # no es candidato
        "caducado": producto(db, negocio, "VENCIDO", 10, 20, 5, caducidad=hoy() - timedelta(days=1)),  # a merma, no a oferta
        "sin_costo": producto(db, negocio, "SIN COSTO", 0, 20, 5, maximo=1),
    }


def test_candidatos(db, negocio, datos):
    r = ofertas.candidatos(db, negocio.id)
    nombres = [c["nombre"] for c in r["candidatos"]]
    assert set(nombres) == {"JARABE TOS", "AMOXICILINA", "PAÑALES"}
    assert nombres[-1] == "PAÑALES"  # lo que caduca va primero
    assert r["sin_costo"] == 1 and r["con_historial"] is False
    jarabe = next(c for c in r["candidatos"] if c["nombre"] == "JARABE TOS")
    assert jarabe["razones"] == ["caduca"] and jarabe["piezas_por_caducar"] == 8
    assert jarabe["precio_minimo"] == 50  # caduca: al costo
    panales = next(c for c in r["candidatos"] if c["nombre"] == "PAÑALES")
    assert panales["razones"] == ["sobreinventario"]
    assert panales["precio_minimo"] == 128  # 100 × 1.10 × 1.16 = 127.60 → a pesos hacia arriba


def test_interpretar_revisa_las_propuestas(db, negocio, datos):
    lista = ofertas.candidatos(db, negocio.id)["candidatos"]
    ids = {c["nombre"]: c["producto_id"] for c in lista}
    texto = json.dumps({"ofertas": [
        {"producto_id": ids["JARABE TOS"], "tipo": "descuento", "precio_oferta": 30, "paquete_con_id": 0, "motivo": " Caduca  "},
        {"producto_id": ids["AMOXICILINA"], "tipo": "2x1", "precio_oferta": 0, "paquete_con_id": 0, "motivo": "receta"},
        {"producto_id": ids["PAÑALES"], "tipo": "2x1", "precio_oferta": 0, "paquete_con_id": 0, "motivo": "deja menos del mínimo"},
        {"producto_id": ids["PAÑALES"], "tipo": "3x2", "precio_oferta": 0, "paquete_con_id": 0, "motivo": "Sobran"},
        {"producto_id": 999999, "tipo": "descuento", "precio_oferta": 1, "paquete_con_id": 0, "motivo": "no existe"},
    ]})
    r = {o["producto_id"]: o for o in ofertas.interpretar(texto, lista, D(1))}
    assert set(r) == {ids["JARABE TOS"], ids["PAÑALES"]}
    jarabe = r[ids["JARABE TOS"]]
    assert jarabe["precio_oferta"] == 50 and jarabe["ajustada"] and jarabe["motivo"] == "Caduca"
    assert jarabe["ganancia"] == 0 and jarabe["descuento_porcentaje"] == 38
    panales = r[ids["PAÑALES"]]
    # 2x1 dejaría $125 por pieza (menos del mínimo de $128); 3x2 deja $166.67.
    assert panales["tipo"] == "3x2" and panales["precio_oferta"] == 500 and panales["precio_normal"] == 750
    assert panales["ganancia"] == D("131.03") and panales["descuento_porcentaje"] == 33
    with pytest.raises(OperacionInvalida):
        ofertas.interpretar("no es json", lista, D(1))


def test_paquete_y_receta(db, negocio, datos):
    lista = ofertas.candidatos(db, negocio.id)["candidatos"]
    ids = {c["nombre"]: c["producto_id"] for c in lista}
    texto = json.dumps({"ofertas": [
        {"producto_id": ids["JARABE TOS"], "tipo": "paquete", "precio_oferta": 200, "paquete_con_id": ids["PAÑALES"], "motivo": "Juntos"},
        {"producto_id": ids["PAÑALES"], "tipo": "descuento", "precio_oferta": 140, "paquete_con_id": 0, "motivo": "ya va en paquete"},
        {"producto_id": ids["AMOXICILINA"], "tipo": "descuento", "precio_oferta": 135.4, "paquete_con_id": 0, "motivo": "Caduca"},
    ]})
    r = ofertas.interpretar(texto, lista, D(1))
    assert [o["tipo"] for o in r] == ["paquete", "descuento"]
    assert r[0]["paquete_con"]["nombre"] == "PAÑALES" and r[0]["precio_oferta"] == 200 and r[0]["precio_normal"] == 330
    assert r[1]["precio_oferta"] == 136  # redondeo del negocio hacia arriba


def test_api_solo_admin(como_bodega):
    assert como_bodega.get("/reportes/ofertas/candidatos").status_code == 403
    assert como_bodega.post("/reportes/ofertas").status_code == 403


def test_recomendar_y_excel(como_admin, datos, monkeypatch):
    enviado = {}

    def falsa(cliente, texto):
        enviado["texto"] = texto
        return json.dumps({"ofertas": [{"producto_id": datos["jarabe"].id, "tipo": "precio_especial", "precio_oferta": 59,
                                        "paquete_con_id": 0, "motivo": "Caducan 8 en 2 meses"}]})

    monkeypatch.setattr(configuracion_ia, "cliente", lambda: object())
    monkeypatch.setattr(ofertas, "_llamar", falsa)
    candidatos = como_admin.get("/reportes/ofertas/candidatos").json()
    assert len(candidatos["candidatos"]) == 3 and "factor_impuestos" not in candidatos["candidatos"][0]
    r = como_admin.post("/reportes/ofertas")
    assert r.status_code == 200, r.text
    assert "JARABE TOS" in enviado["texto"] and "GASAS" not in enviado["texto"]
    o = r.json()["ofertas"]
    assert len(o) == 1 and o[0]["precio_oferta"] == "59.00" and o[0]["nombre"] == "JARABE TOS" and not o[0]["ajustada"]

    x = como_admin.post("/reportes/ofertas/excel", json={"ofertas": o})
    assert x.status_code == 200, x.text
    wb = load_workbook(BytesIO(x.content))
    fila = [c.value for c in wb["Ofertas sugeridas"][5]]
    assert fila[1] == "JARABE TOS" and fila[-1] == "Caducan 8 en 2 meses"
    assert wb["Todos los candidatos"].max_row == 4


def test_sin_clave_avisa(como_admin, datos, monkeypatch):
    monkeypatch.setattr(configuracion_ia.settings, "anthropic_api_key", None)
    r = como_admin.post("/reportes/ofertas")
    assert r.status_code in (400, 409, 422) and "clave" in r.json()["detail"]
