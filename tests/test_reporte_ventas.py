"""Reporte de ventas del día / periodo."""

from datetime import timedelta
from decimal import Decimal as D
from io import BytesIO

from openpyxl import load_workbook

from app.models import Lote
from app.services.reporte_ventas import hoy
from tests.test_devoluciones import cancelar, devolver, pieza
from tests.test_ventas import amoxicilina, caja, producto, r, shampoo, vender  # noqa: F401  (fixtures)


def test_reporte_del_dia(como_admin, caja, shampoo, amoxicilina, db):
    # Costo del shampoo: 60 por pieza (sin impuestos); la amoxicilina no tiene costo.
    for lote in db.query(Lote).filter(Lote.producto_id == shampoo.id):
        lote.costo_unitario = D(60)
    db.commit()
    v1 = vender(como_admin, caja, [r(shampoo, 1)], efectivo="200").json()  # 116
    vender(como_admin, caja, [r(shampoo, 2), r(amoxicilina, 1)], tarjeta="317.50")  # 232 + 85.50
    v3 = vender(como_admin, caja, [r(amoxicilina, 1)], efectivo="100").json()
    cancelar(como_admin, v3, caja)
    devolver(como_admin, v1, caja, [pieza(v1, 0, 1)])  # regresa 116

    x = como_admin.get("/reportes/ventas").json()
    assert (x["tickets"], x["vendido"], x["devuelto"], x["venta_neta"]) == (2, "433.50", "116.00", "317.50")
    assert (x["canceladas"], x["total_canceladas"]) == (1, "85.50")
    assert (x["efectivo"], x["tarjeta"], x["piezas"]) == ("116.00", "317.50", "4")
    assert (x["subtotal"], x["iva"]) == ("385.50", "48.00")
    # Solo el shampoo tiene costo: 3 × (100 − 60) = 120; cubre 300 de 385.50 (78%).
    assert (x["ganancia_estimada"], x["piezas_sin_costo"], x["porcentaje_con_costo"]) == ("120.00", 1, "78")
    assert [p["nombre"] for p in x["productos"]][0] == "SHAMPOO 400ML"
    assert x["por_caja"] == [{"grupo": "Mostrador 1", "tickets": 2, "vendido": "433.50"}]
    assert len(x["por_hora"]) == 1 and x["por_hora"][0]["tickets"] == 2

    ayer = (hoy() - timedelta(days=1)).isoformat()
    assert como_admin.get("/reportes/ventas", params={"desde": ayer}).json()["tickets"] == 0
    periodo = como_admin.get("/reportes/ventas", params={"desde": ayer, "hasta": hoy().isoformat()}).json()
    assert periodo["tickets"] == 2 and len(periodo["por_dia"]) == 1


def test_excel(como_admin, caja, shampoo):
    vender(como_admin, caja, [r(shampoo, 1)], efectivo="200")
    resp = como_admin.get("/reportes/ventas/excel")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames[:3] == ["Resumen", "Productos", "Por categoría"]
    assert ["Vendido (con impuestos)", 116] in [[c.value for c in f][:2] for f in wb["Resumen"].iter_rows()]


def test_fechas_y_permisos(como_admin, como_mostrador):
    assert como_admin.get("/reportes/ventas", params={"desde": "2026-09-10", "hasta": "2026-09-01"}).status_code in (400, 409, 422)
    assert como_mostrador.get("/reportes/ventas").status_code == 403
