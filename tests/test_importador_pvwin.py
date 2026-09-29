"""Importador de PVWin con archivos pequeños que imitan el formato real
(título arriba, encabezados en la fila 4, datos desde la 6, renglones
"Letra: X" y total general)."""

from datetime import datetime
from decimal import Decimal as D

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.importador.pvwin import FormatoInvalido, leer_pvwin
from app.models import AjusteInventario, Categoria, Lote, Producto, TipoAjuste
from app.scripts.importar_pvwin import _borrar_importacion, escribir_reporte, importar

ENC_CATALOGO = [None, "Clave", "Alterna", "ProdServ", "Descripción", "SAT", "UMV", "Factor", "Localiza",
                "Gpo ", "Dpto", "USD", "$Pcio Compra", "Mínimo", "Máximo", "Exto", "%IEPS", "%IVA"]
ENC_INVENTARIO = [None, "Clave", "Descripción", "Localización", "Grupo", "Depto", "Existencia", "UMV"]


def _libro(ruta, titulo, encabezados, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append([None, "FARMACIA LA FE"])
    ws.append([None, titulo])
    ws.append([])
    ws.append(encabezados)
    ws.append([])
    for fila in filas:
        ws.append(fila)
    wb.save(ruta)
    return ruta


def cat(clave, desc, gpo=1, dpto=1, costo=10, iva=0, ieps=0, localiza="        ", minimo=0, maximo=0):
    # Relleno con espacios como en PVWin.
    return [None, clave, " " * 20, 51101500, f"{desc:<50}", "H87", "PZA", 1, localiza,
            gpo, dpto, "No", costo, minimo, maximo, "No", ieps, iva]


def inv(clave, desc, existencia, grupo=1, depto=1, localizacion="        "):
    return [None, clave, f"{desc:<50}", localizacion, grupo, depto, existencia, "PZA"]


@pytest.fixture
def archivos(tmp_path):
    catalogo = _libro(tmp_path / "catalogo.xlsx", "Catálogo de artículos", ENC_CATALOGO, [
        cat(7501000000017, "AMOXICILINA 500MG C/12", costo=35.123, minimo=2, maximo=6, localiza="MAVER   "),
        cat(7501000000024, "SHAMPOO 400ML", gpo=2, dpto=6, costo=50, iva=16),
        cat(4006000014395, "CRA EUCERIN 200 ML", gpo=2, dpto=6, costo=200, iva=16),
        cat(4006000014395, "CRA EUCERIN 200 ML", gpo=2, dpto=6, costo=200, iva=16),  # repetido
        cat("BRONCORUB     ", "A GRANEL BRONCORUB LATA", dpto="_DND", costo=0, iva=20),  # IVA mal, costo 0
        cat(1, "ARTICULO DE PRUEBA"),
        cat(7501000000031, "SIN GRUPO", gpo="_GND", dpto="_DND", iva=0),
        cat(7501000000048, "CADUCIDAD EN LOCALIZA", localiza=datetime(2023, 3, 1)),
        cat(7501000000055, "FECHA EN TEXTO", localiza="C/05/16 "),
        ["Letra: A"] + [None] * 17,
        cat(0, "BEBIDA ARANDANO 1LT", gpo=2, dpto=9, iva=16),  # clave 0 compartida
        cat(13117000894, "PAÑAL AFECTIVE GDE", gpo=2, dpto=5, iva=16),  # perdió el 0
    ])
    inventario = _libro(tmp_path / "inventario.xlsx", "Reporte de inventarios", ENC_INVENTARIO, [
        inv(7501000000017, "AMOXICILINA 500MG C/12", 5),
        inv(4006000014395, "CRA EUCERIN 200 ML", 1, 2, 6),
        inv(4006000014395, "CRA EUCERIN 200ML", 2, 2, 6),  # repetido con otro espaciado
        inv("BRONCORUB", "A GRANEL BRONCORUB LATA", -4),  # negativo
        inv(1, "ARTICULO DE PRUEBA", -2),
        inv(1, "TONICO CANNABIS 150 ML", 10, 2, 21),  # comparte clave "1" con el de prueba
        inv(0, "PEINE ESTILISTA", 3, 2, 9),  # comparte clave "0" con la bebida
        inv(7509999999999, "SOLO EN INVENTARIO MEDICAMENTO", 4, 1, 3, "LIOMONT "),
        inv(7508888888888, "SOLO EN INVENTARIO PERFUMERIA", 6, 2, 6),
        inv(7507777777777, "OFERTA ESTERILIZADOR", 1250, 2, 6),  # sospechoso
        ["Letra: A", None, None, None, None, None, 25, None],
        [None] * 8,
        [None, None, None, None, None, None, 1288, None],  # total general
    ])
    return catalogo, inventario


@pytest.fixture
def lectura(archivos):
    return leer_pvwin(*archivos)


def por_nombre(lectura, nombre):
    [p] = [p for p in lectura.productos if p.nombre == nombre]
    return p


# --- Lectura y limpieza ------------------------------------------------------

def test_ignora_titulos_letras_y_totales(lectura):
    assert lectura.renglones_catalogo == 11
    assert lectura.renglones_inventario == 10
    assert not any(p.nombre.startswith("Letra") for p in lectura.productos)


def test_quita_espacios_de_relleno(lectura):
    p = por_nombre(lectura, "A GRANEL BRONCORUB LATA")
    assert p.clave == "BRONCORUB"


def test_une_catalogo_e_inventario_por_clave(lectura):
    p = por_nombre(lectura, "AMOXICILINA 500MG C/12")
    assert (p.clave, p.existencia, p.costo) == ("7501000000017", D(5), D("35.123"))
    assert (p.minimo, p.maximo, p.laboratorio) == (D(2), D(6), "MAVER")


def test_omite_articulo_de_prueba(lectura):
    assert [(o.nombre, o.detalle.split(" (")[0]) for o in lectura.omitidos] == [
        ("ARTICULO DE PRUEBA", "producto de prueba"), ("ARTICULO DE PRUEBA", "producto de prueba"),
    ]


def test_fusiona_duplicados_aunque_cambie_el_espaciado(lectura):
    p = por_nombre(lectura, "CRA EUCERIN 200 ML")
    assert p.existencia == D(3)
    assert len(lectura.fusionados) == 2  # uno del catálogo, uno del inventario


def test_clave_compartida_se_quita_y_se_marca(lectura):
    for nombre, clave in [("TONICO CANNABIS 150 ML", "1"), ("PEINE ESTILISTA", "0"), ("BEBIDA ARANDANO 1LT", "0")]:
        p = por_nombre(lectura, nombre)
        assert p.clave is None
        assert f"la clave '{clave}'" in p.revision[0]
    assert "ARTICULO DE PRUEBA" in por_nombre(lectura, "TONICO CANNABIS 150 ML").revision[0]


def test_corrige_iva_20_y_marca_costo_cero(lectura):
    p = por_nombre(lectura, "A GRANEL BRONCORUB LATA")
    assert p.iva == D(16)
    assert p.costo is None
    assert any("precio de compra en cero" in m for m in p.revision)
    assert [c.nombre for c in lectura.corregidos] == ["A GRANEL BRONCORUB LATA"]


def test_negativo_entra_en_cero_y_va_a_conteo(lectura):
    p = por_nombre(lectura, "A GRANEL BRONCORUB LATA")
    assert p.existencia == 0
    assert [(n.nombre, n.detalle) for n in lectura.negativos] == [("A GRANEL BRONCORUB LATA", "-4")]


def test_iva_deducido_por_grupo_si_no_viene_en_catalogo(lectura):
    assert por_nombre(lectura, "SOLO EN INVENTARIO MEDICAMENTO").iva == D(0)
    assert por_nombre(lectura, "SOLO EN INVENTARIO PERFUMERIA").iva == D(16)
    assert por_nombre(lectura, "SOLO EN INVENTARIO MEDICAMENTO").costo is None
    assert por_nombre(lectura, "SOLO EN INVENTARIO MEDICAMENTO").laboratorio == "LIOMONT"


def test_sin_grupo_se_marca(lectura):
    assert any("sin grupo" in m for m in por_nombre(lectura, "SIN GRUPO").revision)


def test_fechas_en_localiza_no_son_laboratorio(lectura):
    assert por_nombre(lectura, "CADUCIDAD EN LOCALIZA").laboratorio is None
    assert por_nombre(lectura, "FECHA EN TEXTO").laboratorio is None


def test_existencia_sospechosa_se_marca(lectura):
    assert any("muy alta" in m for m in por_nombre(lectura, "OFERTA ESTERILIZADOR").revision)


def test_claves_resultantes_son_unicas(lectura):
    claves = [p.clave for p in lectura.productos if p.clave]
    assert len(claves) == len(set(claves))


def test_archivo_equivocado_da_error_claro(tmp_path, archivos):
    catalogo, inventario = archivos
    with pytest.raises(FormatoInvalido, match="Existencia"):
        leer_pvwin(catalogo, catalogo)  # el catálogo no tiene columna Existencia


# --- Guardado en la base de datos -----------------------------------------

def test_importar_guarda_productos_lotes_y_ajustes(db, negocio, admin, lectura):
    totales = importar(db, negocio.id, admin.id, lectura)
    con_existencia = [p for p in lectura.productos if p.existencia > 0]

    assert totales["productos_creados"] == len(lectura.productos)
    assert totales["lotes_creados"] == len(con_existencia)
    assert totales["piezas"] == sum(p.existencia for p in con_existencia)

    piezas_lotes = db.scalar(select(func.sum(Lote.cantidad)).where(Lote.negocio_id == negocio.id))
    piezas_ajustes = db.scalar(select(func.sum(AjusteInventario.cantidad)).where(AjusteInventario.negocio_id == negocio.id))
    assert piezas_lotes == piezas_ajustes == totales["piezas"]
    assert db.scalar(select(func.count()).where(Lote.caducidad.is_not(None))) == 0
    ajustes = db.scalars(select(AjusteInventario).where(AjusteInventario.negocio_id == negocio.id)).all()
    assert {(a.tipo, a.usuario_id) for a in ajustes} == {(TipoAjuste.IMPORTACION, admin.id)}


def test_importar_crea_categorias_por_depto(db, negocio, admin, lectura):
    importar(db, negocio.id, admin.id, lectura)
    nombres = set(db.scalars(select(Categoria.nombre).where(Categoria.negocio_id == negocio.id)))
    assert nombres == {"Depto 1", "Depto 3", "Depto 5", "Depto 6", "Depto 9", "Depto 21"}
    sin_depto = db.scalar(select(Producto).where(Producto.nombre == "SIN GRUPO"))
    assert sin_depto.categoria_id is None


def test_producto_marcado_guarda_motivo(db, negocio, admin, lectura):
    importar(db, negocio.id, admin.id, lectura)
    p = db.scalar(select(Producto).where(Producto.nombre == "OFERTA ESTERILIZADOR"))
    assert p.requiere_revision is True
    assert "muy alta" in p.motivo_revision


def test_reimportar_conserva_nombre_y_margen_de_categorias(db, negocio, admin, lectura):
    importar(db, negocio.id, admin.id, lectura)
    depto6 = db.scalar(select(Categoria).where(Categoria.negocio_id == negocio.id, Categoria.pvwin_depto == 6))
    depto6.nombre, depto6.margen_porcentaje = "Perfumería", D(30)
    db.commit()

    borrados = _borrar_importacion(db, negocio.id)
    totales = importar(db, negocio.id, admin.id, lectura)

    assert borrados == len(lectura.productos)
    assert totales["categorias_creadas"] == 0
    shampoo = db.scalar(select(Producto).where(Producto.nombre == "SHAMPOO 400ML"))
    assert shampoo.categoria_id == depto6.id
    assert db.get(Categoria, depto6.id).nombre == "Perfumería"


def test_busqueda_encuentra_codigo_que_perdio_el_cero(db, negocio, admin, lectura, como_admin):
    importar(db, negocio.id, admin.id, lectura)
    r = como_admin.get("/productos", params={"q": "013117000894"}).json()
    assert [p["nombre"] for p in r] == ["PAÑAL AFECTIVE GDE"]


def test_reporte_excel(tmp_path, lectura):
    totales = {"productos_creados": 1, "categorias_creadas": 0, "lotes_creados": 0, "piezas": 0}
    ruta = tmp_path / "reporte.xlsx"
    escribir_reporte(ruta, lectura, totales, simulado=True, borrados=0)
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Revisar", "Conteo fisico", "Fusionados", "Omitidos", "Correcciones"]
    assert wb["Conteo fisico"]["B2"].value == "A GRANEL BRONCORUB LATA"
    assert wb["Conteo fisico"]["C2"].value == -4
