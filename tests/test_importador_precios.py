"""Lista de precios de PVWin: precio sin impuestos + IVA/IEPS marcados en el
catálogo, redondeo del negocio, productos nuevos y casos raros."""

from decimal import Decimal as D

import pytest
from openpyxl import load_workbook

from app.importador.pvwin import impuestos_del_catalogo, leer_lista_precios
from app.models import Producto
from app.scripts.importar_precios_pvwin import aplicar_precios, escribir_reporte, precio_final
from tests.test_importador_pvwin import ENC_CATALOGO, _libro, cat

ENC_PRECIOS = [None, "Clave", "Descripción", "UMV", "USD", "$Precio Venta 1", "$Precio Venta 2", "$Precio Venta 3", "$Precio Venta 4"]


def precio(clave, desc, p1, p2=0):
    return [None, clave, f"{desc:<50}", "PZA", "No", p1, p2, 0, 0]


@pytest.fixture
def lista(tmp_path):
    return leer_lista_precios(_libro(tmp_path / "precios.xlsx", "Reporte de lista de precios", ENC_PRECIOS, [
        precio(7501000000017, "AMOXICILINA 500MG C/12", 85.5),
        precio(7501000000024, "SHAMPOO 400ML", 100),  # IVA 16 -> 116
        precio(7501000000048, "GALLETAS", 51.29),  # IVA 16 -> 59.50 -> 60 con redondeo
        precio(7501000000055, "REFRESCO", 10, 12),  # IEPS 8 + IVA 16; trae precio 2
        precio(7508888888888, "CREMA SIN IVA MARCADO", 50),  # tenía 16 deducido por grupo
        precio(7502222222222, "PRODUCTO NUEVO E", 20.4),
        precio(7503333333333, "SIN PRECIO", 0),
        precio(7501000000024, "SHAMPOO 400ML", 100),  # repetido
        precio(1, "ARTICULO DE PRUEBA", 1),
        precio(7501000000017, "OTRO CON LA MISMA CLAVE", 5),  # clave ya usada
        ["Letra: Z", None, None, None, None, None, None, None, None],
    ]))


@pytest.fixture
def impuestos(tmp_path):
    return impuestos_del_catalogo(_libro(tmp_path / "catalogo.xlsx", "Catálogo de artículos", ENC_CATALOGO, [
        cat(7501000000017, "AMOXICILINA 500MG C/12", iva=0),
        cat(7501000000024, "SHAMPOO 400ML", gpo=2, iva=16),
        cat(7501000000048, "GALLETAS", gpo=2, iva=16),
        cat(7501000000055, "REFRESCO", gpo=2, iva=16, ieps=8),
    ]))


@pytest.fixture
def productos(db, negocio):
    negocio.redondeo_precio_venta = D(1)
    base = [
        ("7501000000017", "AMOXICILINA 500MG C/12", 0),
        ("7501000000024", "SHAMPOO 400ML", 16),
        ("7501000000048", "GALLETAS", 16),
        ("7501000000055", "REFRESCO", 16),
        ("7508888888888", "CREMA SIN IVA MARCADO", 16),  # IVA deducido por grupo en la importación anterior
        ("7503333333333", "SIN PRECIO", 0),
        ("7509999999999", "NO VIENE EN LA LISTA", 0),
    ]
    for clave, nombre, iva in base:
        db.add(Producto(negocio_id=negocio.id, clave=clave, nombre=nombre, iva_porcentaje=D(iva)))
    db.commit()


def producto(db, nombre):
    return db.query(Producto).filter_by(nombre=nombre).one()


def test_precio_final():
    assert precio_final(D("100"), D(16), D(0), None, None) == D("116.00")
    assert precio_final(D("51.29"), D(16), D(0), None, None) == D("59.50")
    assert precio_final(D("51.29"), D(16), D(0), D(1), None) == D("60")
    assert precio_final(D("10"), D(16), D(8), None, None) == D("12.53")  # 10 × 1.08 × 1.16
    assert precio_final(D("293.1"), D(16), D(0), D(1), None) == D("340")  # 339.996: no sube a 341
    assert precio_final(D("51.29"), D(16), D(0), D(1), D("59.50")) == D("59")  # no pasa el máximo


def test_pone_precios_con_impuestos_marcados_y_redondeo(db, negocio, productos, lista, impuestos):
    aplicar_precios(db, negocio, lista, impuestos)
    assert producto(db, "AMOXICILINA 500MG C/12").precio_venta == D("86")
    assert producto(db, "SHAMPOO 400ML").precio_venta == D("116")
    assert producto(db, "GALLETAS").precio_venta == D("60")
    assert producto(db, "REFRESCO").precio_venta == D("13")


def test_iva_no_marcado_se_quita(db, negocio, productos, lista, impuestos):
    r = aplicar_precios(db, negocio, lista, impuestos)
    crema = producto(db, "CREMA SIN IVA MARCADO")
    assert (crema.iva_porcentaje, crema.precio_venta) == (D(0), D(50))
    cambios = {c[1]: c[2:] for c in r.impuestos_corregidos}
    assert cambios == {
        "CREMA SIN IVA MARCADO": (D(16), D(0), D(0), D(0)),  # IVA quitado
        "REFRESCO": (D(16), D(16), D(0), D(8)),  # IEPS marcado en el catálogo
    }


def test_crea_los_nuevos_sin_iva_y_marcados(db, negocio, productos, lista, impuestos):
    r = aplicar_precios(db, negocio, lista, impuestos)
    nuevo = producto(db, "PRODUCTO NUEVO E")
    assert (nuevo.clave, nuevo.iva_porcentaje, nuevo.precio_venta, nuevo.requiere_revision) == (
        "7502222222222", D(0), D(21), True)
    assert "lista de precios" in nuevo.motivo_revision
    assert {n[1] for n in r.nuevos} == {"PRODUCTO NUEVO E", "OTRO CON LA MISMA CLAVE"}


def test_nuevo_con_clave_ya_usada_queda_sin_clave(db, negocio, productos, lista, impuestos):
    aplicar_precios(db, negocio, lista, impuestos)
    otro = producto(db, "OTRO CON LA MISMA CLAVE")
    assert otro.clave is None
    assert "AMOXICILINA" in otro.motivo_revision


def test_casos_raros(db, negocio, productos, lista, impuestos):
    r = aplicar_precios(db, negocio, lista, impuestos)
    sin = producto(db, "SIN PRECIO")
    assert (sin.precio_venta, sin.requiere_revision) == (None, True)
    assert [x[1] for x in r.repetidos] == ["SHAMPOO 400ML"]
    assert [x[1] for x in r.omitidos] == ["ARTICULO DE PRUEBA"]
    assert [x[1] for x in r.otros_precios] == ["REFRESCO"]
    assert [x[1] for x in r.no_en_lista] == ["NO VIENE EN LA LISTA"]


def test_se_puede_repetir_sin_duplicar(db, negocio, productos, lista, impuestos):
    aplicar_precios(db, negocio, lista, impuestos)
    db.commit()
    total = db.query(Producto).count()
    r = aplicar_precios(db, negocio, lista, impuestos)
    assert db.query(Producto).count() == total
    # "OTRO CON LA MISMA CLAVE" quedó sin clave y se encuentra por nombre.
    assert r.nuevos == []
    assert r.impuestos_corregidos == []


def test_reporte(tmp_path, db, negocio, productos, lista, impuestos):
    r = aplicar_precios(db, negocio, lista, impuestos)
    ruta = tmp_path / "reporte.xlsx"
    escribir_reporte(ruta, r, simulado=True, paso=D(1))
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Precios", "Impuestos corregidos", "Nuevos", "Sin precio",
                             "Otros precios", "Repetidos", "Omitidos", "No en la lista"]
    assert wb["Resumen"]["B2"].value == "SIMULACIÓN (no se guardó nada)"
