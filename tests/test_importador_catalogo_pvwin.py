"""Exportación completa de PVWin (zip con DBF): lectura de los archivos y
actualización del catálogo sin tocar existencias."""

import struct
import zipfile
from decimal import Decimal as D

import pytest
from openpyxl import load_workbook

from app.importador.pvwin import FormatoInvalido
from app.importador.pvwin_dbf import leer_catalogo, leer_dbf
from app.models import Categoria, Lote, PrecioHistorial, Producto
from app.scripts.importar_catalogo_pvwin import aplicar_catalogo, escribir_reporte
from app.scripts.importar_precios_pvwin import MOTIVO_NUEVO

CAMPOS = {
    "art": [("ARTICULO", 14), ("DESCRIP", 50), ("IEPS", 5), ("IVA", 5), ("GRUPO", 4), ("DEPTO", 4),
            ("MINIMO", 12), ("MAXIMO", 12)],
    "pre": [("ARTICULO", 14), ("DOLARES", 1), ("PCIO_COM", 12), ("PCIO_VTA", 12), ("PCIO_VTA2", 12),
            ("PCIO_VTA3", 12), ("PCIO_VTA4", 12)],
    "sat": [("ARTICULO", 14), ("C_UNIDAD", 3), ("PRODSERV", 8)],
    "loc": [("ARTICULO", 14), ("LOCALIZA", 8)],
}


def dbf(campos: list[tuple[str, int]], renglones: list[dict], borrados: tuple[int, ...] = ()) -> bytes:
    """Un DBF (dBase III) mínimo, como los que exporta PVWin (cp1252)."""
    largo_renglon = 1 + sum(largo for _, largo in campos)
    largo_encabezado = 32 + 32 * len(campos) + 1
    datos = bytearray(struct.pack("<BBBBIHH20x", 3, 126, 10, 2, len(renglones), largo_encabezado, largo_renglon))
    for nombre, largo in campos:
        datos += struct.pack("<11sc4xBB14x", nombre.encode(), b"C", largo, 0)
    datos += b"\r"
    for i, renglon in enumerate(renglones):
        datos += b"*" if i in borrados else b" "
        for nombre, largo in campos:
            datos += str(renglon.get(nombre, "")).encode("cp1252").ljust(largo)[:largo]
    return bytes(datos + b"\x1a")


def articulo(clave, nombre, *, iva="0.00", ieps="0.00", grupo="1", depto="1", costo="10.000",
             precio="20.00", precio2="0.00", dolares="N", sat="51101500", lab=""):
    return clave, nombre, iva, ieps, grupo, depto, costo, precio, precio2, dolares, sat, lab


def exportar(carpeta, articulos, borrados=()):
    """Escribe art001/pre001/sat001/loc001.zip como los deja PVWin."""
    filas = {"art": [], "pre": [], "sat": [], "loc": []}
    for clave, nombre, iva, ieps, grupo, depto, costo, precio, precio2, dolares, sat, lab in articulos:
        filas["art"].append(dict(ARTICULO=clave, DESCRIP=nombre, IVA=iva, IEPS=ieps, GRUPO=grupo, DEPTO=depto,
                                 MINIMO="2.000", MAXIMO="6.000"))
        filas["pre"].append(dict(ARTICULO=clave, DOLARES=dolares, PCIO_COM=costo, PCIO_VTA=precio,
                                 PCIO_VTA2=precio2, PCIO_VTA3="0.00", PCIO_VTA4="0.00"))
        filas["sat"].append(dict(ARTICULO=clave, C_UNIDAD="H87", PRODSERV=sat))
        filas["loc"].append(dict(ARTICULO=clave, LOCALIZA=lab))
    for clave, nombre in (("art", "exp_art.dbf"), ("pre", "exp_pre.dbf"), ("sat", "exp_sat.dbf"), ("loc", "exp_loc.dbf")):
        with zipfile.ZipFile(carpeta / f"{clave}001.zip", "w") as z:
            z.writestr(nombre, dbf(CAMPOS[clave], filas[clave], borrados))
    return carpeta


ARTICULOS = [
    articulo("7501000000017", "AMOXICILINA 500MG C/12", precio="85.50", lab="AMSA"),
    articulo("7501000000024", "SHAMPOO 400ML", iva="16.00", grupo="2", depto="06", precio="100.00"),
    articulo("7501000000055", "REFRESCO", iva="16.00", ieps="8.00", grupo="2", depto="2", precio="10.00"),
    articulo("070942302388", "PALILLOS GUM C/50", iva="16.00", grupo="2", precio="50.00"),  # Excel: sin el 0
    articulo("041388284316", "KANKA SOLUCION 9.7ML", grupo="1", precio="80.00"),
    articulo("41388284316", "HUMULIN 70/30 10ML SUSP INY", precio="550.00"),  # su clave la tenía KANKA
    articulo("7502222222222", "BOLSITA RECOLECTORA P/NIÑA", iva="16.00", grupo="2", precio="30.00"),
    articulo("7503333333333", "PRECIO PUESTO A MANO", precio="100.00"),
    articulo("7504444444444", "VENÍA SOLO EN LA LISTA C/10 TABS", precio="40.00"),
    articulo("7505555555555", "CREMA CON IVA RARO", iva="8.00", grupo="_GND", costo="0.000", precio="10.00"),
    articulo("7506666666666", "SIN PRECIO", precio="0.00", dolares="S"),
    articulo("1", "ARTICULO DE PRUEBA"),
    articulo("7501000000024", "SHAMPOO 400ML", iva="16.00", grupo="2", precio="100.00"),  # repetido
    articulo("7509999999999", "BORRADO EN PVWIN"),
]


@pytest.fixture
def carpeta(tmp_path):
    return exportar(tmp_path, ARTICULOS, borrados=(len(ARTICULOS) - 1,))


@pytest.fixture
def base(db, negocio, crear_usuario):
    """Lo que dejaron las importaciones anteriores desde Excel."""
    negocio.redondeo_precio_venta = D(1)
    otros = Categoria(negocio_id=negocio.id, nombre="Otros", margen_porcentaje=D(50))
    db.add(otros)
    db.flush()
    productos = {
        "amox": Producto(negocio_id=negocio.id, clave="7501000000017", nombre="AMOXICILINA 500MG C/12",
                         precio_venta=D(86), laboratorio="PATENTE LAB"),
        "shampoo": Producto(negocio_id=negocio.id, clave="7501000000024", nombre="SHAMPOO 400ML", precio_venta=D(100)),
        "palillos": Producto(negocio_id=negocio.id, clave="70942302388", nombre="PALILLOS GUM C/50", precio_venta=D(58)),
        "kanka": Producto(negocio_id=negocio.id, clave="41388284316", nombre="KANKA SOLUCION 9.7ML", precio_venta=D(80)),
        "humulin": Producto(negocio_id=negocio.id, clave=None, nombre="HUMULIN 70/30 10ML SUSP INY",
                            precio_venta=D(500), requiere_revision=True,
                            motivo_revision="la clave '41388284316' también la usa KANKA SOLUCION 9.7ML; asignar clave nueva"),
        "bolsita": Producto(negocio_id=negocio.id, clave="BOLSA NIÑA", nombre="BOLSITA RECOLECTORA P/NIÑA"),
        "a_mano": Producto(negocio_id=negocio.id, clave="7503333333333", nombre="PRECIO PUESTO A MANO", precio_venta=D(95)),
        "lista": Producto(negocio_id=negocio.id, clave="7504444444444", nombre="VENÍA SOLO EN LA LISTA C/10 TABS",
                          categoria_id=otros.id, precio_venta=D(40), requiere_revision=True,
                          motivo_revision=f"{MOTIVO_NUEVO}; existencia negativa en PVWin (-1); se importó en cero, contar físicamente"),
        "fuera": Producto(negocio_id=negocio.id, clave="7508888888888", nombre="YA NO ESTÁ EN PVWIN", precio_venta=D(10)),
    }
    db.add_all(productos.values())
    db.flush()
    db.add(Lote(negocio_id=negocio.id, producto_id=productos["shampoo"].id, cantidad=D(7)))
    admin = crear_usuario(negocio)
    db.add(PrecioHistorial(negocio_id=negocio.id, producto_id=productos["a_mano"].id, usuario_id=admin.id,
                           precio_anterior=D(100), precio_nuevo=D(95), origen="Catálogo"))
    db.commit()
    return productos


def por_nombre(db, nombre):
    return db.query(Producto).filter_by(nombre=nombre).one()


def test_lee_los_dbf(carpeta):
    c = leer_catalogo(carpeta)
    assert c.renglones == len(ARTICULOS) - 1  # el borrado no cuenta
    nombres = [a.nombre for a in c.articulos]
    assert "BOLSITA RECOLECTORA P/NIÑA" in nombres  # Ñ en cp1252
    assert "ARTICULO DE PRUEBA" not in nombres and "BORRADO EN PVWIN" not in nombres
    assert nombres.count("SHAMPOO 400ML") == 1 and len(c.repetidos) == 1
    shampoo = next(a for a in c.articulos if a.nombre == "SHAMPOO 400ML")
    assert (shampoo.iva, shampoo.depto, shampoo.grupo) == (D(16), None, 2)  # depto '06' no es el 6
    amox = next(a for a in c.articulos if a.nombre.startswith("AMOX"))
    assert (amox.costo, amox.precio, amox.laboratorio, amox.clave_sat, amox.depto) == (D(10), D("85.50"), "AMSA", "51101500", 1)
    raro = next(a for a in c.articulos if a.nombre == "CREMA CON IVA RARO")
    assert raro.iva is None and raro.costo is None
    assert any("IVA de 8%" in m for m in raro.revision) and any("sin grupo" in m for m in raro.revision)
    assert any("dólares" in m for m in next(a for a in c.articulos if a.nombre == "SIN PRECIO").revision)


def test_archivos_de_dias_distintos(tmp_path):
    carpeta = exportar(tmp_path, ARTICULOS[:3])
    (tmp_path / "otra").mkdir()
    otra = exportar(tmp_path / "otra", ARTICULOS[:2])
    (carpeta / "sat001.zip").write_bytes((otra / "sat001.zip").read_bytes())
    with pytest.raises(FormatoInvalido, match="mismos artículos"):
        leer_catalogo(carpeta)


def test_falta_un_zip(tmp_path):
    exportar(tmp_path, ARTICULOS[:2])
    (tmp_path / "pre001.zip").unlink()
    with pytest.raises(FormatoInvalido, match="pre001.zip"):
        leer_catalogo(tmp_path)


def test_dbf_vacio():
    with pytest.raises(FormatoInvalido):
        leer_dbf(b"")


def test_actualiza_impuestos_precio_costo_y_sat(db, negocio, base, carpeta):
    r = aplicar_catalogo(db, negocio, leer_catalogo(carpeta))
    shampoo = base["shampoo"]
    assert (shampoo.iva_porcentaje, shampoo.precio_venta, shampoo.costo, shampoo.clave_sat) == (D(16), D(116), D(10), "51101500")
    assert (shampoo.minimo, shampoo.maximo) == (D(2), D(6))
    amox = base["amox"]
    assert amox.precio_venta == D(86) and amox.laboratorio == "PATENTE LAB"  # el laboratorio que ya tenía se queda
    assert ("7501000000024", "SHAMPOO 400ML", D(0), D(16), D(0), D(0)) in r.impuestos
    historial = db.query(PrecioHistorial).filter_by(producto_id=shampoo.id).one()
    assert (historial.precio_anterior, historial.precio_nuevo, historial.usuario_id) == (D(100), D(116), None)


def test_no_toca_existencias(db, negocio, base, carpeta):
    aplicar_catalogo(db, negocio, leer_catalogo(carpeta))
    assert db.query(Lote).filter_by(producto_id=base["shampoo"].id).one().cantidad == D(7)
    assert db.query(Lote).count() == 1


def test_corrige_claves_sin_ceros_sin_mezclar_productos(db, negocio, base, carpeta):
    r = aplicar_catalogo(db, negocio, leer_catalogo(carpeta))
    assert base["palillos"].clave == "070942302388"
    # KANKA tenía la clave de HUMULIN sin su cero: cada uno queda con la suya y su nombre.
    assert (base["kanka"].clave, base["kanka"].nombre) == ("041388284316", "KANKA SOLUCION 9.7ML")
    assert (base["humulin"].clave, base["humulin"].precio_venta) == ("41388284316", D(550))
    assert not base["humulin"].requiere_revision and base["humulin"].motivo_revision is None
    assert base["bolsita"].clave == "7502222222222"  # se encontró por nombre
    assert not r.nombres
    assert sorted(n[1] for n in r.nuevos) == ["CREMA CON IVA RARO", "REFRESCO", "SIN PRECIO"]


def test_respeta_precio_puesto_a_mano(db, negocio, base, carpeta):
    r = aplicar_catalogo(db, negocio, leer_catalogo(carpeta))
    assert base["a_mano"].precio_venta == D(95)
    assert r.precio_manual == [("7503333333333", "PRECIO PUESTO A MANO", D(95), D(100))]


def test_los_que_venian_solo_en_la_lista(db, negocio, base, carpeta):
    r = aplicar_catalogo(db, negocio, leer_catalogo(carpeta))
    p = base["lista"]
    # Se quita la marca de "solo venía en la lista", pero no la de existencia negativa.
    assert p.motivo_revision == "existencia negativa en PVWin (-1); se importó en cero, contar físicamente"
    assert db.get(Categoria, p.categoria_id).nombre == "Patente"  # ahora con el Depto 1 de PVWin
    assert any(fila[1] == p.nombre for fila in r.categorias)


def test_nuevos_y_casos_raros(db, negocio, base, carpeta):
    r = aplicar_catalogo(db, negocio, leer_catalogo(carpeta))
    nuevo = por_nombre(db, "SIN PRECIO")
    assert nuevo.precio_venta is None and nuevo.requiere_revision and nuevo.categoria_id is not None
    assert "dólares" in nuevo.motivo_revision and "precio de venta en cero" in nuevo.motivo_revision
    raro = por_nombre(db, "CREMA CON IVA RARO")
    assert raro.iva_porcentaje == D(0)  # el IVA raro no se aplica: se queda el que tenía
    assert "IVA de 8%" in raro.motivo_revision and "precio de compra en cero" in raro.motivo_revision
    assert r.no_en_pvwin == [("7508888888888", "YA NO ESTÁ EN PVWIN")]
    assert base["fuera"].activo and base["fuera"].precio_venta == D(10)


def test_se_puede_repetir(db, negocio, base, carpeta):
    aplicar_catalogo(db, negocio, leer_catalogo(carpeta))
    productos = db.query(Producto).count()
    r = aplicar_catalogo(db, negocio, leer_catalogo(carpeta))
    assert db.query(Producto).count() == productos
    assert not (r.nuevos or r.precios or r.impuestos or r.claves or r.costos or r.categorias)


def test_reporte(tmp_path, db, negocio, base, carpeta):
    catalogo = leer_catalogo(carpeta)
    r = aplicar_catalogo(db, negocio, catalogo)
    ruta = tmp_path / "reporte.xlsx"
    escribir_reporte(ruta, catalogo, r, True, D(1))
    wb = load_workbook(ruta)
    resumen = {fila[0]: fila[1] for fila in wb["Resumen"].iter_rows(values_only=True)}
    assert resumen["Productos nuevos (sin existencia)"] == 3
    assert {"Precios", "Impuestos", "Nuevos", "Claves corregidas", "Revisar", "No vienen en PVWin"} <= set(wb.sheetnames)
