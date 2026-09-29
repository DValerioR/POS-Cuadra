"""Reporte de faltantes por proveedor: historial de proveedores desde las
entradas y comparación del último costo de cada uno."""

from decimal import Decimal as D

import pytest
from openpyxl import load_workbook
from io import BytesIO

from app.models import Lote, Producto, Proveedor


@pytest.fixture
def datos(db, negocio):
    nadro = Proveedor(negocio_id=negocio.id, nombre="Nadro")
    marzam = Proveedor(negocio_id=negocio.id, nombre="Marzam")
    db.add_all([nadro, marzam])
    db.flush()
    p = {
        "amox": Producto(negocio_id=negocio.id, nombre="AMOXICILINA", clave="1", minimo=D(5), maximo=D(20)),
        "shampoo": Producto(negocio_id=negocio.id, nombre="SHAMPOO", clave="2", minimo=D(2), maximo=D(10)),
        "gasas": Producto(negocio_id=negocio.id, nombre="GASAS", clave="3", minimo=D(3), maximo=D(12)),
        "lleno": Producto(negocio_id=negocio.id, nombre="LLENO", clave="4", minimo=D(1), maximo=D(5)),
        "sin_max": Producto(negocio_id=negocio.id, nombre="SIN MAXIMO", clave="5", minimo=D(0), maximo=D(0)),
    }
    db.add_all(p.values())
    db.commit()
    return nadro, marzam, p


def entrada(cliente, proveedor, folio, fecha, renglones):
    r = cliente.post("/entradas", json={"proveedor_id": proveedor.id, "folio": folio, "fecha_recepcion": fecha,
                                        "renglones": renglones})
    assert r.status_code == 201, r.text
    return r


def renglon(producto, cantidad, costo):
    return {"producto_id": producto.id, "cantidad": str(cantidad), "costo_unitario": str(costo)}


def test_reporte_por_proveedor_con_mejor_precio(como_bodega, datos, db, negocio):
    nadro, marzam, p = datos
    # Amoxicilina: Nadro la dio a 40 hace tiempo y luego a 45; Marzam a 42. Último de Nadro: 45.
    entrada(como_bodega, nadro, "N1", "2026-08-01", [renglon(p["amox"], 1, 40), renglon(p["lleno"], 1, 1)])
    entrada(como_bodega, marzam, "M1", "2026-09-01", [renglon(p["amox"], 1, 42), renglon(p["shampoo"], 1, 30)])
    entrada(como_bodega, nadro, "N2", "2026-09-15", [renglon(p["amox"], 2, 45)])
    # Quedan: amox 4 (≤ 5), shampoo 1 (≤ 2), lleno 1 (≤ 1), gasas 0 sin historial.
    r = como_bodega.get("/reportes/faltantes", params={"proveedor_id": nadro.id}).json()
    amox, lleno = r["del_proveedor"]
    assert (amox["nombre"], amox["existencia"], amox["sugerido"], amox["costo"]) == ("AMOXICILINA", "4.00", "16.00", "45.0000")
    assert (amox["mejor_proveedor"], amox["mejor_costo"], amox["diferencia_porcentaje"]) == ("Marzam", "42.0000", "7.1")
    assert [x["proveedor"] for x in amox["proveedores"]] == ["Marzam", "Nadro"]
    assert (lleno["nombre"], lleno["mejor_proveedor"]) == ("LLENO", None)
    assert [x["nombre"] for x in r["sin_proveedor"]] == ["GASAS"]  # SIN MAXIMO no cuenta
    assert r["total_estimado"] == "724.00"  # 45 × 16 + 1 × 4

    r2 = como_bodega.get("/reportes/faltantes", params={"proveedor_id": marzam.id}).json()
    assert [(x["nombre"], x["mejor_proveedor"]) for x in r2["del_proveedor"]] == [("AMOXICILINA", None), ("SHAMPOO", None)]
    assert [x["nombre"] for x in r2["sin_proveedor"]] == ["GASAS"]  # sale en todos


def test_existencia_suficiente_no_sale(como_bodega, datos, db, negocio):
    nadro, _, p = datos
    entrada(como_bodega, nadro, "N1", "2026-09-01", [renglon(p["amox"], 10, 40)])
    r = como_bodega.get("/reportes/faltantes", params={"proveedor_id": nadro.id}).json()
    assert r["del_proveedor"] == []


def test_excel(como_bodega, datos):
    nadro, marzam, p = datos
    entrada(como_bodega, nadro, "N1", "2026-09-01", [renglon(p["amox"], 1, 45)])
    entrada(como_bodega, marzam, "M1", "2026-09-02", [renglon(p["amox"], 1, 40)])
    r = como_bodega.get("/reportes/faltantes/excel", params={"proveedor_id": nadro.id})
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Pedido", "Sin proveedor registrado"]
    fila = [c.value for c in wb["Pedido"][4]]
    assert (fila[1], fila[8]) == ("AMOXICILINA", "Marzam")


def test_proveedores_del_producto(como_admin, datos):
    nadro, marzam, p = datos
    entrada(como_admin, nadro, "N1", "2026-09-01", [renglon(p["amox"], 1, 45)])
    entrada(como_admin, marzam, "M1", "2026-09-02", [renglon(p["amox"], 1, 40)])
    lista = como_admin.get(f"/productos/{p['amox'].id}/proveedores").json()
    assert [(x["proveedor"], x["costo"]) for x in lista] == [("Marzam", "40.0000"), ("Nadro", "45.0000")]


def test_permisos(como_mostrador, datos):
    nadro, _, p = datos
    assert como_mostrador.get("/reportes/faltantes", params={"proveedor_id": nadro.id}).status_code == 403
    assert como_mostrador.get(f"/productos/{p['amox'].id}/proveedores").status_code == 403
