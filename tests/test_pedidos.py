"""Pedidos a proveedores: armar desde faltantes, guardar, enviar, ligar la
entrada que llega y comparar lo pedido contra lo que llegó."""

from decimal import Decimal as D
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Producto, Proveedor


@pytest.fixture
def datos(db, negocio):
    nadro = Proveedor(negocio_id=negocio.id, nombre="Nadro")
    marzam = Proveedor(negocio_id=negocio.id, nombre="Marzam")
    db.add_all([nadro, marzam])
    db.flush()
    p = {
        "amox": Producto(negocio_id=negocio.id, nombre="AMOXICILINA", clave="1", minimo=D(5), maximo=D(20)),
        "gasas": Producto(negocio_id=negocio.id, nombre="GASAS", clave="3", minimo=D(3), maximo=D(12)),
        "jabon": Producto(negocio_id=negocio.id, nombre="JABON", clave="4", minimo=D(1), maximo=D(5), costo=D(24),
                          factor_conversion=D(12)),
    }
    db.add_all(p.values())
    db.commit()
    return nadro, marzam, p


def entrada(cliente, proveedor, folio, renglones, pedido_id=None):
    r = cliente.post("/entradas", json={"proveedor_id": proveedor.id, "folio": folio, "fecha_recepcion": "2026-09-01",
                                        "renglones": renglones, "pedido_id": pedido_id})
    assert r.status_code == 201, r.text
    return r.json()


def renglon(producto, cantidad, costo):
    return {"producto_id": producto.id, "cantidad": str(cantidad), "costo_unitario": str(costo)}


def crear(cliente, proveedor, renglones, notas=None):
    r = cliente.post("/compras/pedidos", json={"proveedor_id": proveedor.id, "notas": notas, "renglones": [
        {"producto_id": p.id, "cantidad": str(c)} for p, c in renglones]})
    assert r.status_code == 201, r.text
    return r.json()


def test_preparar_desde_faltantes(como_bodega, datos):
    nadro, marzam, p = datos
    entrada(como_bodega, nadro, "N1", [renglon(p["amox"], 2, 40)])
    prep = como_bodega.get("/compras/pedidos/preparar", params={"proveedor_id": nadro.id}).json()
    assert [(x["nombre"], x["cantidad"], x["costo"]) for x in prep["renglones"]] == [("AMOXICILINA", "18.00", "40.0000")]
    assert {x["nombre"] for x in prep["sin_proveedor"]} == {"GASAS", "JABON"}
    # Ya pedido a otro proveedor: sale con cantidad 0 y dice dónde.
    otro = crear(como_bodega, marzam, [(p["amox"], 10)])
    como_bodega.post(f"/compras/pedidos/{otro['id']}/enviar")
    prep = como_bodega.get("/compras/pedidos/preparar", params={"proveedor_id": nadro.id}).json()
    x = prep["renglones"][0]
    assert x["cantidad"] == "0" and x["ya_pedido"] == [{"folio": otro["folio"], "proveedor": "Marzam"}]


def test_crear_editar_y_enviar(como_bodega, datos):
    nadro, _, p = datos
    entrada(como_bodega, nadro, "N1", [renglon(p["amox"], 2, 40)])
    pedido = crear(como_bodega, nadro, [(p["amox"], 18), (p["jabon"], 3), (p["gasas"], 0)], notas="  urgente ")
    assert (pedido["folio"], pedido["estado"], pedido["notas"], pedido["productos"]) == (1, "borrador", "urgente", 2)
    # Costo esperado: último con Nadro, o el del catálogo por pieza (24 / 12).
    assert [(r["nombre"], r["costo_esperado"]) for r in pedido["renglones"]] == [("AMOXICILINA", "40.0000"), ("JABON", "2.0000")]
    assert pedido["total_estimado"] == "726.00"
    r = como_bodega.put(f"/compras/pedidos/{pedido['id']}", json={"proveedor_id": nadro.id, "renglones": [
        {"producto_id": p["amox"].id, "cantidad": "10"}]})
    assert r.json()["productos"] == 1
    assert como_bodega.post(f"/compras/pedidos/{pedido['id']}/enviar").json()["estado"] == "enviado"
    r = como_bodega.put(f"/compras/pedidos/{pedido['id']}", json={"proveedor_id": nadro.id, "renglones": [
        {"producto_id": p["amox"].id, "cantidad": "5"}]})
    assert r.status_code in (400, 409, 422)
    assert crear(como_bodega, nadro, [(p["gasas"], 1)])["folio"] == 2


def test_pedido_vacio(como_bodega, datos):
    nadro, _, p = datos
    r = como_bodega.post("/compras/pedidos", json={"proveedor_id": nadro.id, "renglones": [
        {"producto_id": p["amox"].id, "cantidad": "0"}]})
    assert r.status_code in (400, 409, 422)


def test_llega_incompleto_y_con_cambio_de_costo(como_bodega, datos):
    nadro, _, p = datos
    entrada(como_bodega, nadro, "N0", [renglon(p["amox"], 1, 40)])
    pedido = crear(como_bodega, nadro, [(p["amox"], 10), (p["gasas"], 5)])
    como_bodega.post(f"/compras/pedidos/{pedido['id']}/enviar")
    e = entrada(como_bodega, nadro, "N1", [renglon(p["amox"], 6, 44), renglon(p["jabon"], 1, 20)], pedido_id=pedido["id"])
    assert e["pedido"] == {"id": pedido["id"], "folio": pedido["folio"], "estado": "enviado"}
    d = como_bodega.get(f"/compras/pedidos/{pedido['id']}").json()
    c = {f["nombre"]: f for f in d["comparacion"]["renglones"]}
    assert (c["AMOXICILINA"]["estado"], c["AMOXICILINA"]["faltan"], c["AMOXICILINA"]["cambio_costo_porcentaje"]) == \
        ("incompleto", "4.00", "10.0")
    assert (c["GASAS"]["estado"], c["JABON"]["estado"]) == ("no_llego", "no_pedido")
    assert (d["comparacion"]["faltantes"], d["comparacion"]["cambios_de_costo"]) == (2, 1)
    assert d["estado"] == "enviado"
    # Llega lo demás en otra factura, ligada desde el pedido: se cierra solo.
    e2 = entrada(como_bodega, nadro, "N2", [renglon(p["amox"], 4, 44), renglon(p["gasas"], 5, 3)])
    d = como_bodega.get(f"/compras/pedidos/{pedido['id']}").json()
    assert e2["id"] in [x["id"] for x in d["entradas_disponibles"]]
    d = como_bodega.post(f"/compras/pedidos/{pedido['id']}/entradas", json={"entrada_id": e2["id"]}).json()
    assert d["estado"] == "recibido" and d["comparacion"]["faltantes"] == 0
    # Una entrada no se liga dos veces.
    assert como_bodega.post(f"/compras/pedidos/{pedido['id']}/entradas", json={"entrada_id": e2["id"]}).status_code in (400, 409, 422)


def test_entrada_de_otro_proveedor_no_se_liga(como_bodega, datos):
    nadro, marzam, p = datos
    pedido = crear(como_bodega, nadro, [(p["amox"], 10)])
    como_bodega.post(f"/compras/pedidos/{pedido['id']}/enviar")
    e = entrada(como_bodega, marzam, "M1", [renglon(p["amox"], 1, 40)])
    r = como_bodega.post(f"/compras/pedidos/{pedido['id']}/entradas", json={"entrada_id": e["id"]})
    assert r.status_code in (400, 409, 422) and "otro proveedor" in r.json()["detail"]


def test_cerrar_y_cancelar(como_bodega, datos):
    nadro, _, p = datos
    pedido = crear(como_bodega, nadro, [(p["amox"], 10)])
    assert como_bodega.post(f"/compras/pedidos/{pedido['id']}/cerrar").status_code in (400, 409, 422)  # borrador
    assert como_bodega.post(f"/compras/pedidos/{pedido['id']}/cancelar").json()["estado"] == "cancelado"
    otro = crear(como_bodega, nadro, [(p["amox"], 10)])
    como_bodega.post(f"/compras/pedidos/{otro['id']}/enviar")
    entrada(como_bodega, nadro, "N1", [renglon(p["amox"], 3, 40)], pedido_id=otro["id"])
    assert como_bodega.post(f"/compras/pedidos/{otro['id']}/cancelar").status_code in (400, 409, 422)
    assert como_bodega.post(f"/compras/pedidos/{otro['id']}/cerrar").json()["estado"] == "recibido"
    lista = como_bodega.get("/compras/pedidos", params={"estado": "recibido"}).json()
    assert [x["folio"] for x in lista] == [otro["folio"]]


def test_excel(como_bodega, datos):
    nadro, _, p = datos
    pedido = crear(como_bodega, nadro, [(p["jabon"], 3)], notas="Entregar en la mañana")
    r = como_bodega.get(f"/compras/pedidos/{pedido['id']}/excel")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    filas = [[c.value for c in f] for f in ws.iter_rows()]
    assert filas[1][0] == "Entregar en la mañana"
    assert ["4", "JABON", 3, 2, 6] in [f[:5] for f in filas]


def test_permisos(como_mostrador, datos):
    nadro, _, _ = datos
    assert como_mostrador.get("/compras/pedidos").status_code == 403
    assert como_mostrador.get("/compras/pedidos/preparar", params={"proveedor_id": nadro.id}).status_code == 403
