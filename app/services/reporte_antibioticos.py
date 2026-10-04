"""Libro de control de antibióticos (lo que pide Salubridad): entradas y
salidas de cada antibiótico con sus documentos y la existencia.

No se guarda un registro aparte: los movimientos salen de lo que el sistema
ya guarda y nunca se edita (entradas de mercancía, ventas, devoluciones y
ajustes de inventario), así el libro siempre cuadra con el inventario y
también cubre productos que se marquen como antibióticos después.

- Entrada de mercancía: proveedor, factura, lote y caducidad.
- Venta: ticket, lote y los datos de la receta (médico, cédula, domicilio y
  fecha). La farmacia se queda con la receta y sus datos se capturan después
  desde el reporte, para no hacer esperar al cliente.
- Devolución o cancelación de una venta: regresa piezas (entrada).
- Ajustes: existencia inicial importada, conteo físico (+/-) y mermas. Las
  capturas de caducidad no cuentan: solo pasan piezas de un lote a otro.

La existencia de cada renglón se calcula hacia atrás desde la existencia
actual, así que el saldo final del periodo es el del inventario. Las piezas
vendidas que el sistema no tenía registradas (aviso de inventario) no
bajaron ningún lote, así que no salen como salida hasta que se corrijan con
un ajuste; el ticket sí sale en "Recetas por capturar".
"""

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import (
    AjusteInventario, Devolucion, DevolucionRenglon, Entrada, EntradaRenglon, EstadoVenta, Lote, Medico, Negocio,
    Producto, Proveedor, RecetaVenta, RolUsuario, TipoAjuste, TipoDevolucion, Usuario, Venta, VentaRenglon,
    VentaRenglonLote,
)
from app.services.errores import NoEncontrado, OperacionInvalida, SinPermiso

ROLES = {RolUsuario.ADMIN, RolUsuario.MOSTRADOR}


def _zona() -> ZoneInfo:
    return ZoneInfo(settings.zona_horaria)


def validar_rol(usuario: Usuario) -> None:
    if usuario.rol not in ROLES:
        raise SinPermiso("El reporte de antibióticos lo ven administradores y mostrador")


@dataclass
class Movimiento:
    fecha: datetime
    producto_id: int
    tipo: str  # entrada, venta, devolucion, cancelacion, inicial, ajuste, merma
    cantidad: Decimal  # con signo: + entra, - sale
    documento: str  # "Factura A-123", "Ticket 45", ...
    detalle: str | None = None  # proveedor, motivo del ajuste
    numero_lote: str | None = None
    caducidad: date | None = None
    venta_id: int | None = None
    receta: dict | None = None  # médico, cédula, domicilio, fecha
    saldo: Decimal = Decimal(0)


def piezas(d: Decimal | None) -> Decimal | None:
    """10.00 -> 10; 2.50 -> 2.5 (las piezas sin ceros de sobra)."""
    if d is None:
        return None
    return d.quantize(Decimal(1)) if d == d.to_integral_value() else d.normalize()


TIPOS = {"entrada": "Entrada (compra)", "venta": "Salida (venta)", "devolucion": "Devolución de cliente",
         "cancelacion": "Venta cancelada", "inicial": "Existencia inicial", "ajuste": "Ajuste de inventario",
         "merma": "Merma"}


def _receta(r: RecetaVenta | None) -> dict | None:
    if r is None:
        return None
    return {"medico": r.medico_nombre, "cedula": r.cedula, "domicilio": r.domicilio, "fecha": r.fecha_receta}


def _movimientos(db: Session, negocio_id: int, ids: list[int], desde: datetime) -> list[Movimiento]:
    """Todos los movimientos de esos productos desde `desde` hasta hoy."""
    movs: list[Movimiento] = []
    for r, e, prov, lote in db.execute(
        select(EntradaRenglon, Entrada, Proveedor.nombre, Lote)
        .join(Entrada, Entrada.id == EntradaRenglon.entrada_id)
        .join(Proveedor, Proveedor.id == Entrada.proveedor_id)
        .join(Lote, Lote.id == EntradaRenglon.lote_id)
        .where(Entrada.negocio_id == negocio_id, EntradaRenglon.producto_id.in_(ids), Entrada.created_at >= desde)
    ):
        movs.append(Movimiento(e.created_at, r.producto_id, "entrada", r.piezas, f"Factura {e.folio}", prov,
                               lote.numero_lote, lote.caducidad))

    recetas = {}
    for vrl, renglon, venta, lote in db.execute(
        select(VentaRenglonLote, VentaRenglon, Venta, Lote)
        .join(VentaRenglon, VentaRenglon.id == VentaRenglonLote.renglon_id)
        .join(Venta, Venta.id == VentaRenglon.venta_id)
        .join(Lote, Lote.id == VentaRenglonLote.lote_id)
        .where(Venta.negocio_id == negocio_id, VentaRenglon.producto_id.in_(ids), Venta.created_at >= desde)
    ):
        recetas.setdefault(venta.id, None)
        movs.append(Movimiento(venta.created_at, renglon.producto_id, "venta", -vrl.cantidad, f"Ticket {venta.folio}",
                               None, lote.numero_lote, lote.caducidad, venta_id=venta.id))
    if recetas:
        for receta in db.scalars(select(RecetaVenta).where(RecetaVenta.venta_id.in_(list(recetas)))):
            recetas[receta.venta_id] = _receta(receta)
        for m in movs:
            if m.venta_id:
                m.receta = recetas.get(m.venta_id)

    for dr, dev, folio, lote in db.execute(
        select(DevolucionRenglon, Devolucion, Venta.folio, Lote)
        .join(Devolucion, Devolucion.id == DevolucionRenglon.devolucion_id)
        .join(Venta, Venta.id == Devolucion.venta_id)
        .join(Lote, Lote.id == DevolucionRenglon.lote_id)
        .where(Devolucion.negocio_id == negocio_id, Lote.producto_id.in_(ids), Devolucion.created_at >= desde)
    ):
        tipo = "cancelacion" if dev.tipo == TipoDevolucion.CANCELACION else "devolucion"
        movs.append(Movimiento(dev.created_at, lote.producto_id, tipo, dr.cantidad, f"Ticket {folio}", dev.motivo,
                               lote.numero_lote, lote.caducidad, venta_id=dev.venta_id))

    for a, lote in db.execute(
        select(AjusteInventario, Lote).join(Lote, Lote.id == AjusteInventario.lote_id)
        .where(AjusteInventario.negocio_id == negocio_id, AjusteInventario.producto_id.in_(ids),
               AjusteInventario.created_at >= desde, AjusteInventario.tipo != TipoAjuste.CAPTURA_CADUCIDAD)
    ):
        tipo = {TipoAjuste.IMPORTACION: "inicial", TipoAjuste.MERMA: "merma"}.get(a.tipo, "ajuste")
        movs.append(Movimiento(a.created_at, a.producto_id, tipo, a.cantidad, TIPOS[tipo], a.motivo,
                               lote.numero_lote, lote.caducidad))
    movs.sort(key=lambda m: (m.fecha, m.cantidad < 0))
    return movs


@dataclass
class Libro:
    desde: date
    hasta: date
    productos: dict[int, Producto]
    resumen: list[dict] = field(default_factory=list)  # por antibiótico: inicial, entradas, salidas, final
    movimientos: list[Movimiento] = field(default_factory=list)  # del periodo, con saldo


def _rango(desde: date, hasta: date) -> tuple[datetime, datetime]:
    if hasta < desde:
        raise OperacionInvalida("La fecha final es antes que la inicial")
    if (hasta - desde).days > 400:
        raise OperacionInvalida("El periodo es demasiado largo (máximo un año)")
    return datetime.combine(desde, time.min, _zona()), datetime.combine(hasta + timedelta(days=1), time.min, _zona())


def libro(db: Session, usuario: Usuario, desde: date, hasta: date, producto_id: int | None = None) -> Libro:
    validar_rol(usuario)
    inicio, fin = _rango(desde, hasta)
    filtro = [Producto.negocio_id == usuario.negocio_id, Producto.antibiotico.is_(True)]
    if producto_id is not None:
        filtro.append(Producto.id == producto_id)
    productos = {p.id: p for p in db.scalars(select(Producto).where(*filtro).order_by(Producto.nombre))}
    r = Libro(desde, hasta, productos)
    if not productos:
        return r
    ids = list(productos)
    actual = dict(db.execute(select(Lote.producto_id, func.sum(Lote.cantidad))
                             .where(Lote.producto_id.in_(ids)).group_by(Lote.producto_id)).all())
    movs = _movimientos(db, usuario.negocio_id, ids, inicio)
    despues = defaultdict(Decimal)
    for m in movs:
        if m.fecha >= fin:
            despues[m.producto_id] += m.cantidad
    del_periodo = [m for m in movs if m.fecha < fin]
    saldo_final = {pid: Decimal(actual.get(pid) or 0) - despues[pid] for pid in ids}
    cambio = defaultdict(Decimal)
    entradas, salidas = defaultdict(Decimal), defaultdict(Decimal)
    for m in del_periodo:
        cambio[m.producto_id] += m.cantidad
        (entradas if m.cantidad > 0 else salidas)[m.producto_id] += abs(m.cantidad)
    saldo = {pid: saldo_final[pid] - cambio[pid] for pid in ids}
    for m in del_periodo:
        saldo[m.producto_id] += m.cantidad
        m.saldo = saldo[m.producto_id]
    r.movimientos = sorted(del_periodo, key=lambda m: (productos[m.producto_id].nombre, m.fecha))
    for pid, p in productos.items():
        inicial = saldo_final[pid] - cambio[pid]
        if not (inicial or entradas[pid] or salidas[pid] or saldo_final[pid]) and producto_id is None:
            continue  # sin existencia ni movimientos: no hace falta en el libro
        r.resumen.append({"producto_id": pid, "clave": p.clave, "nombre": p.nombre, "inicial": piezas(inicial),
                          "entradas": piezas(entradas[pid]), "salidas": piezas(salidas[pid]),
                          "final": piezas(saldo_final[pid])})
    return r


def a_dict(m: Movimiento, productos: dict[int, Producto]) -> dict:
    p = productos[m.producto_id]
    return {
        "fecha": m.fecha, "producto_id": m.producto_id, "clave": p.clave, "producto": p.nombre, "tipo": m.tipo,
        "tipo_texto": TIPOS[m.tipo], "entrada": piezas(m.cantidad) if m.cantidad > 0 else None,
        "salida": piezas(-m.cantidad) if m.cantidad < 0 else None, "saldo": piezas(m.saldo), "documento": m.documento,
        "detalle": m.detalle, "numero_lote": m.numero_lote, "caducidad": m.caducidad, "venta_id": m.venta_id,
        "receta": m.receta, "falta_receta": m.tipo == "venta" and m.receta is None,
    }


# --- Recetas -------------------------------------------------------------------

def recetas_pendientes(db: Session, usuario: Usuario, limite: int = 300) -> list[dict]:
    """Ventas con antibióticos (no canceladas) a las que les falta capturar la receta."""
    validar_rol(usuario)
    filas = db.execute(
        select(Venta, VentaRenglon.nombre, VentaRenglon.cantidad)
        .join(VentaRenglon, VentaRenglon.venta_id == Venta.id)
        .join(Producto, Producto.id == VentaRenglon.producto_id)
        .outerjoin(RecetaVenta, RecetaVenta.venta_id == Venta.id)
        .where(Venta.negocio_id == usuario.negocio_id, Venta.estado == EstadoVenta.COMPLETADA,
               Producto.antibiotico.is_(True), RecetaVenta.id.is_(None))
        .order_by(Venta.created_at.desc())
    ).all()
    ventas: dict[int, dict] = {}
    for venta, nombre, cantidad in filas:
        v = ventas.setdefault(venta.id, {"venta_id": venta.id, "folio": venta.folio, "fecha": venta.created_at,
                                         "antibioticos": []})
        v["antibioticos"].append({"nombre": nombre, "cantidad": piezas(cantidad)})
    return list(ventas.values())[:limite]


def contar_pendientes(db: Session, negocio_id: int) -> int:
    return db.scalar(
        select(func.count(func.distinct(Venta.id)))
        .join(VentaRenglon, VentaRenglon.venta_id == Venta.id)
        .join(Producto, Producto.id == VentaRenglon.producto_id)
        .outerjoin(RecetaVenta, RecetaVenta.venta_id == Venta.id)
        .where(Venta.negocio_id == negocio_id, Venta.estado == EstadoVenta.COMPLETADA,
               Producto.antibiotico.is_(True), RecetaVenta.id.is_(None))
    )


def guardar_receta(db: Session, usuario: Usuario, venta_id: int, medico: str, cedula: str,
                   domicilio: str | None, fecha_receta: date | None) -> RecetaVenta:
    """Guarda (o corrige) los datos de la receta de una venta y recuerda al
    médico por su cédula. No hace commit."""
    validar_rol(usuario)
    venta = db.get(Venta, venta_id)
    if venta is None or venta.negocio_id != usuario.negocio_id:
        raise NoEncontrado("Venta no encontrada")
    medico = " ".join((medico or "").upper().split())
    cedula = "".join((cedula or "").upper().split())
    domicilio = " ".join((domicilio or "").split()) or None
    if not medico:
        raise OperacionInvalida("Falta el nombre del médico")
    if not cedula or len(cedula) > 20:
        raise OperacionInvalida("Falta la cédula profesional del médico")
    if fecha_receta and fecha_receta > venta.created_at.astimezone(_zona()).date():
        raise OperacionInvalida("La fecha de la receta es posterior a la venta")

    m = db.scalar(select(Medico).where(Medico.negocio_id == usuario.negocio_id, Medico.cedula == cedula))
    if m is None:
        m = Medico(negocio_id=usuario.negocio_id, cedula=cedula, nombre=medico, domicilio=domicilio)
        db.add(m)
    else:
        m.nombre = medico
        m.domicilio = domicilio or m.domicilio
    db.flush()
    receta = db.scalar(select(RecetaVenta).where(RecetaVenta.venta_id == venta.id))
    if receta is None:
        receta = RecetaVenta(negocio_id=usuario.negocio_id, venta_id=venta.id, medico_nombre=medico, cedula=cedula,
                             usuario_id=usuario.id)
        db.add(receta)
    receta.medico_id, receta.medico_nombre, receta.cedula = m.id, medico, cedula
    receta.domicilio, receta.fecha_receta, receta.usuario_id = domicilio or m.domicilio, fecha_receta, usuario.id
    db.flush()
    return receta


def buscar_medicos(db: Session, usuario: Usuario, q: str) -> list[dict]:
    validar_rol(usuario)
    texto = (q or "").strip().upper()
    if len(texto) < 2:
        return []
    filas = db.scalars(
        select(Medico).where(Medico.negocio_id == usuario.negocio_id,
                             (Medico.cedula.startswith(texto)) | (func.unaccent(Medico.nombre).ilike(func.unaccent(f"%{texto}%"))))
        .order_by(Medico.nombre).limit(8)
    )
    return [{"cedula": m.cedula, "nombre": m.nombre, "domicilio": m.domicilio} for m in filas]


# --- Excel ---------------------------------------------------------------------

def excel(db: Session, usuario: Usuario, desde: date, hasta: date) -> bytes:
    r = libro(db, usuario, desde, hasta)
    negocio = db.get(Negocio, usuario.negocio_id)
    wb = Workbook()
    negrita, blanca = Font(bold=True), Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="0A4D40")
    falta = PatternFill("solid", fgColor="FDECEA")

    def encabezado(ws, titulos, anchos):
        ws.append(titulos)
        for celda in ws[ws.max_row]:
            celda.font, celda.fill = blanca, relleno
            celda.alignment = Alignment(wrap_text=True, vertical="center")
        for i, ancho in enumerate(anchos):
            ws.column_dimensions[chr(65 + i)].width = ancho

    ws = wb.active
    ws.title = "Libro de control"
    ws.append([f"Libro de control de antibióticos · {negocio.razon_social or negocio.nombre}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"RFC: {negocio.rfc or ''}", f"Periodo: {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}",
               f"Generado: {datetime.now(_zona()):%d/%m/%Y %H:%M}"])
    ws.append([])
    encabezado(ws, ["Fecha", "Clave", "Antibiótico (nombre y presentación)", "Movimiento", "Documento",
                    "Proveedor / motivo", "Lote", "Caducidad", "Entrada", "Salida", "Existencia",
                    "Médico", "Cédula profesional", "Domicilio del médico", "Fecha de la receta"],
               (17, 15, 46, 20, 16, 26, 12, 12, 9, 9, 10, 30, 16, 36, 14))
    ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=4)
    for m in r.movimientos:
        d = a_dict(m, r.productos)
        rec = d["receta"] or {}
        ws.append([m.fecha.astimezone(_zona()).replace(tzinfo=None), d["clave"], d["producto"], d["tipo_texto"],
                   d["documento"], d["detalle"], d["numero_lote"], d["caducidad"], d["entrada"], d["salida"],
                   d["saldo"], rec.get("medico") or ("FALTA CAPTURAR" if d["falta_receta"] else None),
                   rec.get("cedula"), rec.get("domicilio"), rec.get("fecha")])
        fila = ws[ws.max_row]
        fila[0].number_format = "dd/mm/yyyy hh:mm"
        fila[7].number_format = fila[14].number_format = "dd/mm/yyyy"
        if d["falta_receta"]:
            for celda in fila[11:15]:
                celda.fill = falta

    res = wb.create_sheet("Resumen")
    encabezado(res, ["Clave", "Antibiótico", "Existencia inicial", "Entradas", "Salidas", "Existencia final"],
               (16, 50, 16, 11, 11, 16))
    for x in r.resumen:
        res.append([x["clave"], x["nombre"], x["inicial"], x["entradas"], x["salidas"], x["final"]])
    res.freeze_panes = "A2"

    pend = wb.create_sheet("Recetas por capturar")
    encabezado(pend, ["Ticket", "Fecha", "Antibióticos"], (10, 18, 80))
    for v in recetas_pendientes(db, usuario, limite=5000):
        pend.append([v["folio"], v["fecha"].astimezone(_zona()).replace(tzinfo=None),
                     ", ".join(f"{a['cantidad']} × {a['nombre']}" for a in v["antibioticos"])])
        pend.cell(row=pend.max_row, column=2).number_format = "dd/mm/yyyy hh:mm"
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
