"""Reporte de revisión del catálogo: qué productos tienen algún problema y
cuál es, sin cambiar nada. Incluye la clave SAT que le correspondería a cada
producto según su nombre.

Problemas que se revisan: sin costo, sin precio, precio menor al costo, sin
clave SAT, clave SAT que no existe en el catálogo del SAT o que no parece
corresponder, IVA dudoso para ese tipo de producto, sin código de barras,
marcado para revisar (desde PVWin o al darlo de alta), nombre repetido y
categorías sin productos.

La clave SAT sugerida sale de reglas (no de IA): primero la sustancia activa
que trae el nombre ("... (IBUPROFENO)" -> antiinflamatorios no esteroideos),
después el tipo de producto ("SH ..." -> champús). Cada sugerencia dice qué la
decidió. Si no hay regla que aplique, no se sugiere nada.
El IVA esperado es orientativo (medicinas 0 %, higiene y cosméticos 16 %...):
confirmarlo con el contador.
"""

import csv
import gzip
import re
from dataclasses import dataclass, field
from decimal import Decimal
from functools import cache
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Categoria, Producto
from app.services.errores import NoEncontrado, OperacionInvalida
from app.services.parecidos import normalizar

CATALOGO_SAT = Path(__file__).resolve().parents[1] / "facturacion" / "datos" / "claves_prod_serv.csv.gz"
# Claves que no son de un producto: 01010101 "no existe en el catálogo" y los servicios (segmentos 70 a 95).
SIN_CLAVE_REAL = {"01010101"}


@cache
def claves_sat() -> dict[str, str]:
    with gzip.open(CATALOGO_SAT, "rt", encoding="utf-8", newline="") as f:
        return {fila["clave"]: fila["descripcion"] for fila in csv.DictReader(f)}


# --- Reglas para sugerir la clave SAT ------------------------------------------------
# (palabras, clave, tipo, IVA esperado). Las palabras son una expresión regular
# sobre el nombre normalizado (mayúsculas, sin acentos, "SH" -> "SHAMPOO"...).

MEDICINA = "medicamento"
IVA_MEDICINA = Decimal(0)
IVA_GENERAL = Decimal(16)

SUSTANCIAS: list[tuple[str, str]] = [
    # Antiinflamatorios no esteroideos
    ("IBUPROFENO|NAPROXENO|DICLOFENACO|KETOROLACO|MELOXICAM|CELECOXIB|PIROXICAM|INDOMETACINA|NIMESULIDA|KETOPROFENO|"
     "ACEMETACINA|ETORICOXIB|ACIDO MEFENAMICO|DEXKETOPROFENO", "51142100"),
    # Analgésicos no narcóticos
    ("PARACETAMOL|ACETAMINOFEN|METAMIZOL|ACIDO ACETILSALICILICO|ASPIRINA", "51142000"),
    ("TRAMADOL|BUPRENORFINA|TAPENTADOL", "51142200"),
    ("ERGOTAMINA|SUMATRIPTAN|RIZATRIPTAN", "51142400"),
    # Antiinfecciosos
    ("AMOXICILINA|AMPICILINA|CLAVULAN|DICLOXACILINA|PENICILINA|CEFALEXINA|CEFTRIAXONA|CEFUROXIMA|CEFIXIMA|CEFADROXILO|"
     "CEFACLOR|AZITROMICINA|CLARITROMICINA|ERITROMICINA|CIPROFLOXACINO|LEVOFLOXACINO|MOXIFLOXACINO|NORFLOXACINO|"
     "DOXICICLINA|CLINDAMICINA|GENTAMICINA|NEOMICINA|TRIMETOPRIMA|SULFAMETOXAZOL|NITROFURANTOINA|FOSFOMICINA|"
     "LINCOMICINA|METRONIDAZOL|RIFAXIMINA|MUPIROCINA", "51101500"),
    ("ALBENDAZOL|MEBENDAZOL|NITAZOXANIDA|IVERMECTINA|PIRANTEL|PRAZICUANTEL|SECNIDAZOL|TINIDAZOL|PERMETRINA", "51101700"),
    ("CLOTRIMAZOL|KETOCONAZOL|FLUCONAZOL|ITRACONAZOL|MICONAZOL|TERBINAFINA|NISTATINA|ISOCONAZOL|BIFONAZOL", "51101800"),
    ("ACICLOVIR|VALACICLOVIR|OSELTAMIVIR", "51102300"),
    ("FENAZOPIRIDINA", "51102200"),
    # Corazón, presión y colesterol
    ("LOSARTAN|TELMISARTAN|VALSARTAN|IRBESARTAN|CANDESARTAN|OLMESARTAN|CAPTOPRIL|ENALAPRIL|LISINOPRIL|RAMIPRIL|"
     "AMLODIPINO|NIFEDIPINO|FELODIPINO|VERAPAMILO|DILTIAZEM|METOPROLOL|PROPRANOLOL|ATENOLOL|CARVEDILOL|BISOPROLOL|"
     "NEBIVOLOL|PRAZOSINA|HIDRALAZINA|METILDOPA", "51121700"),
    ("HIDROCLOROTIAZIDA|CLORTALIDONA|FUROSEMIDA|ESPIRONOLACTONA|BUMETANIDA|INDAPAMIDA", "51191500"),
    ("ATORVASTATINA|SIMVASTATINA|ROSUVASTATINA|PRAVASTATINA|BEZAFIBRATO|FENOFIBRATO|GEMFIBROZILO|EZETIMIBA", "51121800"),
    ("DIGOXINA", "51121900"),
    ("ISOSORBIDA|TRIMETAZIDINA", "51121600"),
    ("CLOPIDOGREL|CILOSTAZOL|PENTOXIFILINA", "51131700"),
    ("WARFARINA|ACENOCUMAROL|RIVAROXABAN|APIXABAN|DABIGATRAN|ENOXAPARINA", "51131600"),
    # Diabetes, tiroides, hormonas
    ("METFORMINA|GLIBENCLAMIDA|GLIMEPIRIDA|SITAGLIPTINA|LINAGLIPTINA|VILDAGLIPTINA|SAXAGLIPTINA|DAPAGLIFLOZINA|"
     "EMPAGLIFLOZINA|PIOGLITAZONA|INSULINA|JARDIANZ", "51181500"),
    ("LEVOTIROXINA|TIAMAZOL|METIMAZOL", "51181600"),
    ("PREDNISONA|PREDNISOLONA|DEXAMETASONA|BETAMETASONA|HIDROCORTISONA|DEFLAZACORT|METILPREDNISOLONA", "51181700"),
    ("LEVONORGESTREL|ETINILESTRADIOL|DROSPIRENONA|DESOGESTREL|GESTODENO|NORETISTERONA|MEDROXIPROGESTERONA", "51181800"),
    # Aparato digestivo
    ("OMEPRAZOL|ESOMEPRAZOL|PANTOPRAZOL|LANSOPRAZOL|RABEPRAZOL|RANITIDINA|FAMOTIDINA|SUCRALFATO|METOCLOPRAMIDA|"
     "DOMPERIDONA|CISAPRIDA|ITOPRIDA", "51171900"),
    ("BUTILHIOSCINA|HIOSCINA|TRIMEBUTINA|PINAVERIO|DICICLOVERINA|OTILONIO|ALVERINA|DROTAVERINA", "51172100"),
    ("HIDROXIDO DE ALUMINIO|HIDROXIDO DE MAGNESIO|SIMETICONA|DIMETICONA|BICARBONATO|SAL DE UVAS|MAGALDRATO", "51171500"),
    ("LOPERAMIDA|BISMUTO|RACECADOTRILO|PEPTO", "51171700"),
    ("SENOSIDOS|SENOSIDO|BISACODILO|LACTULOSA|POLIETILENGLICOL|PSYLLIUM|PLANTAGO|MAGNESIA", "51171600"),
    ("ONDANSETRON|DIFENIDOL|MECLIZINA|DIMENHIDRINATO|BETAHISTINA|CINARIZINA|FLUNARIZINA", "51171800"),
    # Vías respiratorias y alergias
    ("LORATADINA|DESLORATADINA|CETIRIZINA|LEVOCETIRIZINA|FEXOFENADINA|CLORFENAMINA|CLORFENIRAMINA|DIFENHIDRAMINA|"
     "EPINASTINA|BILASTINA|KETOTIFENO|RUPATADINA|HIDROXICINA", "51161600"),
    ("AMBROXOL|BROMHEXINA|DEXTROMETORFANO|GUAIFENESINA|FENILEFRINA|PSEUDOEFEDRINA|ACETILCISTEINA|CARBOCISTEINA|"
     "BENZONATATO|OXOLAMINA|CLOBUTINOL|DESENFRIOL|ANTIFLU|TABCIN|XL3|VICK|VAPORUB|DESENFRIOL|NEOMELUBRINA", "51161800"),
    # Sistema nervioso
    ("CARBAMAZEPINA|OXCARBAZEPINA|VALPROATO|VALPROICO|LEVETIRACETAM|TOPIRAMATO|GABAPENTINA|PREGABALINA|FENITOINA|"
     "LAMOTRIGINA", "51141500"),
    ("SERTRALINA|FLUOXETINA|PAROXETINA|ESCITALOPRAM|CITALOPRAM|VENLAFAXINA|DESVENLAFAXINA|DULOXETINA|AMITRIPTILINA|"
     "IMIPRAMINA|MIRTAZAPINA|BUPROPION|TRAZODONA", "51141600"),
    ("ALPRAZOLAM|CLONAZEPAM|DIAZEPAM|LORAZEPAM|BROMAZEPAM", "51141900"),
    ("OLANZAPINA|QUETIAPINA|RISPERIDONA|HALOPERIDOL|ARIPIPRAZOL", "51141700"),
    ("METOCARBAMOL|TIOCOLCHICOSIDO|ORFENADRINA|CICLOBENZAPRINA|TIZANIDINA|CARISOPRODOL", "51151900"),
    ("ALOPURINOL|COLCHICINA|FEBUXOSTAT", "51211500"),
    # Vitaminas, minerales, suero
    ("VITAMINA|VITAMINAS|COMPLEJO B|NEUROBION|MULTIVITAMINICO|ACIDO FOLICO|OMEGA|CENTRUM|SUPRADYN|DOLO NEUROBION", "51191905"),
    ("SULFATO FERROSO|FUMARATO FERROSO|HIERRO", "51131500"),
    ("ELECTROLIT|SUERO ORAL|VIDA SUERO|HIDRATACION ORAL|PEDIALYTE|FLORALYTE", "51191906"),
    # Uso externo
    ("AGUA OXIGENADA|PEROXIDO DE HIDROGENO", "51102709"),
    ("ALCOHOL", "51102710"),
    ("BENZAL|ISODINE|YODOPOVIDONA|CLORHEXIDINA|MERTHIOLATE|ANTISEPTICO", "51102700"),
    ("LAGRIMAS|HIPROMELOSA|HIALURONATO OFT|CARBOXIMETILCELULOSA", "51241120"),
]

# Tipos de producto que no son medicina: (palabras, clave, IVA esperado).
TIPOS: list[tuple[str, str, Decimal | None]] = [
    (r"CONDON|CONDONES|PRESERVATIVO|PRESERVATIVOS", "53131622", IVA_GENERAL),
    (r"PANAL.*ADULTO|ADULTO.*PANAL|TENA", "53102306", IVA_GENERAL),
    (r"PANAL|PANALES|HUGGIES|KLEEN BEBE|ABSORSEC", "53102305", IVA_GENERAL),
    (r"TOALLA F|TOALLAS FEMENINAS|NATURELLA|KOTEX|SABA|PROTECTOR DIARIO|PANTIPROTECTOR|TAMPON|TAMPONES|COPA MENSTRUAL",
     "53131615", Decimal(0)),
    (r"SHAMPOO", "53131628", IVA_GENERAL),
    (r"ACONDICIONADOR|TRATAMIENTO|TINTE|MOUSSE|P PEINAR|PARA PEINAR|MASC|MASCARILLA P CABELLO|GEL P CABELLO|GEL PARA CABELLO|SILICA|KERATINA|SPRAY CAPRICE", "53131602", IVA_GENERAL),
    (r"PASTA DENTAL|PASTA COLGATE|PAST DENTAL|DENTIFRICO", "53131502", IVA_GENERAL),
    (r"CEP DENTAL|CEPILLO DENTAL|CEPILLO DE DIENTES", "53131503", IVA_GENERAL),
    (r"ENJUAGUE|LISTERINE|ENJ BUCAL", "53131501", IVA_GENERAL),
    (r"HILO DENTAL|SEDA DENTAL", "53131504", IVA_GENERAL),
    (r"\bDES\b|DESODORANTE|ANTITRANSPIRANTE", "53131606", IVA_GENERAL),
    (r"JABON", "53131608", IVA_GENERAL),
    (r"BLOQ|BLOQUEADOR|PROTECTOR SOLAR|PROTEC SOLAR|FPS", "53131609", IVA_GENERAL),
    (r"LABIAL|CHAP STICK|LABELLO|BALSAMO LABIAL", "53131630", IVA_GENERAL),
    (r"PERFUME|COLONIA|FRAGANCIA|LOCION", "53131620", IVA_GENERAL),
    (r"ESMALTE", "53131638", IVA_GENERAL),
    (r"MAQUILLAJE|RIMEL|DELINEADOR|POLVO COMPACTO|BASE DE MAQUILLAJE|SOMBRA", "53131619", IVA_GENERAL),
    (r"PEINE|CEPILLO P CABELLO|CEPILLO PARA CABELLO", "53131604", IVA_GENERAL),
    (r"RASTRILLO|PRESTOBARBA|NAVAJA|ESPUMA DE AFEITAR|GEL DE AFEITAR", "53131600", IVA_GENERAL),
    (r"TALCO|DEPIL|DEPILATORIA|NAIR|VEET", "53131600", IVA_GENERAL),
    (r"CREMA|HINDS|NIVEA|CERAVE|EUCERIN|AVENE|VASELINA|ACEITE P BEBE|ACEITE PARA BEBE", "53131613", IVA_GENERAL),
    (r"TOALLITAS|HUMEDAS|DODYS", "53131600", IVA_GENERAL),
    (r"KLEENEX|PANUELOS", "14111701", IVA_GENERAL),
    (r"HIGIENICO|PAPEL HIGIENICO", "14111704", IVA_GENERAL),
    (r"CUBREBOCAS|TAPABOCAS", "42131606", IVA_GENERAL),
    (r"JERINGA|JERINGAS", "42142600", IVA_GENERAL),
    (r"GASA|GASAS", "42311512", IVA_GENERAL),
    (r"ALGODON|COTONETES|HISOPOS", "42141500", IVA_GENERAL),
    (r"TERMOMETRO", "42182200", IVA_GENERAL),
    (r"CURITAS|APOSITO|VENDA|VENDAS|VENDAJE", "42311505", IVA_GENERAL),
    (r"CINTA MICROPORE|MICROPORE|LEUKOPLAST|LEUKOFIX|CINTA ADH|NEXCARE", "42311708", IVA_GENERAL),
    (r"PILA|PILAS|DURACELL|ENERGIZER", "26111702", IVA_GENERAL),
    (r"INSECTICIDA|RAID|BAYGON|REPELENTE|OFF\b|VAPE", "10191509", IVA_GENERAL),
    (r"CERVEZA", "50202201", IVA_GENERAL),
    (r"BEBIDA AGUA|AGUA CIEL|AGUA BONAFONT|AGUA EPURA|AGUA NATURAL", "50202301", IVA_GENERAL),
    (r"GATORADE|POWERADE|ENERGETICA|RED BULL|MONSTER|VOLT", "50202309", IVA_GENERAL),
    (r"JUGO|NECTAR|JUMEX|DEL VALLE|BOING", "50202304", IVA_GENERAL),  # la ley del IVA grava jugos y néctares
    (r"BEBIDA|REFRESCO|COCA|PEPSI|SQUIRT|FANTA|SPRITE|SIDRAL", "50202306", IVA_GENERAL),
    (r"LECHE", "50131704", Decimal(0)),
    (r"CHOCOLATE|CARLOS V|KINDER|MILKY WAY|SNICKERS|KIT KAT|HERSHEY", "50161813", IVA_GENERAL),
    (r"SABRITAS|RUFFLES|CHEETOS|DORITOS|TOSTITOS|PAPAS|CACAHUATE|BOTANA", "50192100", IVA_GENERAL),
    (r"GALLETAS|GALLETA|OREO|RITZ|EMPERADOR|MARIAS", "50181900", None),
    (r"DULCE|DULCES|CARAMELO|PALETA|CHICLE|TRIDENT|TIC TAC|HALLS|MENTAS|GOMITAS|BUBBALOO", "50161800", IVA_GENERAL),
]

# Formas de medicamento: si el nombre dice alguna y no hay otra regla, es medicina (IVA 0 %).
FORMAS_MEDICINA = re.compile(
    r"\b(TAB|CAP|GRAGEAS|GRAG|JARABE|JBE|SUSP|SUSPENSION|INY|INYECTABLE|AMP|AMPOLLETAS|GTS|GOTAS|OVULOS|"
    r"SUPOSITORIOS|SUP|UNG|UNGUENTO|OFT|OFTALMICA|COMPR|MCG|MG|SOBRES|EFERV|OTICA|NEBUL|AEROSOL|INHALADOR|"
    r"PARCHE|PARCHES|AGRANEL|OFTENO|OV|OVULO)\b"
)
# Una concentración ("CREMA 5%", "SOL 0.3%") casi siempre es de un medicamento.
CONCENTRACION = re.compile(r"\d\s*%")

SIN_IVA_SEGURO = {"51102700", "51102709", "51102710", "51191905", "51191906", "51131500"}

_SUSTANCIAS = [(re.compile(rf"\b({p})\b"), c) for p, c in SUSTANCIAS]
_TIPOS = [(re.compile(rf"\b({p})\b" if not p.startswith("\\b") else p), c, iva) for p, c, iva in TIPOS]


@dataclass
class Sugerencia:
    clave: str | None
    descripcion: str | None
    motivo: str  # qué la decidió
    tipo: str | None  # "medicamento" u otro
    iva_esperado: Decimal | None


def sugerir_clave(nombre: str, requiere_receta: bool = False) -> Sugerencia:
    """La clave SAT que le correspondería al producto por su nombre (o None)."""
    texto = normalizar(nombre)
    for patron, clave in _SUSTANCIAS:
        m = patron.search(texto)
        if m:
            # Antisépticos, vitaminas y suero oral: su IVA depende de cómo está registrado el producto.
            iva = None if clave in SIN_IVA_SEGURO else IVA_MEDICINA
            return Sugerencia(clave, claves_sat().get(clave), f"el nombre dice {m.group(1).lower()}", MEDICINA, iva)
    es_medicina = requiere_receta or bool(FORMAS_MEDICINA.search(texto)) or bool(CONCENTRACION.search(nombre))
    if not es_medicina:
        for patron, clave, iva in _TIPOS:
            m = patron.search(texto)
            if m:
                return Sugerencia(clave, claves_sat().get(clave), f"el nombre dice {m.group(0).strip().lower()}", "otro", iva)
    if es_medicina:
        return Sugerencia(None, None, "es medicamento, pero no se reconoció la sustancia activa", MEDICINA, IVA_MEDICINA)
    return Sugerencia(None, None, "no se reconoció el tipo de producto", None, None)


# --- Problemas de cada producto --------------------------------------------------------


@dataclass
class Revision:
    producto: Producto
    categoria: str
    sugerencia: Sugerencia
    problemas: list[str] = field(default_factory=list)
    marcas: dict[str, bool] = field(default_factory=dict)


COLUMNAS = [
    ("sin_costo", "Sin costo"),
    ("sin_precio", "Sin precio"),
    ("precio_bajo", "Precio menor al costo"),
    ("sin_clave_sat", "Sin clave SAT"),
    ("clave_sat_dudosa", "Clave SAT dudosa"),
    ("iva_dudoso", "IVA dudoso"),
    ("sin_codigo", "Sin código de barras"),
    ("para_revisar", "Marcado para revisar"),
    ("nombre_repetido", "Nombre repetido"),
]


def _mismo_grupo(a: str, b: str) -> bool:
    """Misma clase del SAT (los primeros 6 dígitos): "51101500" y "51101508" sí; "51101500" y "51171900" no."""
    return a[:6] == b[:6]


def _describe_lo_mismo(clave: str, nombre: str) -> bool:
    """La clave actual es más específica y habla de lo mismo que el nombre
    ("51131517 Ácido fólico" para "ACIDO FOLICO 0.5 MG"): no es dudosa."""
    palabras = {w for w in normalizar(claves_sat()[clave]).split() if len(w) >= 5}
    return bool(palabras & set(normalizar(nombre).split()))


def revisar(p: Producto, categoria: str) -> Revision:
    s = sugerir_clave(p.nombre, p.requiere_receta)
    r = Revision(p, categoria, s)

    def marcar(clave: str, texto: str) -> None:
        r.marcas[clave] = True
        r.problemas.append(texto)

    iva = Decimal(p.iva_porcentaje or 0)
    if p.costo is None or p.costo <= 0:
        marcar("sin_costo", "sin costo")
    if p.precio_venta is None or p.precio_venta <= 0:
        marcar("sin_precio", "sin precio de venta")
    elif p.costo and p.precio_venta < (p.costo * (1 + iva / 100)).quantize(Decimal("0.01")):
        marcar("precio_bajo", f"se vende en ${p.precio_venta} y cuesta ${p.costo * (1 + iva / 100):.2f} con IVA")
    actual = (p.clave_sat or "").strip()
    if not actual:
        marcar("sin_clave_sat", "sin clave SAT")
    elif actual not in claves_sat():
        marcar("clave_sat_dudosa", f"la clave SAT {actual} no existe en el catálogo del SAT")
    elif actual in SIN_CLAVE_REAL or actual[:2] >= "70":
        marcar("clave_sat_dudosa", f"la clave SAT {actual} ({claves_sat()[actual]}) no es de un producto")
    elif s.clave and not _mismo_grupo(actual, s.clave) and not _describe_lo_mismo(actual, p.nombre):
        marcar("clave_sat_dudosa", f"la clave SAT {actual} ({claves_sat()[actual]}) no parece corresponder")
    if s.iva_esperado is not None and iva != s.iva_esperado:
        marcar("iva_dudoso", f"tiene IVA {iva:g} % y para {'un medicamento' if s.tipo == MEDICINA else 'este tipo de producto'} "
                             f"normalmente es {s.iva_esperado:g} %")
    if not (p.clave or "").strip():
        marcar("sin_codigo", "sin código de barras ni clave")
    if p.requiere_revision:
        marcar("para_revisar", f"marcado para revisar: {p.motivo_revision or 'sin motivo'}")
    return r


def revisar_catalogo(db: Session, negocio_id: int) -> tuple[list[Revision], list[str]]:
    """(productos activos con algún problema, categorías sin productos activos)."""
    categorias = dict(db.execute(select(Categoria.id, Categoria.nombre).where(Categoria.negocio_id == negocio_id)).all())
    productos = db.scalars(select(Producto).where(Producto.negocio_id == negocio_id, Producto.activo.is_(True))
                           .order_by(Producto.nombre)).all()
    revisiones = [revisar(p, categorias.get(p.categoria_id, "")) for p in productos]
    # Nombres iguales (sin contar espacios ni signos): el mismo producto dado de alta dos veces, o
    # variantes que conviene distinguir en el nombre para no confundirlas al vender.
    iguales: dict[str, list[Revision]] = {}
    for r in revisiones:
        iguales.setdefault(re.sub(r"[^0-9A-Z]", "", normalizar(r.producto.nombre)), []).append(r)
    for grupo in iguales.values():
        if len(grupo) > 1:
            for r in grupo:
                otros = ", ".join(x.producto.clave or "sin código" for x in grupo if x is not r)
                r.marcas["nombre_repetido"] = True
                r.problemas.append(f"hay {len(grupo)} productos con este nombre (otros códigos: {otros})")
    con_productos = {cid for (cid,) in db.execute(select(Producto.categoria_id).where(
        Producto.negocio_id == negocio_id, Producto.activo.is_(True)).group_by(Producto.categoria_id))}
    vacias = sorted((n for cid, n in categorias.items() if cid not in con_productos), key=_orden_natural)
    return [r for r in revisiones if r.problemas], vacias


def _orden_natural(texto: str):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", texto)]


# --- Excel --------------------------------------------------------------------------------


def excel(db: Session, negocio_id: int) -> bytes:
    revisiones, vacias = revisar_catalogo(db, negocio_id)
    total = db.scalar(select(func.count()).select_from(Producto).where(
        Producto.negocio_id == negocio_id, Producto.activo.is_(True)))
    negrita = Font(bold=True)
    encabezado = PatternFill("solid", fgColor="DCEFE9")
    wb = Workbook()

    res = wb.active
    res.title = "Resumen"
    res.append(["Revisión del catálogo"])
    res["A1"].font = Font(bold=True, size=14)
    res.append([f"{total} productos activos; {len(revisiones)} con algo que revisar. No se cambió nada."])
    res.append([])
    res.append(["Problema", "Productos", "Qué hacer"])
    for c in res[4]:
        c.font, c.fill = negrita, encabezado
    ayuda = {
        "sin_costo": "Se llena solo al registrar sus facturas en Entradas, o a mano en Productos y precios.",
        "sin_precio": "Ponerle precio: no se puede vender sin precio.",
        "precio_bajo": "Revisar el precio o el costo: se vende perdiendo.",
        "sin_clave_sat": "Necesaria para facturar. La hoja Productos trae la sugerida cuando se pudo reconocer.",
        "clave_sat_dudosa": "La clave que tiene no parece de ese producto (muchas vienen así de PVWin). Ver la sugerida.",
        "iva_dudoso": "Orientativo: medicinas 0 %, higiene y cosméticos 16 %... Confirmarlo con el contador.",
        "sin_codigo": "Ponerle su código de barras para poder escanearlo.",
        "para_revisar": "Motivos que dejó la importación de PVWin (existencias negativas, sin costo...).",
        "nombre_repetido": "El mismo producto dado de alta dos veces, o variantes que conviene distinguir en el nombre.",
    }
    for clave, titulo in COLUMNAS:
        res.append([titulo, sum(1 for r in revisiones if r.marcas.get(clave)), ayuda[clave]])
    res.append(["Categorías sin productos", len(vacias), "Hoja Categorías vacías: se pueden borrar o reutilizar."])
    res.append([])
    res.append(["La clave SAT sugerida sale del nombre del producto (la sustancia activa o el tipo de producto); "
                "la columna \"Por qué\" dice qué la decidió. Revisar antes de usarla."])
    res.column_dimensions["A"].width = 26
    res.column_dimensions["B"].width = 12
    res.column_dimensions["C"].width = 100

    hoja = wb.create_sheet("Productos")
    titulos = ["Código", "Producto", "Categoría", "Costo", "Precio", "IVA %", "Clave SAT actual", "Descripción SAT actual",
               "Clave SAT sugerida", "Descripción SAT sugerida", "Por qué", "Problemas"] + [t for _, t in COLUMNAS]
    hoja.append(titulos)
    for c in hoja[1]:
        c.font, c.fill = negrita, encabezado
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in revisiones:
        p, s = r.producto, r.sugerencia
        actual = (p.clave_sat or "").strip()
        hoja.append([
            p.clave, p.nombre, r.categoria,
            float(p.costo) if p.costo is not None else None,
            float(p.precio_venta) if p.precio_venta is not None else None,
            float(p.iva_porcentaje or 0),
            actual or None, claves_sat().get(actual) if actual else None,
            s.clave, s.descripcion, s.motivo, "; ".join(r.problemas),
        ] + ["Sí" if r.marcas.get(clave) else "" for clave, _ in COLUMNAS])
    anchos = [16, 45, 18, 10, 10, 7, 12, 30, 12, 34, 34, 60] + [11] * len(COLUMNAS)
    for i, ancho in enumerate(anchos):
        hoja.column_dimensions[hoja.cell(1, i + 1).column_letter].width = ancho
    hoja.freeze_panes = "C2"
    hoja.auto_filter.ref = hoja.dimensions

    cat = wb.create_sheet("Categorías vacías")
    cat.append(["Categoría sin productos activos"])
    cat["A1"].font = negrita
    for n in vacias:
        cat.append([n])
    cat.column_dimensions["A"].width = 40

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# --- Aplicar claves SAT sugeridas ------------------------------------------------------


def claves_sugeridas(db: Session, negocio_id: int) -> list[dict]:
    """Productos con clave SAT sugerida que no tienen clave o la tienen dudosa."""
    revisiones, _ = revisar_catalogo(db, negocio_id)
    return [{
        "producto_id": r.producto.id, "nombre": r.producto.nombre, "clave": r.producto.clave,
        "clave_actual": (r.producto.clave_sat or "").strip() or None,
        "descripcion_actual": claves_sat().get((r.producto.clave_sat or "").strip()),
        "sugerida": r.sugerencia.clave, "descripcion": r.sugerencia.descripcion, "motivo": r.sugerencia.motivo,
        "caso": "sin_clave" if r.marcas.get("sin_clave_sat") else "dudosa",
    } for r in revisiones
        if r.sugerencia.clave and (r.marcas.get("sin_clave_sat") or r.marcas.get("clave_sat_dudosa"))]


def aplicar_claves(db: Session, negocio_id: int, cambios: list[tuple[int, str]]) -> int:
    """Pone la clave SAT elegida a cada producto. Solo claves que existen en el
    catálogo del SAT. Regresa cuántos cambiaron. No hace commit."""
    claves = claves_sat()
    malas = sorted({c for _, c in cambios if c not in claves})
    if malas:
        raise OperacionInvalida(f"Estas claves no existen en el catálogo del SAT: {', '.join(malas)}")
    productos = {p.id: p for p in db.scalars(select(Producto).where(
        Producto.negocio_id == negocio_id, Producto.id.in_([pid for pid, _ in cambios])))}
    if len(productos) != len({pid for pid, _ in cambios}):
        raise NoEncontrado("Algún producto no existe")
    cambiados = 0
    for pid, clave in cambios:
        if productos[pid].clave_sat != clave:
            productos[pid].clave_sat = clave
            cambiados += 1
    db.flush()
    return cambiados
