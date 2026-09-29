"""Lectura y limpieza de los reportes de Excel que exporta PVWin.

Solo lee y decide; no toca la base de datos (eso lo hace
app/scripts/importar_pvwin.py). Las reglas están en CONTEXTO.md, sección
"Importación desde PVWin".

Se usan dos reportes, tal como salen de PVWin, que se unen por Clave:
- "Catálogo de artículos": costo, impuestos, mínimo/máximo. Sin existencias.
- "Reporte de inventarios - Detallado - Sin mostrar ceros": existencias.
"""

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

FILA_ENCABEZADO = 4
FILA_DATOS = 6

# Columnas (base 0; la columna 0 va vacía salvo en renglones "Letra: X").
CAT = dict(clave=1, prodserv=3, descripcion=4, localiza=8, gpo=9, dpto=10, costo=12,
           minimo=13, maximo=14, ieps=16, iva=17)
INV = dict(clave=1, descripcion=2, localizacion=3, grupo=4, depto=5, existencia=6)

# Productos que no se importan.
OMITIR = {"ARTICULO DE PRUEBA"}
# Existencias por encima de esto se marcan como sospechosas (ej. 1,250 esterilizadores).
EXISTENCIA_SOSPECHOSA = Decimal(500)
# IVA por grupo de PVWin, para productos que no vienen en el catálogo.
IVA_POR_GRUPO = {1: Decimal(0), 2: Decimal(16)}
# IVA capturado mal en PVWin -> valor correcto.
CORRECCIONES_IVA = {Decimal(20): Decimal(16)}


class FormatoInvalido(Exception):
    """El archivo no tiene la estructura esperada (¿se exportó otro reporte?)."""


@dataclass
class ProductoPVWin:
    clave: str | None
    nombre: str
    clave_sat: str | None = None
    laboratorio: str | None = None
    grupo: int | None = None
    depto: int | None = None
    costo: Decimal | None = None
    iva: Decimal | None = None
    ieps: Decimal = Decimal(0)
    minimo: Decimal | None = None
    maximo: Decimal | None = None
    existencia: Decimal = Decimal(0)
    en_catalogo: bool = False
    en_inventario: bool = False
    revision: list[str] = field(default_factory=list)


@dataclass
class Incidencia:
    clave: str | None
    nombre: str
    detalle: str


@dataclass
class ResultadoLectura:
    productos: list[ProductoPVWin]
    omitidos: list[Incidencia]
    fusionados: list[Incidencia]
    corregidos: list[Incidencia]
    # Existencias negativas puestas en cero; detalle = existencia original.
    negativos: list[Incidencia]
    renglones_catalogo: int
    renglones_inventario: int


# --- Conversión de celdas -------------------------------------------------

def _texto(valor) -> str:
    """Quita espacios de relleno y colapsa espacios dobles."""
    return " ".join(str(valor).split()) if valor is not None else ""


def _comparable(nombre: str) -> str:
    """Nombre para comparar: sin espacios ni signos, así "0.5 MG" y "0.5MG"
    cuentan como el mismo producto."""
    return re.sub(r"[^0-9A-ZÁÉÍÓÚÜÑ]", "", nombre.upper())


def _clave(valor) -> str | None:
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return _texto(valor) or None


def _decimal(valor) -> Decimal | None:
    if valor is None or _texto(valor) == "":
        return None
    return Decimal(str(valor))


def _numero_o_none(valor) -> int | None:
    """Grupo/Depto: número, o '_GND'/'_DND' (sin grupo/depto) -> None."""
    return valor if isinstance(valor, int) else None


def _laboratorio(valor) -> str | None:
    """La columna Localiza/Localización trae nombres de laboratorio (MAVER,
    BAYER) mezclados con fechas viejas ("C/05/16", celdas de fecha, números).
    Solo se acepta texto sin dígitos."""
    if isinstance(valor, (datetime, date)) or not isinstance(valor, str):
        return None
    texto = _texto(valor)
    if not texto or re.search(r"\d", texto):
        return None
    return texto


# --- Lectura de archivos --------------------------------------------------

def _renglones(ruta: Path, columnas: dict[str, int], encabezados: dict[str, str]):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb["Hoja1"]
    except KeyError:
        raise FormatoInvalido(f"{ruta.name}: no tiene la hoja 'Hoja1'")

    encabezado = next(ws.iter_rows(min_row=FILA_ENCABEZADO, max_row=FILA_ENCABEZADO, values_only=True))
    for nombre, esperado in encabezados.items():
        encontrado = _texto(encabezado[columnas[nombre]])
        if encontrado != esperado:
            raise FormatoInvalido(
                f"{ruta.name}: en la fila {FILA_ENCABEZADO} se esperaba la columna "
                f"'{esperado}' y se encontró '{encontrado}'"
            )

    for fila in ws.iter_rows(min_row=FILA_DATOS, values_only=True):
        # Columna 0 ocupada = subtotal "Letra: X"; sin descripción = vacío o total general.
        if fila[0] is None and _texto(fila[columnas["descripcion"]]):
            yield fila
    wb.close()


def leer_pvwin(ruta_catalogo: Path, ruta_inventario: Path) -> ResultadoLectura:
    catalogo = list(_renglones(ruta_catalogo, CAT, {"clave": "Clave", "costo": "$Pcio Compra", "iva": "%IVA"}))
    inventario = list(_renglones(ruta_inventario, INV, {"clave": "Clave", "existencia": "Existencia"}))

    omitidos: list[Incidencia] = []
    fusionados: list[Incidencia] = []
    corregidos: list[Incidencia] = []
    negativos: list[Incidencia] = []

    # Una clave que aparece con descripciones distintas no identifica a un
    # producto (ej. clave "1" y "0" en PVWin): esos productos se importan sin
    # clave y quedan para asignarles una nueva.
    # clave -> {nombre comparable: nombre como viene en PVWin}
    nombres_por_clave: dict[str, dict[str, str]] = defaultdict(dict)
    for filas, col in ((catalogo, CAT), (inventario, INV)):
        for fila in filas:
            if clave := _clave(fila[col["clave"]]):
                nombre = _texto(fila[col["descripcion"]])
                nombres_por_clave[clave].setdefault(_comparable(nombre), nombre)
    claves_compartidas = {
        c: list(nombres.values()) for c, nombres in nombres_por_clave.items() if len(nombres) > 1
    }

    # Un producto es (clave, nombre comparable): así dos productos distintos con
    # la misma clave no se mezclan, y un renglón repetido sí se fusiona.
    productos: dict[tuple[str | None, str], ProductoPVWin] = {}

    for fila in catalogo:
        clave, nombre = _clave(fila[CAT["clave"]]), _texto(fila[CAT["descripcion"]])
        if nombre in OMITIR:
            omitidos.append(Incidencia(clave, nombre, "producto de prueba (catálogo)"))
            continue
        if (clave, _comparable(nombre)) in productos:
            fusionados.append(Incidencia(clave, nombre, "renglón repetido en el catálogo"))
            continue
        productos[(clave, _comparable(nombre))] = ProductoPVWin(
            clave=clave,
            nombre=nombre,
            clave_sat=_clave(fila[CAT["prodserv"]]),
            laboratorio=_laboratorio(fila[CAT["localiza"]]),
            grupo=_numero_o_none(fila[CAT["gpo"]]),
            depto=_numero_o_none(fila[CAT["dpto"]]),
            costo=_decimal(fila[CAT["costo"]]),
            iva=_decimal(fila[CAT["iva"]]),
            ieps=_decimal(fila[CAT["ieps"]]) or Decimal(0),
            minimo=_decimal(fila[CAT["minimo"]]),
            maximo=_decimal(fila[CAT["maximo"]]),
            en_catalogo=True,
        )

    for fila in inventario:
        clave, nombre = _clave(fila[INV["clave"]]), _texto(fila[INV["descripcion"]])
        existencia = _decimal(fila[INV["existencia"]]) or Decimal(0)
        if nombre in OMITIR:
            omitidos.append(Incidencia(clave, nombre, f"producto de prueba (inventario, existencia {existencia})"))
            continue
        producto = productos.get((clave, _comparable(nombre)))
        if producto is None:
            producto = productos[(clave, _comparable(nombre))] = ProductoPVWin(
                clave=clave,
                nombre=nombre,
                laboratorio=_laboratorio(fila[INV["localizacion"]]),
                grupo=_numero_o_none(fila[INV["grupo"]]),
                depto=_numero_o_none(fila[INV["depto"]]),
            )
        elif producto.en_inventario:
            fusionados.append(Incidencia(
                clave, producto.nombre,
                f"repetido en inventario ('{nombre}'): existencias {producto.existencia} + {existencia}",
            ))
        else:
            # Viene del catálogo; el laboratorio puede venir solo en el inventario.
            producto.laboratorio = producto.laboratorio or _laboratorio(fila[INV["localizacion"]])
        producto.existencia += existencia
        producto.en_inventario = True

    for producto in productos.values():
        _limpiar(producto, claves_compartidas, corregidos, negativos)

    return ResultadoLectura(
        productos=list(productos.values()),
        omitidos=omitidos,
        fusionados=fusionados,
        corregidos=corregidos,
        negativos=negativos,
        renglones_catalogo=len(catalogo),
        renglones_inventario=len(inventario),
    )


def _limpiar(
    p: ProductoPVWin,
    claves_compartidas: dict[str, list[str]],  # clave -> nombres de los productos que la usan
    corregidos: list[Incidencia],
    negativos: list[Incidencia],
) -> None:
    """Aplica las reglas de limpieza y anota en p.revision lo que necesita ojo humano."""
    if p.clave is None:
        p.revision.append("sin clave en PVWin; asignar clave")
    elif p.clave in claves_compartidas:
        otros = [n for n in claves_compartidas[p.clave] if _comparable(n) != _comparable(p.nombre)]
        p.revision.append(f"la clave '{p.clave}' también la usa {', '.join(otros)}; asignar clave nueva")
        p.clave = None

    if p.grupo is None:
        p.revision.append("sin grupo en PVWin; revisar IVA y categoría")

    if p.iva is None:
        # No viene en el catálogo: se deduce del grupo (1 = medicamento 0%, 2 = 16%).
        p.iva = IVA_POR_GRUPO.get(p.grupo, Decimal(0))
    elif p.iva in CORRECCIONES_IVA:
        corregido = CORRECCIONES_IVA[p.iva]
        corregidos.append(Incidencia(p.clave, p.nombre, f"IVA {p.iva}% corregido a {corregido}%"))
        p.iva = corregido

    if p.en_catalogo and not p.costo:
        p.costo = None
        p.revision.append("precio de compra en cero en PVWin")

    if p.existencia < 0:
        negativos.append(Incidencia(p.clave, p.nombre, str(p.existencia)))
        p.revision.append(f"existencia negativa en PVWin ({p.existencia}); se importó en cero, contar físicamente")
        p.existencia = Decimal(0)
    elif p.existencia > EXISTENCIA_SOSPECHOSA:
        p.revision.append(f"existencia muy alta ({p.existencia}); verificar")
