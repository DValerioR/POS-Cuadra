"""Actualiza el catálogo con la exportación completa de PVWin (los zip con
archivos DBF: art001.zip, pre001.zip, sat001.zip y loc001.zip) y crea los
productos que falten.

Uso:
    python -m app.scripts.importar_catalogo_pvwin --negocio 1 --carpeta datos/catalogos [--simular]

Reglas (ver CONTEXTO.md, "Importación desde PVWin"):
- PVWin manda en nombre, costo, IVA, IEPS, clave SAT, mínimo y máximo. El
  laboratorio solo se llena si el producto no tenía.
- Precio de venta = "Precio Venta 1" (sin impuestos) + IEPS + IVA, con el
  redondeo del negocio. Si el último cambio de precio lo hizo una persona en
  el sistema, se respeta su precio.
- Las existencias no se tocan (la exportación no las trae).
- La categoría no se toca, salvo en los productos que se habían creado solo
  con la lista de precios (sin departamento): esos se clasifican otra vez,
  ahora con el departamento de PVWin. Los nuevos se clasifican igual.
- Corrige las claves a las que el Excel les quitó los ceros de la izquierda
  (ej. 70942302388 -> 070942302388), para que el escáner las encuentre.

Se puede correr otra vez: solo cambia lo que sea distinto.
Al final escribe un reporte en Excel (por defecto en datos/).
"""

import argparse
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.importador.pvwin import FormatoInvalido, _comparable
from app.importador.pvwin_dbf import ArticuloPVWin, CatalogoPVWin, leer_catalogo
from app.models import Categoria, Negocio, PrecioHistorial, Producto
from app.scripts.clasificar_categorias import asegurar_categorias, clasificar
from app.scripts.importar_precios_pvwin import MOTIVO_NUEVO, precio_final
from app.services.catalogo import registrar_precio

ORIGEN = "Catálogo completo de PVWin"
# Motivos de revisión que vienen de datos de PVWin: se quitan y se vuelven a
# poner según lo que diga la exportación nueva. Los demás (existencias
# negativas, marcas puestas a mano) se respetan.
MOTIVOS_DE_PVWIN = (
    MOTIVO_NUEVO, "precio de compra en cero en PVWin", "sin grupo en PVWin", "IVA de ",
    "PVWin lo tiene en dólares", "precio de venta en cero en PVWin", "sin clave en PVWin", "la clave '",
    "asignar clave nueva",  # segunda parte de "la clave 'X' también la usa ...; asignar clave nueva"
)


@dataclass
class Resultado:
    actualizados: int = 0
    precios: list[tuple] = field(default_factory=list)  # clave, nombre, sin impuestos, IVA, IEPS, antes, ahora
    impuestos: list[tuple] = field(default_factory=list)  # clave, nombre, IVA antes, IVA ahora, IEPS antes, IEPS ahora
    claves: list[tuple] = field(default_factory=list)  # clave antes, clave ahora, nombre
    nombres: list[tuple] = field(default_factory=list)  # clave, antes, ahora
    costos: list[tuple] = field(default_factory=list)  # clave, nombre, antes, ahora
    claves_sat: int = 0
    categorias: list[tuple] = field(default_factory=list)  # clave, nombre, antes, ahora, regla
    nuevos: list[tuple] = field(default_factory=list)  # clave, nombre, categoría, precio
    precio_manual: list[tuple] = field(default_factory=list)  # clave, nombre, precio actual, precio PVWin
    sin_precio: list[tuple] = field(default_factory=list)  # clave, nombre
    otros_precios: list[tuple] = field(default_factory=list)  # clave, nombre, precio 1, otros
    revisar: list[tuple] = field(default_factory=list)  # clave, nombre, motivos
    revision_quitada: int = 0
    no_en_pvwin: list[tuple] = field(default_factory=list)  # clave, nombre


def _motivos(producto: Producto) -> list[str]:
    return [m for m in (producto.motivo_revision or "").split("; ") if m]


def _poner_motivos(producto: Producto, motivos: list[str]) -> None:
    producto.motivo_revision = "; ".join(motivos) or None
    producto.requiere_revision = bool(motivos)


def _precio_puesto_a_mano(db: Session, negocio_id: int) -> set[int]:
    """Productos cuyo último cambio de precio lo hizo una persona (no un script)."""
    ultimos: dict[int, int | None] = {}
    for producto_id, usuario_id in db.execute(
        select(PrecioHistorial.producto_id, PrecioHistorial.usuario_id)
        .where(PrecioHistorial.negocio_id == negocio_id)
        .order_by(PrecioHistorial.created_at, PrecioHistorial.id)
    ):
        ultimos[producto_id] = usuario_id
    return {p for p, u in ultimos.items() if u is not None}


def aplicar_catalogo(db: Session, negocio: Negocio, catalogo: CatalogoPVWin) -> Resultado:
    """Guarda en la sesión (sin commit) los cambios y los productos nuevos."""
    r = Resultado()
    productos = db.scalars(select(Producto).where(Producto.negocio_id == negocio.id).order_by(Producto.id)).all()
    a_mano = _precio_puesto_a_mano(db, negocio.id)
    categorias = asegurar_categorias(db, negocio.id)
    nombre_categoria = {c.id: c.nombre for c in categorias.values()}
    depto_de_categoria = {c.id: c.pvwin_depto for c in categorias.values() if c.pvwin_depto is not None}

    por_par = {(p.clave, _comparable(p.nombre)): p for p in productos}
    por_clave: dict[str, list[Producto]] = defaultdict(list)
    por_nombre: dict[str, list[Producto]] = defaultdict(list)
    for p in productos:
        if p.clave:
            por_clave[p.clave].append(p)
        por_nombre[_comparable(p.nombre)].append(p)

    nombres_por_clave: dict[str, set[str]] = defaultdict(set)
    for a in catalogo.articulos:
        if a.clave:
            nombres_por_clave[a.clave].add(a.comparable)
    claves_compartidas = {c for c, nombres in nombres_por_clave.items() if len(nombres) > 1}
    usados: set[int] = set()
    asignados: dict[int, Producto] = {}  # id(artículo) -> producto de la base
    nombres_pvwin = defaultdict(int)
    for a in catalogo.articulos:
        nombres_pvwin[a.comparable] += 1

    def sin_ceros(clave: str | None) -> str | None:
        """La clave como la dejó el Excel, que les quitó los ceros de la izquierda."""
        if clave and clave.isdigit() and clave.startswith("0"):
            return clave.lstrip("0") or None
        return None

    def asignar(a: ArticuloPVWin, p: Producto | None) -> bool:
        if p is None or p.id in usados:
            return False
        usados.add(p.id)
        asignados[id(a)] = p
        return True

    def solo_libre(lista: list[Producto]) -> Producto | None:
        libres = [p for p in lista if p.id not in usados]
        return libres[0] if len(libres) == 1 else None

    # 1) Misma clave y mismo nombre (también con la clave sin los ceros).
    for a in catalogo.articulos:
        asignar(a, por_par.get((a.clave, a.comparable))) or asignar(a, por_par.get((sin_ceros(a.clave), a.comparable)))
    # 2) Mismo nombre, sin ambigüedad: productos sin clave o con una clave que
    #    el Excel echó a perder (ej. "#NAME?").
    for a in catalogo.articulos:
        if id(a) not in asignados and nombres_pvwin[a.comparable] == 1:
            asignar(a, solo_libre(por_nombre.get(a.comparable, [])))
    # 3) Misma clave con el nombre cambiado en PVWin, solo si el nombre viejo
    #    ya no existe en PVWin (si existe, es otro producto con la clave mal).
    for a in catalogo.articulos:
        if id(a) in asignados or not a.clave or a.clave in claves_compartidas:
            continue
        for clave in (a.clave, sin_ceros(a.clave)):
            p = solo_libre(por_clave.get(clave, [])) if clave else None
            if p and _comparable(p.nombre) not in nombres_pvwin and asignar(a, p):
                break

    def buscar(a: ArticuloPVWin) -> Producto | None:
        return asignados.get(id(a))

    def clave_para(a: ArticuloPVWin, producto: Producto | None) -> tuple[str | None, str | None]:
        """La clave que le toca y, si no se puede usar, el motivo para revisar."""
        if a.clave is None:
            return None, "sin clave en PVWin; asignar clave"
        if a.clave in claves_compartidas:
            return None, f"la clave '{a.clave}' la usan varios productos en PVWin; asignar clave nueva"
        if any(p is not producto for p in por_clave.get(a.clave, [])):
            return None, f"la clave '{a.clave}' ya la tiene otro producto; asignar clave nueva"
        return a.clave, None

    def poner_precio(producto: Producto, a: ArticuloPVWin) -> None:
        if producto.id in a_mano:
            if a.precio > 0:
                pvwin = precio_final(a.precio, producto.iva_porcentaje, producto.ieps_porcentaje,
                                     negocio.redondeo_precio_venta, producto.precio_maximo_publico)
                if pvwin != producto.precio_venta:
                    r.precio_manual.append((producto.clave, producto.nombre, producto.precio_venta, pvwin))
            return
        anterior = producto.precio_venta
        if a.precio <= 0:
            producto.precio_venta = None
            r.sin_precio.append((producto.clave, producto.nombre))
        else:
            producto.precio_venta = precio_final(a.precio, producto.iva_porcentaje, producto.ieps_porcentaje,
                                                 negocio.redondeo_precio_venta, producto.precio_maximo_publico)
        if anterior != producto.precio_venta:
            r.precios.append((producto.clave, producto.nombre, a.precio, producto.iva_porcentaje,
                              producto.ieps_porcentaje, anterior, producto.precio_venta))
            registrar_precio(db, producto, anterior, None, ORIGEN)

    def categoria_para(a: ArticuloPVWin) -> tuple[Categoria, str]:
        c = clasificar(a.nombre, a.laboratorio, a.depto)
        return categorias[c.categoria or "Otros"], c.regla

    # Claves corregidas, antes que todo lo demás: primero se sueltan todas las
    # que cambian y luego se ponen las nuevas, para que un producto no le
    # estorbe a otro (ej. KANKA tenía la clave de HUMULIN sin su cero).
    cambios = [(a, p) for a in catalogo.articulos if (p := asignados.get(id(a))) and p.clave != a.clave]
    anteriores = {p.id: p.clave for _, p in cambios}
    for _, p in cambios:
        if p.clave:
            por_clave[p.clave].remove(p)
        p.clave = None
    db.flush()
    motivo_clave: dict[int, str] = {}
    for a, p in cambios:
        clave, motivo = clave_para(a, p)
        if clave is None and anteriores[p.id] and not por_clave.get(anteriores[p.id]):
            clave = anteriores[p.id]  # se queda con la que tenía
        elif clave is not None:
            r.claves.append((anteriores[p.id], clave, a.nombre))
        if clave is None and motivo:
            motivo_clave[p.id] = motivo
        p.clave = clave
        if clave:
            por_clave[clave].append(p)
    db.flush()

    for a in catalogo.articulos:
        motivos_nuevos = list(a.revision)
        if a.precio <= 0:
            motivos_nuevos.append("precio de venta en cero en PVWin; capturar precio")
        if a.otros_precios:
            r.otros_precios.append((a.clave, a.nombre, a.precio, ", ".join(str(o) for o in a.otros_precios)))

        producto = buscar(a)
        if producto is None:
            clave, motivo = clave_para(a, None)
            if motivo:
                motivos_nuevos.append(motivo)
            categoria, _ = categoria_para(a)
            producto = Producto(
                negocio_id=negocio.id, categoria_id=categoria.id, clave=clave, nombre=a.nombre,
                clave_sat=a.clave_sat, laboratorio=a.laboratorio, costo=a.costo,
                iva_porcentaje=a.iva if a.iva is not None else Decimal(0), ieps_porcentaje=a.ieps,
                minimo=a.minimo, maximo=a.maximo,
            )
            _poner_motivos(producto, motivos_nuevos)
            db.add(producto)
            db.flush()
            usados.add(producto.id)
            if clave:
                por_clave[clave].append(producto)
            poner_precio(producto, a)
            r.nuevos.append((producto.clave, producto.nombre, categoria.nombre, producto.precio_venta))
            if motivos_nuevos:
                r.revisar.append((producto.clave, producto.nombre, "; ".join(motivos_nuevos)))
            continue

        usados.add(producto.id)
        r.actualizados += 1
        motivos_antes = _motivos(producto)
        venia_sin_depto = MOTIVO_NUEVO in motivos_antes

        if motivo := motivo_clave.get(producto.id):
            motivos_nuevos.append(motivo)
        if producto.nombre != a.nombre:
            r.nombres.append((producto.clave, producto.nombre, a.nombre))
            producto.nombre = a.nombre
        if a.clave_sat and producto.clave_sat != a.clave_sat:
            producto.clave_sat = a.clave_sat
            r.claves_sat += 1
        if a.laboratorio and not producto.laboratorio:
            producto.laboratorio = a.laboratorio
        if a.costo is not None and producto.costo != a.costo:
            r.costos.append((producto.clave, producto.nombre, producto.costo, a.costo))
            producto.costo = a.costo
        if a.iva is not None and (producto.iva_porcentaje, producto.ieps_porcentaje) != (a.iva, a.ieps):
            r.impuestos.append((producto.clave, producto.nombre, producto.iva_porcentaje, a.iva,
                                producto.ieps_porcentaje, a.ieps))
            producto.iva_porcentaje, producto.ieps_porcentaje = a.iva, a.ieps
        producto.minimo, producto.maximo = a.minimo, a.maximo

        sin_categoria = producto.categoria_id is None or producto.categoria_id in depto_de_categoria
        if venia_sin_depto or sin_categoria:
            categoria, regla = categoria_para(a)
            if categoria.id != producto.categoria_id:
                r.categorias.append((producto.clave, producto.nombre,
                                     nombre_categoria.get(producto.categoria_id, ""), categoria.nombre, regla))
                producto.categoria_id = categoria.id

        poner_precio(producto, a)

        conservados = [m for m in motivos_antes if not m.startswith(MOTIVOS_DE_PVWIN)]
        motivos = conservados + [m for m in motivos_nuevos if m not in conservados]
        if motivos != motivos_antes:
            if motivos_antes and not motivos:
                r.revision_quitada += 1
            _poner_motivos(producto, motivos)
        if motivos_nuevos:
            r.revisar.append((producto.clave, producto.nombre, "; ".join(motivos_nuevos)))

    r.no_en_pvwin = [(p.clave, p.nombre) for p in productos if p.id not in usados and p.activo]
    db.flush()
    return r


def escribir_reporte(ruta: Path, catalogo: CatalogoPVWin, r: Resultado, simulado: bool, paso: Decimal | None) -> None:
    wb = Workbook()
    negrita = Font(bold=True)

    def hoja(titulo: str, encabezados: list[str], filas, anchos=(18, 50)) -> None:
        ws = wb.create_sheet(titulo)
        ws.append(encabezados)
        for celda in ws[1]:
            celda.font = negrita
        for fila in filas:
            ws.append(list(fila))
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[chr(64 + i)].width = ancho
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Resumen"
    subio = sum(1 for f in r.precios if f[5] is not None and f[6] is not None and f[6] > f[5])
    bajo = sum(1 for f in r.precios if f[5] is not None and f[6] is not None and f[6] < f[5])
    for fila in (
        ("Fecha", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Modo", "SIMULACIÓN (no se guardó nada)" if simulado else "Cambios guardados"),
        ("Redondeo del negocio", f"a {paso} hacia arriba" if paso else "sin redondeo"),
        ("Artículos leídos de PVWin", catalogo.renglones),
        ("Productos que ya estaban (actualizados)", r.actualizados),
        ("Productos nuevos (sin existencia)", len(r.nuevos)),
        ("Precios de venta que cambian", len(r.precios)),
        ("  suben", subio),
        ("  bajan", bajo),
        ("IVA o IEPS que cambian", len(r.impuestos)),
        ("Claves corregidas (ceros a la izquierda y otras)", len(r.claves)),
        ("Nombres que cambian", len(r.nombres)),
        ("Costos que cambian o se llenan", len(r.costos)),
        ("Claves SAT puestas o cambiadas", r.claves_sat),
        ("Categorías asignadas de nuevo (venían sin departamento)", len(r.categorias)),
        ("Precios puestos a mano que se respetaron (PVWin dice otro)", len(r.precio_manual)),
        ("Productos sin precio (en cero en PVWin)", len(r.sin_precio)),
        ("Productos con precio 2, 3 o 4 (se usó el precio 1)", len(r.otros_precios)),
        ("Productos para revisar por datos de PVWin", len(r.revisar)),
        ("Productos a los que se les quitó la marca de revisar", r.revision_quitada),
        ("Productos de la base que no vienen en PVWin (sin cambios)", len(r.no_en_pvwin)),
        ("Renglones repetidos en PVWin", len(catalogo.repetidos)),
        ("Renglones omitidos", len(catalogo.omitidos)),
        ("Existencias", "no se tocan (la exportación no las trae)"),
    ):
        ws.append(fila)
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 34

    hoja("Precios", ["Clave", "Producto", "Precio sin impuestos", "IVA %", "IEPS %", "Precio anterior", "Precio nuevo"],
         r.precios, (18, 50, 18, 8, 8, 15, 15))
    hoja("Impuestos", ["Clave", "Producto", "IVA antes", "IVA ahora", "IEPS antes", "IEPS ahora"],
         r.impuestos, (18, 50, 10, 10, 10, 10))
    hoja("Nuevos", ["Clave", "Producto", "Categoría", "Precio de venta"], r.nuevos, (18, 50, 26, 15))
    hoja("Claves corregidas", ["Clave antes", "Clave ahora", "Producto"], r.claves, (18, 18, 50))
    hoja("Revisar", ["Clave", "Producto", "Qué revisar"], r.revisar, (18, 50, 70))
    hoja("Categorias", ["Clave", "Producto", "Antes", "Ahora", "Por qué"], r.categorias, (18, 50, 24, 24, 45))
    hoja("Costos", ["Clave", "Producto", "Costo antes", "Costo ahora"], r.costos, (18, 50, 14, 14))
    hoja("Nombres", ["Clave", "Nombre antes", "Nombre ahora"], r.nombres, (18, 50, 50))
    hoja("Precio a mano", ["Clave", "Producto", "Precio actual (se queda)", "Precio según PVWin"],
         r.precio_manual, (18, 50, 22, 20))
    hoja("Sin precio", ["Clave", "Producto"], r.sin_precio)
    hoja("Otros precios", ["Clave", "Producto", "Precio 1 (usado)", "Precios 2-4"], r.otros_precios, (18, 50, 15, 20))
    hoja("No vienen en PVWin", ["Clave", "Producto"], r.no_en_pvwin)
    hoja("Repetidos y omitidos", ["Clave", "Producto", "Detalle"],
         [(i.clave, i.nombre, i.detalle) for i in catalogo.repetidos + catalogo.omitidos + catalogo.corregidos],
         (18, 50, 50))

    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def main() -> None:
    parser = argparse.ArgumentParser(description="Actualizar el catálogo con la exportación completa de PVWin (zip con DBF)")
    parser.add_argument("--negocio", type=int, required=True)
    parser.add_argument("--carpeta", type=Path, required=True, help="carpeta con art001.zip, pre001.zip, sat001.zip y loc001.zip")
    parser.add_argument("--simular", action="store_true", help="no guarda nada, solo genera el reporte")
    parser.add_argument("--reporte", type=Path, help="ruta del reporte Excel (por defecto en datos/)")
    args = parser.parse_args()

    try:
        catalogo = leer_catalogo(args.carpeta)
    except FormatoInvalido as e:
        sys.exit(f"Formato inesperado: {e}")

    with SessionLocal() as db:
        negocio = db.get(Negocio, args.negocio)
        if negocio is None:
            sys.exit(f"No existe el negocio {args.negocio}.")
        resultado = aplicar_catalogo(db, negocio, catalogo)
        paso = negocio.redondeo_precio_venta
        if args.simular:
            db.rollback()
        else:
            db.commit()

    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    reporte = args.reporte or Path("datos") / f"reporte_catalogo_pvwin_{sello}{'_simulacion' if args.simular else ''}.xlsx"
    escribir_reporte(reporte, catalogo, resultado, args.simular, paso)

    print("SIMULACIÓN: no se guardó nada." if args.simular else "Cambios guardados.")
    print(f"  Productos actualizados:    {resultado.actualizados}")
    print(f"  Productos nuevos:          {len(resultado.nuevos)}")
    print(f"  Precios que cambian:       {len(resultado.precios)}")
    print(f"  IVA/IEPS que cambian:      {len(resultado.impuestos)}")
    print(f"  Claves corregidas:         {len(resultado.claves)}")
    print(f"  Categorías reasignadas:    {len(resultado.categorias)}")
    print(f"  Para revisar:              {len(resultado.revisar)}")
    print(f"  No vienen en PVWin:        {len(resultado.no_en_pvwin)}")
    print(f"  Reporte: {reporte}")


if __name__ == "__main__":
    main()
