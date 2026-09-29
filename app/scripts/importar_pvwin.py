"""Importa catálogo y existencias desde los reportes de Excel de PVWin.

Uso:
    python -m app.scripts.importar_pvwin --negocio 1 --usuario diego \\
        --catalogo "datos/Catalogo de articulos.xlsx" \\
        --inventario "datos/Reporte de inventario.xlsx" [--simular] [--reemplazar]

--simular     hace todo y genera el reporte, pero no guarda nada.
--reemplazar  borra productos, lotes y ajustes del negocio antes de importar.
              Solo para antes de arrancar con ventas reales. Las categorías no
              se borran, para no perder el nombre/margen que se les haya puesto.

Al final escribe un reporte en Excel (por defecto en datos/) con lo que hizo:
resumen, productos a revisar, lista para conteo físico, fusionados, omitidos
y correcciones.
"""

import argparse
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.importador.pvwin import FormatoInvalido, Incidencia, ResultadoLectura, leer_pvwin
from app.models import AjusteInventario, Categoria, Lote, Negocio, Producto, RolUsuario, TipoAjuste, Usuario

MOTIVO_IMPORTACION = "Existencia inicial importada de PVWin (sin lote ni caducidad)"


def _categorias_por_depto(db: Session, negocio_id: int, deptos: set[int]) -> tuple[dict[int, int], int]:
    """Busca o crea una categoría por departamento de PVWin. Regresa
    {depto: categoria_id} y cuántas se crearon."""
    existentes = {
        c.pvwin_depto: c
        for c in db.scalars(
            select(Categoria).where(Categoria.negocio_id == negocio_id, Categoria.pvwin_depto.in_(deptos))
        )
    }
    creadas = 0
    for depto in sorted(deptos - existentes.keys()):
        categoria = Categoria(negocio_id=negocio_id, nombre=f"Depto {depto}", pvwin_depto=depto)
        db.add(categoria)
        existentes[depto] = categoria
        creadas += 1
    db.flush()
    return {depto: c.id for depto, c in existentes.items()}, creadas


def _borrar_importacion(db: Session, negocio_id: int) -> int:
    db.execute(delete(AjusteInventario).where(AjusteInventario.negocio_id == negocio_id))
    db.execute(delete(Lote).where(Lote.negocio_id == negocio_id))
    return db.execute(delete(Producto).where(Producto.negocio_id == negocio_id)).rowcount


def importar(db: Session, negocio_id: int, usuario_id: int, lectura: ResultadoLectura) -> dict[str, int]:
    """Guarda en la sesión (sin commit) los productos, lotes y ajustes."""
    deptos = {p.depto for p in lectura.productos if p.depto is not None}
    categoria_de, categorias_creadas = _categorias_por_depto(db, negocio_id, deptos)

    pares: list[tuple[Producto, Decimal]] = []
    for p in lectura.productos:
        producto = Producto(
            negocio_id=negocio_id,
            categoria_id=categoria_de.get(p.depto),
            clave=p.clave,
            nombre=p.nombre,
            clave_sat=p.clave_sat,
            laboratorio=p.laboratorio,
            costo=p.costo,
            iva_porcentaje=p.iva,
            ieps_porcentaje=p.ieps,
            minimo=p.minimo,
            maximo=p.maximo,
            requiere_revision=bool(p.revision),
            motivo_revision="; ".join(p.revision) or None,
        )
        db.add(producto)
        pares.append((producto, p.existencia))
    db.flush()  # asigna ids a los productos

    # Cada existencia entra en un lote especial sin lote ni caducidad, para
    # poder vender desde el primer día (ver CONTEXTO.md, migración gradual).
    lotes: list[tuple[Lote, Producto]] = []
    for producto, existencia in pares:
        if existencia > 0:
            lote = Lote(
                negocio_id=negocio_id,
                producto_id=producto.id,
                cantidad=existencia,
                costo_unitario=producto.costo,
            )
            db.add(lote)
            lotes.append((lote, producto))
    db.flush()

    for lote, producto in lotes:
        db.add(AjusteInventario(
            negocio_id=negocio_id,
            producto_id=producto.id,
            lote_id=lote.id,
            usuario_id=usuario_id,
            tipo=TipoAjuste.IMPORTACION,
            cantidad=lote.cantidad,
            motivo=MOTIVO_IMPORTACION,
        ))
    db.flush()

    return {
        "categorias_creadas": categorias_creadas,
        "productos_creados": len(pares),
        "lotes_creados": len(lotes),
        "piezas": sum(l.cantidad for l, _ in lotes),
    }


def escribir_reporte(ruta: Path, lectura: ResultadoLectura, totales: dict, simulado: bool, borrados: int) -> None:
    wb = Workbook()
    negrita = Font(bold=True)

    def hoja(titulo: str, encabezados: list[str], filas) -> None:
        ws = wb.create_sheet(titulo)
        ws.append(encabezados)
        for celda in ws[1]:
            celda.font = negrita
        for fila in filas:
            ws.append(list(fila))
        for i, ancho in enumerate([18, 50, 80][: len(encabezados)], start=1):
            ws.column_dimensions[chr(64 + i)].width = ancho
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Resumen"
    revisar = [p for p in lectura.productos if p.revision]
    filas_resumen = [
        ("Fecha", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Modo", "SIMULACIÓN (no se guardó nada)" if simulado else "Importación guardada"),
        ("Renglones leídos del catálogo", lectura.renglones_catalogo),
        ("Renglones leídos del inventario", lectura.renglones_inventario),
        ("Productos borrados antes de importar (--reemplazar)", borrados),
        ("Productos creados", totales["productos_creados"]),
        ("  que venían en catálogo e inventario", sum(1 for p in lectura.productos if p.en_catalogo and p.en_inventario)),
        ("  solo en catálogo (existencia cero)", sum(1 for p in lectura.productos if p.en_catalogo and not p.en_inventario)),
        ("  solo en inventario (sin costo; IVA 0% porque no viene marcado)", sum(1 for p in lectura.productos if not p.en_catalogo)),
        ("Categorías creadas (Depto N)", totales["categorias_creadas"]),
        ("Lotes 'sin caducidad' creados", totales["lotes_creados"]),
        ("Piezas importadas", totales["piezas"]),
        ("Renglones fusionados", len(lectura.fusionados)),
        ("Renglones omitidos", len(lectura.omitidos)),
        ("Existencias negativas puestas en cero", len(lectura.negativos)),
        ("Correcciones automáticas", len(lectura.corregidos)),
        ("Productos marcados para revisar", len(revisar)),
    ]
    for fila in filas_resumen:
        ws.append(fila)
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 30

    hoja("Revisar", ["Clave", "Producto", "Qué revisar"], ((p.clave, p.nombre, "; ".join(p.revision)) for p in revisar))
    hoja(
        "Conteo fisico",
        ["Clave", "Producto", "Existencia en PVWin", "Conteo físico"],
        ((i.clave, i.nombre, Decimal(i.detalle), None) for i in lectura.negativos),
    )
    for titulo, incidencias in (
        ("Fusionados", lectura.fusionados),
        ("Omitidos", lectura.omitidos),
        ("Correcciones", lectura.corregidos),
    ):
        hoja(titulo, ["Clave", "Producto", "Detalle"], _filas(incidencias))

    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def _filas(incidencias: list[Incidencia]):
    return ((i.clave, i.nombre, i.detalle) for i in incidencias)


def main() -> None:
    parser = argparse.ArgumentParser(description="Importar catálogo y existencias de PVWin")
    parser.add_argument("--negocio", type=int, required=True)
    parser.add_argument("--usuario", required=True, help="usuario administrador que hace la importación")
    parser.add_argument("--catalogo", type=Path, required=True)
    parser.add_argument("--inventario", type=Path, required=True)
    parser.add_argument("--simular", action="store_true", help="no guarda nada, solo genera el reporte")
    parser.add_argument("--reemplazar", action="store_true", help="borra lo importado antes (productos, lotes, ajustes)")
    parser.add_argument("--reporte", type=Path, help="ruta del reporte Excel (por defecto en datos/)")
    args = parser.parse_args()

    for ruta in (args.catalogo, args.inventario):
        if not ruta.exists():
            sys.exit(f"No existe el archivo: {ruta}")

    with SessionLocal() as db:
        if db.get(Negocio, args.negocio) is None:
            sys.exit(f"No existe el negocio {args.negocio}.")
        usuario = db.scalar(
            select(Usuario).where(Usuario.negocio_id == args.negocio, Usuario.nombre_usuario == args.usuario)
        )
        if usuario is None or usuario.rol != RolUsuario.ADMIN:
            sys.exit(f"'{args.usuario}' no es un usuario administrador de ese negocio.")

        existentes = db.scalar(select(func.count()).select_from(Producto).where(Producto.negocio_id == args.negocio))
        if existentes and not args.reemplazar:
            sys.exit(
                f"El negocio ya tiene {existentes} productos. Usa --reemplazar para borrarlos e importar de nuevo "
                "(solo antes de tener ventas reales)."
            )

        try:
            lectura = leer_pvwin(args.catalogo, args.inventario)
        except FormatoInvalido as e:
            sys.exit(f"Formato inesperado: {e}")

        borrados = _borrar_importacion(db, args.negocio) if args.reemplazar else 0
        totales = importar(db, args.negocio, usuario.id, lectura)

        if args.simular:
            db.rollback()
        else:
            db.commit()

    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    reporte = args.reporte or Path("datos") / f"reporte_importacion_{sello}{'_simulacion' if args.simular else ''}.xlsx"
    escribir_reporte(reporte, lectura, totales, args.simular, borrados)

    print("SIMULACIÓN: no se guardó nada." if args.simular else "Importación guardada.")
    print(f"  Productos creados:        {totales['productos_creados']}")
    print(f"  Categorías creadas:       {totales['categorias_creadas']}")
    print(f"  Piezas importadas:        {totales['piezas']} en {totales['lotes_creados']} lotes sin caducidad")
    print(f"  Negativos puestos en 0:   {len(lectura.negativos)}")
    print(f"  Productos para revisar:   {sum(1 for p in lectura.productos if p.revision)}")
    print(f"  Reporte: {reporte}")


if __name__ == "__main__":
    main()
