"""Pone los precios de venta desde la lista de precios de PVWin, y crea los
productos de la lista que todavía no existen.

Uso:
    python -m app.scripts.importar_precios_pvwin --negocio 1 \\
        --lista "datos/Catalogo completo con precios.xlsx" \\
        --catalogo "datos/Catalogo de articulos.xlsx" [--simular]

Reglas (ver CONTEXTO.md, "Importación desde PVWin"):
- La lista trae el precio sin impuestos ("Precio Venta 1"). El precio de
  venta es ese precio + IEPS + IVA, redondeado como diga el negocio (por
  ejemplo a pesos enteros hacia arriba, sin pasar el precio máximo).
- IVA e IEPS: solo los que el catálogo de artículos marca explícitamente;
  si no vienen marcados son 0%. Esto también corrige a los productos que ya
  estaban en la base con un IVA deducido por su grupo.
- Los productos que no están en la base se crean sin existencia, con IVA 0%
  y marcados para revisar (no traen costo ni departamento).
- Precio en cero: el producto queda sin precio (no se puede vender) y marcado.

Se puede correr otra vez: actualiza precios y ya no vuelve a crear productos.
Al final escribe un reporte en Excel (por defecto en datos/).
"""

import argparse
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.importador.pvwin import (
    OMITIR, FormatoInvalido, PrecioPVWin, _comparable, impuestos_del_catalogo, leer_lista_precios,
)
from app.models import Negocio, Producto
from app.services.precios import redondear_precio_venta

CENTAVO = Decimal("0.01")
MOTIVO_NUEVO = "solo venía en la lista de precios de PVWin: sin costo ni departamento"


@dataclass
class ResultadoPrecios:
    renglones: int = 0
    con_precio: list[tuple] = field(default_factory=list)  # clave, nombre, base, IVA, IEPS, anterior, nuevo
    impuestos_corregidos: list[tuple] = field(default_factory=list)  # clave, nombre, IVA antes, IVA ahora, IEPS antes, IEPS ahora
    nuevos: list[tuple] = field(default_factory=list)  # clave, nombre, precio
    sin_precio: list[tuple] = field(default_factory=list)  # clave, nombre, detalle
    otros_precios: list[tuple] = field(default_factory=list)  # clave, nombre, precio 1, otros
    repetidos: list[tuple] = field(default_factory=list)  # clave, nombre, detalle
    omitidos: list[tuple] = field(default_factory=list)
    no_en_lista: list[tuple] = field(default_factory=list)  # clave, nombre


def precio_final(base: Decimal, iva: Decimal, ieps: Decimal, paso: Decimal | None, maximo: Decimal | None) -> Decimal:
    """Precio sin impuestos -> precio de venta: IEPS sobre la base, IVA sobre
    base + IEPS (igual que el desglose del ticket), y el redondeo del negocio."""
    iva, ieps = Decimal(iva), Decimal(ieps)  # en productos recién creados llegan como 0 (int)
    con_impuestos = (base * (1 + ieps / 100) * (1 + iva / 100)).quantize(CENTAVO, ROUND_HALF_UP)
    return redondear_precio_venta(con_impuestos, paso, maximo)


def _marcar(producto: Producto, motivo: str) -> None:
    motivos = [m for m in (producto.motivo_revision or "").split("; ") if m]
    if motivo not in motivos:
        motivos.append(motivo)
    producto.requiere_revision = True
    producto.motivo_revision = "; ".join(motivos)


def aplicar_precios(
    db: Session, negocio: Negocio, lista: list[PrecioPVWin],
    impuestos: dict[tuple[str | None, str], tuple[Decimal, Decimal]],
) -> ResultadoPrecios:
    """Guarda en la sesión (sin commit) precios, impuestos y productos nuevos."""
    r = ResultadoPrecios(renglones=len(lista))
    productos = db.scalars(select(Producto).where(Producto.negocio_id == negocio.id)).all()

    # Impuestos marcados: por clave + nombre, o solo por nombre si es único
    # (productos a los que se les quitó la clave por estar repetida).
    nombres_catalogo = Counter(nombre for _, nombre in impuestos)
    impuestos_por_nombre = {nombre: v for (_, nombre), v in impuestos.items() if nombres_catalogo[nombre] == 1}

    def impuestos_de(clave: str | None, nombre: str) -> tuple[Decimal, Decimal]:
        comparable = _comparable(nombre)
        return impuestos.get((clave, comparable)) or impuestos_por_nombre.get(comparable) or (Decimal(0), Decimal(0))

    # 1) Impuestos de lo que ya está en la base: solo lo explícito.
    for p in productos:
        iva, ieps = impuestos_de(p.clave, p.nombre)
        if (p.iva_porcentaje, p.ieps_porcentaje) != (iva, ieps):
            r.impuestos_corregidos.append((p.clave, p.nombre, p.iva_porcentaje, iva, p.ieps_porcentaje, ieps))
            p.iva_porcentaje, p.ieps_porcentaje = iva, ieps

    # 2) A qué producto de la base corresponde cada renglón de la lista.
    por_par = {(p.clave, _comparable(p.nombre)): p for p in productos}
    por_clave = defaultdict(list)
    for p in productos:
        if p.clave:
            por_clave[p.clave].append(p)
    sin_clave_por_nombre = defaultdict(list)
    for p in productos:
        if not p.clave:
            sin_clave_por_nombre[_comparable(p.nombre)].append(p)
    nombres_por_clave_lista = defaultdict(set)
    for fila in lista:
        if fila.clave:
            nombres_por_clave_lista[fila.clave].add(_comparable(fila.nombre))

    def buscar(fila: PrecioPVWin) -> Producto | None:
        comparable = _comparable(fila.nombre)
        if (fila.clave, comparable) in por_par:
            return por_par[(fila.clave, comparable)]
        # Mismo código de barras con el nombre cambiado.
        if fila.clave and len(por_clave.get(fila.clave, [])) == 1 and len(nombres_por_clave_lista[fila.clave]) == 1:
            return por_clave[fila.clave][0]
        # Se le quitó la clave al importar porque otro producto la usaba.
        candidatos = sin_clave_por_nombre.get(comparable, [])
        return candidatos[0] if len(candidatos) == 1 else None

    vistos: dict[int, Decimal] = {}  # producto -> precio sin impuestos con el que quedó
    for fila in lista:
        if fila.nombre in OMITIR:
            r.omitidos.append((fila.clave, fila.nombre, "producto de prueba"))
            continue
        producto = buscar(fila)
        nuevo = producto is None
        if nuevo:
            clave = fila.clave
            motivo = MOTIVO_NUEVO
            if clave and (clave in por_clave or len(nombres_por_clave_lista[clave]) > 1):
                otros = [p.nombre for p in por_clave.get(clave, [])] or ["otro producto de la lista"]
                motivo += f"; la clave '{clave}' también la usa {', '.join(otros)}; asignar clave nueva"
                clave = None
            producto = Producto(negocio_id=negocio.id, clave=clave, nombre=fila.nombre)
            _marcar(producto, motivo)
            db.add(producto)
            db.flush()
            por_par[(producto.clave, _comparable(producto.nombre))] = producto
            if producto.clave:
                por_clave[producto.clave].append(producto)
            else:
                sin_clave_por_nombre[_comparable(producto.nombre)].append(producto)
        elif producto.id in vistos:
            if vistos[producto.id] != fila.precio:
                r.repetidos.append((fila.clave, fila.nombre, f"repetido con otro precio ({vistos[producto.id]} y {fila.precio}); se usó el primero"))
            else:
                r.repetidos.append((fila.clave, fila.nombre, "renglón repetido en la lista"))
            continue
        vistos[producto.id] = fila.precio

        if fila.otros_precios:
            r.otros_precios.append((fila.clave, fila.nombre, fila.precio, ", ".join(str(o) for o in fila.otros_precios)))

        anterior = producto.precio_venta
        if fila.precio <= 0:
            producto.precio_venta = None
            _marcar(producto, "precio de venta en cero en PVWin; capturar precio")
            r.sin_precio.append((producto.clave, producto.nombre, "precio en cero en PVWin"))
        else:
            producto.precio_venta = precio_final(
                fila.precio, producto.iva_porcentaje, producto.ieps_porcentaje,
                negocio.redondeo_precio_venta, producto.precio_maximo_publico,
            )
            r.con_precio.append((producto.clave, producto.nombre, fila.precio, producto.iva_porcentaje,
                                 producto.ieps_porcentaje, anterior, producto.precio_venta))
        if nuevo:
            r.nuevos.append((producto.clave, producto.nombre, producto.precio_venta))

    r.no_en_lista = [(p.clave, p.nombre) for p in productos if p.id not in vistos]
    db.flush()
    return r


def escribir_reporte(ruta: Path, r: ResultadoPrecios, simulado: bool, paso: Decimal | None) -> None:
    wb = Workbook()
    negrita = Font(bold=True)

    def hoja(titulo: str, encabezados: list[str], filas, anchos=(18, 50)) -> None:
        ws = wb.create_sheet(titulo)
        ws.append(encabezados)
        for celda in ws[1]:
            celda.font = negrita
        for fila in filas:
            ws.append(list(fila))
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[chr(64 + i)].width = ancho
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Resumen"
    quitado = sum(1 for c in r.impuestos_corregidos if c[3] == 0 and c[2] > 0)
    for fila in (
        ("Fecha", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Modo", "SIMULACIÓN (no se guardó nada)" if simulado else "Precios guardados"),
        ("Redondeo del negocio", f"a {paso} hacia arriba" if paso else "sin redondeo"),
        ("Renglones leídos de la lista", r.renglones),
        ("Productos con precio", len(r.con_precio)),
        ("Productos nuevos creados (sin existencia, IVA 0%, para revisar)", len(r.nuevos)),
        ("Productos sin precio (en cero en PVWin)", len(r.sin_precio)),
        ("Impuestos corregidos (solo lo marcado en el catálogo)", len(r.impuestos_corregidos)),
        ("  de ellos, IVA quitado porque no venía marcado", quitado),
        ("Productos con precio 2, 3 o 4 (se usó el precio 1)", len(r.otros_precios)),
        ("Renglones repetidos en la lista", len(r.repetidos)),
        ("Renglones omitidos", len(r.omitidos)),
        ("Productos de la base que no vienen en la lista (sin cambio de precio)", len(r.no_en_lista)),
    ):
        ws.append(fila)
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 30

    hoja("Precios", ["Clave", "Producto", "Precio sin impuestos", "IVA %", "IEPS %", "Precio anterior", "Precio de venta"],
         r.con_precio, (18, 50, 18, 8, 8, 15, 15))
    hoja("Impuestos corregidos", ["Clave", "Producto", "IVA antes", "IVA ahora", "IEPS antes", "IEPS ahora"],
         r.impuestos_corregidos, (18, 50, 10, 10, 10, 10))
    hoja("Nuevos", ["Clave", "Producto", "Precio de venta"], r.nuevos, (18, 50, 15))
    hoja("Sin precio", ["Clave", "Producto", "Detalle"], r.sin_precio, (18, 50, 40))
    hoja("Otros precios", ["Clave", "Producto", "Precio 1 (usado)", "Precios 2-4"], r.otros_precios, (18, 50, 15, 20))
    hoja("Repetidos", ["Clave", "Producto", "Detalle"], r.repetidos, (18, 50, 60))
    hoja("Omitidos", ["Clave", "Producto", "Detalle"], r.omitidos, (18, 50, 30))
    hoja("No en la lista", ["Clave", "Producto"], r.no_en_lista)

    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def main() -> None:
    parser = argparse.ArgumentParser(description="Poner precios de venta desde la lista de precios de PVWin")
    parser.add_argument("--negocio", type=int, required=True)
    parser.add_argument("--lista", type=Path, required=True, help="Reporte de lista de precios (impuestos no incluidos)")
    parser.add_argument("--catalogo", type=Path, required=True, help="Catálogo de artículos (de ahí salen IVA e IEPS)")
    parser.add_argument("--simular", action="store_true", help="no guarda nada, solo genera el reporte")
    parser.add_argument("--reporte", type=Path, help="ruta del reporte Excel (por defecto en datos/)")
    args = parser.parse_args()

    for ruta in (args.lista, args.catalogo):
        if not ruta.exists():
            sys.exit(f"No existe el archivo: {ruta}")
    try:
        lista = leer_lista_precios(args.lista)
        impuestos = impuestos_del_catalogo(args.catalogo)
    except FormatoInvalido as e:
        sys.exit(f"Formato inesperado: {e}")

    with SessionLocal() as db:
        negocio = db.get(Negocio, args.negocio)
        if negocio is None:
            sys.exit(f"No existe el negocio {args.negocio}.")
        resultado = aplicar_precios(db, negocio, lista, impuestos)
        paso = negocio.redondeo_precio_venta
        if args.simular:
            db.rollback()
        else:
            db.commit()

    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    reporte = args.reporte or Path("datos") / f"reporte_precios_{sello}{'_simulacion' if args.simular else ''}.xlsx"
    escribir_reporte(reporte, resultado, args.simular, paso)

    print("SIMULACIÓN: no se guardó nada." if args.simular else "Precios guardados.")
    print(f"  Productos con precio:       {len(resultado.con_precio)}")
    print(f"  Productos nuevos:           {len(resultado.nuevos)}")
    print(f"  Sin precio (cero en PVWin): {len(resultado.sin_precio)}")
    print(f"  Impuestos corregidos:       {len(resultado.impuestos_corregidos)}")
    print(f"  Reporte: {reporte}")


if __name__ == "__main__":
    main()
