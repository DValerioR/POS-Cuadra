from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.api.entradas import _exacto
from app.core.auth import usuario_actual
from app.core.database import get_db
from app.models import Usuario
from app.services import pedidos
from app.services.errores import ERRORES_NEGOCIO, a_http

# La pantalla es /pedidos; el API va aparte para no chocar con ella.
router = APIRouter(prefix="/compras/pedidos", tags=["pedidos a proveedores"])


class RenglonIn(BaseModel):
    producto_id: int
    cantidad: Decimal


class PedidoIn(BaseModel):
    proveedor_id: int
    notas: str | None = None
    renglones: list[RenglonIn]


class EntradaLigaIn(BaseModel):
    entrada_id: int


def _renglones(datos: PedidoIn) -> list[pedidos.RenglonPedido]:
    return [pedidos.RenglonPedido(r.producto_id, r.cantidad) for r in datos.renglones]


def _hacer(db: Session, usuario: Usuario, accion) -> dict:
    """Corre la acción, hace commit y regresa el detalle del pedido."""
    try:
        pedido = accion()
    except ERRORES_NEGOCIO as e:
        db.rollback()
        raise a_http(e)
    db.commit()
    return _exacto(pedidos.detalle(db, usuario, pedido.id))


@router.get("")
def listar(estado: str | None = None, proveedor_id: int | None = None, limite: int = Query(50, le=200),
           usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    try:
        return _exacto(pedidos.listar(db, usuario, estado, proveedor_id, limite))
    except ERRORES_NEGOCIO as e:
        raise a_http(e)


@router.get("/preparar")
def preparar(proveedor_id: int, sugerencias: bool = False, usuario: Usuario = Depends(usuario_actual),
             db: Session = Depends(get_db)):
    """Los faltantes del proveedor para armar un pedido nuevo (con las
    cantidades del asistente si sugerencias=true)."""
    try:
        return _exacto(pedidos.preparar(db, usuario, proveedor_id, sugerencias))
    except ERRORES_NEGOCIO as e:
        raise a_http(e)


@router.post("", status_code=201)
def crear(datos: PedidoIn, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    return _hacer(db, usuario, lambda: pedidos.crear(db, usuario, datos.proveedor_id, _renglones(datos), datos.notas))


@router.get("/{pedido_id}")
def detalle(pedido_id: int, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    try:
        return _exacto(pedidos.detalle(db, usuario, pedido_id))
    except ERRORES_NEGOCIO as e:
        raise a_http(e)


@router.put("/{pedido_id}")
def actualizar(pedido_id: int, datos: PedidoIn, usuario: Usuario = Depends(usuario_actual),
               db: Session = Depends(get_db)):
    return _hacer(db, usuario, lambda: pedidos.actualizar(db, usuario, pedido_id, _renglones(datos), datos.notas))


@router.post("/{pedido_id}/enviar")
def enviar(pedido_id: int, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    return _hacer(db, usuario, lambda: pedidos.enviar(db, usuario, pedido_id))


@router.post("/{pedido_id}/cancelar")
def cancelar(pedido_id: int, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    return _hacer(db, usuario, lambda: pedidos.cancelar(db, usuario, pedido_id))


@router.post("/{pedido_id}/cerrar")
def cerrar(pedido_id: int, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    return _hacer(db, usuario, lambda: pedidos.cerrar(db, usuario, pedido_id))


@router.post("/{pedido_id}/entradas")
def ligar_entrada(pedido_id: int, datos: EntradaLigaIn, usuario: Usuario = Depends(usuario_actual),
                  db: Session = Depends(get_db)):
    """Liga una entrada ya registrada (factura de ese proveedor) al pedido."""
    return _hacer(db, usuario, lambda: pedidos.ligar_entrada(db, usuario, pedido_id, datos.entrada_id))


@router.get("/{pedido_id}/excel")
def excel(pedido_id: int, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    """El pedido en Excel, para capturarlo en el portal del proveedor o mandarlo."""
    try:
        p = pedidos.detalle(db, usuario, pedido_id)
    except ERRORES_NEGOCIO as e:
        raise a_http(e)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Pedido {p['folio']}"
    ws.append([f"Pedido {p['folio']} para {p['proveedor']}", f"{p['created_at']:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    if p["notas"]:
        ws.append([p["notas"]])
    ws.append([])
    ws.append(["Clave", "Producto", "Piezas", "Costo esperado por pieza", "Importe estimado"])
    encabezado = ws.max_row
    for celda in ws[encabezado]:
        celda.font = Font(bold=True)
    for r in p["renglones"]:
        ws.append([r["clave"], r["nombre"], r["cantidad"], r["costo_esperado"], r["importe"]])
    ws.append([])
    ws.append(["", "Total estimado (sin impuestos)", p["piezas"], "", p["total_estimado"]])
    ws[ws.max_row][1].font = Font(bold=True)
    for letra, ancho in zip("ABCDE", (16, 52, 10, 22, 16)):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = ws.cell(row=encabezado + 1, column=1)
    salida = BytesIO()
    wb.save(salida)
    nombre = f"pedido_{p['folio']}_{p['proveedor']}.xlsx"
    return Response(salida.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nombre)}"})
