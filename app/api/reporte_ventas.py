from datetime import date
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.api.entradas import _exacto
from app.core.auth import solo_admin
from app.core.database import get_db
from app.models import Usuario
from app.services import reporte_ventas
from app.services.errores import ERRORES_NEGOCIO, a_http

router = APIRouter(prefix="/reportes/ventas", tags=["reporte de ventas"])


def _datos(db: Session, usuario: Usuario, desde: date | None, hasta: date | None) -> dict:
    desde = desde or reporte_ventas.hoy()
    try:
        return reporte_ventas.reporte(db, usuario.negocio_id, desde, hasta or desde)
    except ERRORES_NEGOCIO as e:
        raise a_http(e)


@router.get("")
def reporte(desde: date | None = None, hasta: date | None = None, usuario: Usuario = Depends(solo_admin),
            db: Session = Depends(get_db)):
    """Ventas del día (sin fechas) o de un periodo."""
    return _exacto(_datos(db, usuario, desde, hasta))


@router.get("/excel")
def excel(desde: date | None = None, hasta: date | None = None, usuario: Usuario = Depends(solo_admin),
          db: Session = Depends(get_db)):
    r = _datos(db, usuario, desde, hasta)
    periodo = f"{r['desde']:%d/%m/%Y}" + ("" if r["desde"] == r["hasta"] else f" al {r['hasta']:%d/%m/%Y}")
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Ventas del {periodo}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for texto, valor in [
        ("Vendido (con impuestos)", r["vendido"]), ("Devoluciones y cambios", r["devuelto"]), ("Venta neta", r["venta_neta"]),
        ("Tickets", r["tickets"]), ("Ticket promedio", r["promedio_ticket"]), ("Piezas", r["piezas"]),
        ("Cobrado en efectivo", r["efectivo"]), ("Cobrado con tarjeta", r["tarjeta"]),
        ("Ventas canceladas", r["canceladas"]), ("Total cancelado", r["total_canceladas"]),
        ("Subtotal sin impuestos", r["subtotal"]), ("IVA", r["iva"]), ("IEPS", r["ieps"]),
        ("Ganancia estimada", r["ganancia_estimada"]),
    ]:
        ws.append([texto, valor])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16

    def hoja(titulo, encabezados, filas, anchos):
        h = wb.create_sheet(titulo)
        h.append(encabezados)
        for celda in h[1]:
            celda.font = Font(bold=True)
        for f in filas:
            h.append(f)
        for letra, ancho in zip("ABCD", anchos):
            h.column_dimensions[letra].width = ancho
        h.freeze_panes = "A2"

    hoja("Productos", ["Producto", "Piezas", "Vendido"], [[x["nombre"], x["piezas"], x["vendido"]] for x in r["productos"]], (50, 10, 14))
    hoja("Por categoría", ["Categoría", "Piezas", "Vendido"], [[x["grupo"], x["piezas"], x["vendido"]] for x in r["por_categoria"]], (34, 10, 14))
    hoja("Por cajero", ["Cajero", "Tickets", "Vendido"], [[x["grupo"], x["tickets"], x["vendido"]] for x in r["por_cajero"]], (30, 10, 14))
    hoja("Por caja", ["Caja", "Tickets", "Vendido"], [[x["grupo"], x["tickets"], x["vendido"]] for x in r["por_caja"]], (24, 10, 14))
    hoja("Por hora", ["Hora", "Tickets", "Vendido"], [[f"{x['grupo']:02d}:00", x["tickets"], x["vendido"]] for x in r["por_hora"]], (10, 10, 14))
    if r["desde"] != r["hasta"]:
        hoja("Por día", ["Día", "Tickets", "Vendido"], [[x["grupo"], x["tickets"], x["vendido"]] for x in r["por_dia"]], (14, 10, 14))
    salida = BytesIO()
    wb.save(salida)
    nombre = f"ventas_{r['desde']:%Y%m%d}" + ("" if r["desde"] == r["hasta"] else f"_{r['hasta']:%Y%m%d}") + ".xlsx"
    return Response(salida.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nombre)}"})
