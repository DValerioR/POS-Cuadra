from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.api.entradas import _exacto
from app.core.auth import usuario_actual
from app.core.database import get_db
from app.models import Producto, Usuario
from app.services import faltantes
from app.services.entradas import _validar_rol
from app.services.errores import ERRORES_NEGOCIO, NoEncontrado, a_http

router = APIRouter(tags=["faltantes y pedidos"])


@router.get("/reportes/faltantes")
def reporte(proveedor_id: int, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    """Faltantes para pedir a un proveedor (admin y bodega): lo que ese
    proveedor ha surtido y está en su mínimo o menos, con su último costo y
    quién tiene el mejor último costo; aparte, los faltantes sin proveedor."""
    try:
        return _exacto(faltantes.reporte(db, usuario, proveedor_id))
    except ERRORES_NEGOCIO as e:
        raise a_http(e)


@router.get("/reportes/faltantes/excel")
def reporte_excel(proveedor_id: int, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    try:
        r = faltantes.reporte(db, usuario, proveedor_id)
    except ERRORES_NEGOCIO as e:
        raise a_http(e)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pedido"
    negrita = Font(bold=True)
    amarillo = PatternFill("solid", fgColor="FFF4D6")
    ws.append([f"Faltantes para {r['proveedor']}", f"{r['fecha']:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Clave", "Producto", "Existencia", "Mínimo", "Máximo", "Pedir", "Último costo aquí", "Fecha",
               "Mejor precio con", "Su costo", "Fecha", "Diferencia %"])
    for celda in ws[3]:
        celda.font = negrita
    for x in r["del_proveedor"]:
        ws.append([x["clave"], x["nombre"], x["existencia"], x["minimo"], x["maximo"], x["sugerido"],
                   x["costo"], x["fecha_costo"], x["mejor_proveedor"], x["mejor_costo"], x["mejor_fecha"],
                   x["diferencia_porcentaje"]])
        if x["mejor_proveedor"]:
            for celda in ws[ws.max_row]:
                celda.fill = amarillo
    ws.append([])
    ws.append(["", "Total estimado (sin impuestos)", "", "", "", "", r["total_estimado"]])
    ws[ws.max_row][1].font = negrita
    for letra, ancho in zip("ABCDEFGHIJKL", (16, 50, 11, 9, 9, 8, 16, 12, 22, 10, 12, 12)):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "A4"
    otra = wb.create_sheet("Sin proveedor registrado")
    otra.append(["Clave", "Producto", "Existencia", "Mínimo", "Máximo", "Pedir"])
    for celda in otra[1]:
        celda.font = negrita
    for x in r["sin_proveedor"]:
        otra.append([x["clave"], x["nombre"], x["existencia"], x["minimo"], x["maximo"], x["sugerido"]])
    for letra, ancho in zip("ABCDEF", (16, 50, 11, 9, 9, 8)):
        otra.column_dimensions[letra].width = ancho
    otra.freeze_panes = "A2"
    salida = BytesIO()
    wb.save(salida)
    nombre = f"faltantes_{r['proveedor']}_{r['fecha']:%Y%m%d}.xlsx"
    return Response(salida.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nombre)}"})


@router.get("/productos/{producto_id}/proveedores")
def proveedores_del_producto(producto_id: int, usuario: Usuario = Depends(usuario_actual), db: Session = Depends(get_db)):
    """Proveedores que han surtido el producto, con su último costo por pieza
    (el más barato primero)."""
    try:
        _validar_rol(usuario)
        p = db.get(Producto, producto_id)
        if p is None or p.negocio_id != usuario.negocio_id:
            raise NoEncontrado("Producto no encontrado")
    except ERRORES_NEGOCIO as e:
        raise a_http(e)
    return _exacto(faltantes.ultimos_costos(db, usuario.negocio_id, {producto_id}).get(producto_id, []))
