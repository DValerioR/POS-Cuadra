from datetime import date
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Body, Depends, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.api.entradas import _exacto
from app.asistente.consultas import hoy
from app.core.auth import solo_admin
from app.core.database import get_db
from app.models import Usuario
from app.services import ofertas
from app.services.errores import ERRORES_NEGOCIO, a_http

router = APIRouter(prefix="/reportes/ofertas", tags=["ofertas"])


def _sin_internos(c: dict) -> dict:
    return {k: v for k, v in c.items() if k != "factor_impuestos"}


@router.get("/candidatos")
def candidatos(usuario: Usuario = Depends(solo_admin), db: Session = Depends(get_db)):
    """Productos que conviene poner en oferta (consultas exactas, sin IA)."""
    datos = ofertas.candidatos(db, usuario.negocio_id)
    datos["candidatos"] = [_sin_internos(c) for c in datos["candidatos"]]
    return _exacto({**datos, "maximo_a_la_ia": ofertas.MAXIMO_A_LA_IA})


@router.post("")
def recomendar(usuario: Usuario = Depends(solo_admin), db: Session = Depends(get_db)):
    """Ofertas que propone el asistente para los candidatos más urgentes, ya
    revisadas por el código. Solo recomienda: no cambia ningún precio."""
    try:
        resultado = ofertas.recomendar(db, usuario.negocio_id)
    except ERRORES_NEGOCIO as e:
        db.rollback()
        raise a_http(e)
    db.commit()
    return _exacto(resultado)


def _texto_oferta(o: dict) -> str:
    if o["tipo"] in ("2x1", "3x2"):
        return f"{o['tipo']} (paga ${o['precio_oferta']} por {o['piezas']})"
    if o["tipo"] == "paquete":
        return f"Paquete con {o['paquete_con']['nombre']} a ${o['precio_oferta']}"
    return f"{ofertas.TIPOS[o['tipo']]}: ${o['precio_oferta']} (−{o['descuento_porcentaje']}%)"


@router.post("/excel")
def excel(propuestas: list[dict] = Body(default=[], embed=True, alias="ofertas"),
          usuario: Usuario = Depends(solo_admin), db: Session = Depends(get_db)):
    """Excel con las ofertas que se le enviaron (las que regresó el asistente
    en pantalla, para no volver a pagar la consulta) y, en otra hoja, todos
    los candidatos calculados de nuevo."""
    datos = ofertas.candidatos(db, usuario.negocio_id)
    wb = Workbook()
    negrita = Font(bold=True)
    amarillo = PatternFill("solid", fgColor="FFF4D6")

    ws = wb.active
    ws.title = "Ofertas sugeridas"
    ws.append([f"Ofertas sugeridas por el asistente · {hoy():%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Solo son recomendaciones: ningún precio se cambió. En amarillo, las que el sistema subió al precio mínimo."])
    ws.append([])
    ws.append(["Clave", "Producto", "Por qué", "Existencia", "Caducidad", "Caducan", "Vendidas 12 sem.", "Costo pieza",
               "Precio normal", "Oferta", "Precio oferta", "Descuento %", "Ganancia (sin imp.)", "Motivo del asistente"])
    for celda in ws[4]:
        celda.font = negrita
    for o in propuestas:
        try:
            ws.append([o["clave"], o["nombre"], ", ".join(ofertas.RAZONES.get(r, r) for r in o["razones"]),
                       float(o["existencia"]), date.fromisoformat(o["caducidad"]) if o["caducidad"] else "", float(o["piezas_por_caducar"]) if o["piezas_por_caducar"] else "",
                       o["vendidas_12_semanas"], float(o["costo_pieza"]), float(o["precio_normal"]), _texto_oferta(o),
                       float(o["precio_oferta"]), float(o["descuento_porcentaje"]), float(o["ganancia"]), o["motivo"]])
        except (KeyError, TypeError, ValueError):
            continue  # renglón mal formado: se omite
        if o.get("ajustada"):
            for celda in ws[ws.max_row]:
                celda.fill = amarillo
    for letra, ancho in zip("ABCDEFGHIJKLMN", (16, 45, 28, 11, 12, 9, 10, 11, 12, 40, 12, 11, 13, 55)):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "A5"

    h = wb.create_sheet("Todos los candidatos")
    h.append(["Clave", "Producto", "Categoría", "Receta", "Por qué", "Existencia", "Caducidad", "Caducan",
              "Vendidas 12 sem.", "Costo pieza", "Precio", "Margen %", "Precio mínimo de oferta", "Dinero parado (costo)"])
    for celda in h[1]:
        celda.font = negrita
    for c in datos["candidatos"]:
        h.append([c["clave"], c["nombre"], c["categoria"] or "", "Sí" if c["requiere_receta"] else "",
                  ", ".join(ofertas.RAZONES[r] for r in c["razones"]), c["existencia"], c["caducidad"],
                  c["piezas_por_caducar"], c["vendidas_12_semanas"], c["costo_pieza"], c["precio"], c["margen"],
                  c["precio_minimo"], c["valor_inventario"]])
    for letra, ancho in zip("ABCDEFGHIJKLMN", (16, 45, 18, 7, 28, 11, 12, 9, 10, 11, 10, 10, 14, 14)):
        h.column_dimensions[letra].width = ancho
    h.freeze_panes = "A2"

    salida = BytesIO()
    wb.save(salida)
    nombre = f"ofertas_{hoy():%Y%m%d}.xlsx"
    return Response(salida.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nombre)}"})
